import re
from datetime import datetime
import secrets

import os
import sqlite3
import json
from openpyxl import load_workbook

from flask import (
    Flask,
    render_template,
    redirect,
    url_for,
    request,
    session,
    jsonify,
    send_file
)

from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename


# =========================================================
# FLASK
# =========================================================

app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "darabilim-dev-key-change-later"
)

# =========================================================
# COLAB DEVELOPMENT MODE
# =========================================================
# True  — Colab-та логинсіз әзірлеу
# False — қалыпты логин жүйесі
app.config["DARABILIM_DEV_MODE"] = True


# =========================================================
# PATH
# =========================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DB_PATH = os.path.join(
    BASE_DIR,
    "database",
    "darabilim.db"
)


# =========================================================
# УАҚЫТША МҰҒАЛІМ АККАУНТЫ
# Кейін PostgreSQL базасына көшіреміз
# =========================================================

DEMO_USER = {
    "username": "teacher",
    "password_hash": generate_password_hash("12345"),
    "name": "Мұғалім"
}


# =========================================================
# DATABASE
# =========================================================

def get_db():

    conn = sqlite3.connect(DB_PATH)

    conn.row_factory = sqlite3.Row

    return conn


def init_db():

    os.makedirs(
        os.path.dirname(DB_PATH),
        exist_ok=True
    )

    conn = get_db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS students (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            full_name TEXT NOT NULL,

            class_name TEXT NOT NULL DEFAULT '4 «Д»',

            parent_name TEXT,

            parent_phone TEXT,

            telegram TEXT,

            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    conn.commit()
    conn.close()


init_db()


# =========================================================
# LOGIN CHECK
# =========================================================

def logged_in():

    # ========================================================
    # DARABILIM_LOGIN_PASSWORD_RESTORE_FINAL_V1
    # Production authentication:
    # only a real authenticated session is accepted.
    # ========================================================

    return "user" in session


# =========================================================
# ROOT
# =========================================================

@app.route("/")
def index():

    if logged_in():
        return redirect(url_for("dashboard"))

    return redirect(url_for("login"))


# =========================================================
# LOGIN
# =========================================================

@app.route("/login", methods=["GET", "POST"])
def login():

    if logged_in():
        return redirect(url_for("dashboard"))

    error = None

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        if (
            username == DEMO_USER["username"]
            and
            check_password_hash(
                DEMO_USER["password_hash"],
                password
            )
        ):

            session["user"] = username

            session["teacher_name"] = DEMO_USER["name"]

            return redirect(
                url_for("dashboard")
            )

        error = "Логин немесе құпиясөз қате."

    return render_template(
        "login.html",
        error=error
    )


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/dashboard")
def dashboard():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    student_count = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM students
        WHERE class_name = ?
        """,
        ('4 «Д»',)
    ).fetchone()["total"]

    conn.close()

    return render_template(
        "dashboard.html",

        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),

        student_count=student_count
    )


# =========================================================
# STUDENTS
# =========================================================

@app.route("/students")
def students():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    students_list = conn.execute(
        """
        SELECT *
        FROM students
        WHERE class_name = ?
        ORDER BY full_name COLLATE NOCASE
        """,
        ('4 «Д»',)
    ).fetchall()

    conn.close()

    return render_template(
        "students.html",

        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),

        students=students_list
    )


# =========================================================
# ADD STUDENT
# =========================================================

@app.route(
    "/students/add",
    methods=["POST"]
)
def add_student():

    if not logged_in():
        return redirect(url_for("login"))

    full_name = request.form.get(
        "full_name",
        ""
    ).strip()

    parent_name = request.form.get(
        "parent_name",
        ""
    ).strip()

    parent_phone = request.form.get(
        "parent_phone",
        ""
    ).strip()

    telegram = request.form.get(
        "telegram",
        ""
    ).strip()

    if not full_name:
        return redirect(
            url_for("students")
        )

    conn = get_db()

    conn.execute(
        """
        INSERT INTO students
        (
            full_name,
            class_name,
            parent_name,
            parent_phone,
            telegram
        )

        VALUES (?, ?, ?, ?, ?)
        """,

        (
            full_name,
            '4 «Д»',
            parent_name,
            parent_phone,
            telegram
        )
    )

    conn.commit()
    conn.close()

    return redirect(
        url_for("students")
    )


# =========================================================
# DELETE STUDENT
# =========================================================

@app.route(
    "/students/<int:student_id>/delete",
    methods=["POST"]
)
def delete_student(student_id):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    conn.execute(
        """
        DELETE FROM students
        WHERE id = ?
        """,
        (student_id,)
    )

    conn.commit()
    conn.close()

    return redirect(
        url_for("students")
    )




# =========================================================
# EXCEL АРҚЫЛЫ ОҚУШЫЛАРДЫ ЖАППАЙ ЖҮКТЕУ
# =========================================================

@app.route("/students/import", methods=["POST"])
def import_students():

    if not logged_in():
        return redirect(url_for("login"))

    uploaded_file = request.files.get("excel_file")

    if not uploaded_file or uploaded_file.filename == "":
        return redirect(url_for("students"))

    if not uploaded_file.filename.lower().endswith(".xlsx"):
        return redirect(url_for("students"))

    try:
        workbook = load_workbook(
            uploaded_file,
            data_only=True
        )

        sheet = workbook.active

        conn = get_db()

        added = 0

        # 1-қатар тақырыптар деп есептеледі
        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            if not row:
                continue

            full_name = (
                str(row[0]).strip()
                if len(row) > 0 and row[0]
                else ""
            )

            parent_name = (
                str(row[1]).strip()
                if len(row) > 1 and row[1]
                else ""
            )

            parent_phone = (
                str(row[2]).strip()
                if len(row) > 2 and row[2]
                else ""
            )

            telegram = (
                str(row[3]).strip()
                if len(row) > 3 and row[3]
                else ""
            )

            if not full_name:
                continue

            # Бір оқушыны екі рет енгізбеу
            existing = conn.execute(
                """
                SELECT id
                FROM students
                WHERE full_name = ?
                AND class_name = ?
                """,
                (
                    full_name,
                    '4 «Д»'
                )
            ).fetchone()

            if existing:
                continue

            conn.execute(
                """
                INSERT INTO students
                (
                    full_name,
                    class_name,
                    parent_name,
                    parent_phone,
                    telegram
                )
                VALUES (?, ?, ?, ?, ?)
                """,
                (
                    full_name,
                    '4 «Д»',
                    parent_name,
                    parent_phone,
                    telegram
                )
            )

            added += 1

        conn.commit()
        conn.close()

        print(
            f"✅ Excel арқылы {added} оқушы енгізілді"
        )

    except Exception as e:

        print(
            "❌ Excel импорт қатесі:",
            str(e)
        )

    return redirect(
        url_for("students")
    )


# =========================================================
# LOGOUT
# =========================================================


# ============================================================
# SETTINGS
# ============================================================

@app.route(
    "/settings",
    methods=["GET", "POST"]
)
def settings():

    if not logged_in():
        return redirect(
            url_for("login")
        )

    conn = get_db()

    if request.method == "POST":

        values = {
            "teacher_name":
                request.form.get(
                    "teacher_name",
                    ""
                ).strip()
                or "Мұғалім",

            "school_name":
                request.form.get(
                    "school_name",
                    ""
                ).strip(),

            "city_name":
                request.form.get(
                    "city_name",
                    ""
                ).strip(),

            "default_class":
                request.form.get(
                    "default_class",
                    ""
                ).strip(),

            "default_subject":
                request.form.get(
                    "default_subject",
                    "Математика"
                ).strip()
        }

        for key, value in values.items():

            conn.execute(
                """
                INSERT INTO app_settings (
                    setting_key,
                    setting_value,
                    updated_at
                )
                VALUES (
                    ?,
                    ?,
                    CURRENT_TIMESTAMP
                )

                ON CONFLICT(setting_key)
                DO UPDATE SET

                    setting_value =
                        excluded.setting_value,

                    updated_at =
                        CURRENT_TIMESTAMP
                """,
                (
                    key,
                    value
                )
            )

        conn.commit()

        session["teacher_name"] = (
            values["teacher_name"]
        )

        conn.close()

        return redirect(
            url_for(
                "settings",
                saved=1
            )
        )


    rows = conn.execute(
        """
        SELECT
            setting_key,
            setting_value

        FROM app_settings
        """
    ).fetchall()

    conn.close()

    current_settings = {
        row["setting_key"]:
            row["setting_value"]
        for row in rows
    }

    return render_template(
        "settings.html",

        teacher_name=session.get(
            "teacher_name",
            current_settings.get(
                "teacher_name",
                "Мұғалім"
            )
        ),

        settings=current_settings,

        saved=(
            request.args.get(
                "saved"
            ) == "1"
        )
    )



# ============================================================
# DIAGNOSTIC CRUD - EDIT
# ============================================================

@app.route(
    "/diagnostics/<int:diagnostic_id>/edit",
    methods=["POST"]
)
def edit_diagnostic(diagnostic_id):

    if not logged_in():
        return redirect(
            url_for("login")
        )

    title = request.form.get(
        "title",
        ""
    ).strip()

    subject = request.form.get(
        "subject",
        ""
    ).strip()

    question_count_raw = request.form.get(
        "question_count",
        ""
    ).strip()

    if not title:

        return redirect(
            url_for("diagnostics")
        )

    try:

        question_count = int(
            question_count_raw
        )

    except (TypeError, ValueError):

        question_count = 0


    conn = get_db()

    diagnostic = conn.execute(
        """
        SELECT id
        FROM diagnostics
        WHERE id = ?
        """,
        (diagnostic_id,)
    ).fetchone()


    if diagnostic is None:

        conn.close()

        return redirect(
            url_for("diagnostics")
        )


    conn.execute(
        """
        UPDATE diagnostics

        SET
            title = ?,
            subject = ?,
            question_count = ?

        WHERE id = ?
        """,
        (
            title,
            subject,
            question_count,
            diagnostic_id
        )
    )

    conn.commit()
    conn.close()

    return redirect(
        url_for("diagnostics")
    )


# ============================================================
# DIAGNOSTIC CRUD - DELETE
# ============================================================

@app.route(
    "/diagnostics/<int:diagnostic_id>/delete",
    methods=["POST"]
)
def delete_diagnostic(diagnostic_id):

    if not logged_in():
        return redirect(
            url_for("login")
        )

    conn = get_db()


    # Осы бақылауға қатысты attempt ID
    attempts = conn.execute(
        """
        SELECT id
        FROM diagnostic_attempts
        WHERE diagnostic_id = ?
        """,
        (diagnostic_id,)
    ).fetchall()


    attempt_ids = [
        row["id"]
        for row in attempts
    ]


    # Attempt-терге қатысты dependent data
    for attempt_id in attempt_ids:

        conn.execute(
            """
            DELETE FROM personal_tasks
            WHERE source_attempt_id = ?
            """,
            (attempt_id,)
        )

        conn.execute(
            """
            DELETE FROM objective_results
            WHERE attempt_id = ?
            """,
            (attempt_id,)
        )

        conn.execute(
            """
            DELETE FROM diagnostic_answers
            WHERE attempt_id = ?
            """,
            (attempt_id,)
        )


    # Attempts
    conn.execute(
        """
        DELETE FROM diagnostic_attempts
        WHERE diagnostic_id = ?
        """,
        (diagnostic_id,)
    )


    # Questions
    conn.execute(
        """
        DELETE FROM diagnostic_questions
        WHERE diagnostic_id = ?
        """,
        (diagnostic_id,)
    )


    # Objective links
    conn.execute(
        """
        DELETE FROM diagnostic_objectives
        WHERE diagnostic_id = ?
        """,
        (diagnostic_id,)
    )


    # Diagnostic
    conn.execute(
        """
        DELETE FROM diagnostics
        WHERE id = ?
        """,
        (diagnostic_id,)
    )


    conn.commit()
    conn.close()

    return redirect(
        url_for("diagnostics")
    )



@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )




# =========================================================
# DEVELOPMENT LOGIN
# Тек әзірлеу кезінде
# =========================================================

@app.route("/dev-login")
def dev_login():

    # Development login bypass disabled.
    # Always use the normal username/password login.

    session.clear()

    return redirect(
        url_for("login")
    )


# =========================================================
# МАТЕРИАЛДАР
# =========================================================

@app.route("/materials")
def materials():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    materials_list = conn.execute(
        """
        SELECT *
        FROM materials
        ORDER BY uploaded_at DESC
        """
    ).fetchall()

    material_status = {
        "math_book": None,
        "kazakh_book": None,
        "math_ktj": None,
        "kazakh_ktj": None
    }

    for item in materials_list:

        subject = item["subject"]
        material_type = item["material_type"]

        if (
            subject == "Математика"
            and material_type == "Оқулық"
        ):
            material_status["math_book"] = item

        elif (
            subject == "Қазақ тілі"
            and material_type == "Оқулық"
        ):
            material_status["kazakh_book"] = item

        elif (
            subject == "Математика"
            and material_type == "КТЖ"
        ):
            material_status["math_ktj"] = item

        elif (
            subject == "Қазақ тілі"
            and material_type == "КТЖ"
        ):
            material_status["kazakh_ktj"] = item

    conn.close()

    return render_template(
        "materials.html",
        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),
        materials=materials_list,
        material_status=material_status
    )


@app.route(
    "/materials/upload",
    methods=["POST"]
)
def upload_material():

    if not logged_in():
        return redirect(url_for("login"))

    title = request.form.get(
        "title",
        ""
    ).strip()

    subject = request.form.get(
        "subject",
        ""
    ).strip()

    material_type = request.form.get(
        "material_type",
        ""
    ).strip()

    uploaded_file = request.files.get(
        "material_file"
    )

    if (
        not title
        or not subject
        or not material_type
        or not uploaded_file
        or uploaded_file.filename == ""
    ):
        return redirect(
            url_for("materials")
        )

    extension = os.path.splitext(
        uploaded_file.filename
    )[1].lower()

    allowed = {
        ".pdf",
        ".docx",
        ".xlsx"
    }

    # Кеңейтімі жоқ бұрынғы DOCX секілді
    # файлдарға да мүмкіндік береміз.
    if extension and extension not in allowed:
        return redirect(
            url_for("materials")
        )

    safe_name = secure_filename(
        uploaded_file.filename
    )

    if not safe_name:
        safe_name = "material"

    upload_dir = os.path.join(
        BASE_DIR,
        "uploads",
        "materials"
    )

    os.makedirs(
        upload_dir,
        exist_ok=True
    )

    final_name = safe_name

    base_name = os.path.splitext(
        safe_name
    )[0]

    counter = 1

    while os.path.exists(
        os.path.join(
            upload_dir,
            final_name
        )
    ):

        final_name = (
            f"{base_name}_{counter}"
            f"{extension}"
        )

        counter += 1

    save_path = os.path.join(
        upload_dir,
        final_name
    )

    uploaded_file.save(
        save_path
    )

    conn = get_db()

    conn.execute(
        """
        INSERT INTO materials
        (
            title,
            subject,
            material_type,
            filename,
            file_path,
            processing_status,
            objectives_extracted,
            topics_extracted
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            title,
            subject,
            material_type,
            final_name,
            save_path,
            "uploaded",
            0,
            0
        )
    )

    conn.commit()
    conn.close()

    return redirect(
        url_for("materials")
    )


# =========================================================
# КІРІС БАҚЫЛАУ
# =========================================================

@app.route("/diagnostics")
def diagnostics():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    diagnostics_list = conn.execute(
        """
        SELECT *
        FROM diagnostics
        ORDER BY created_at DESC
        """
    ).fetchall()

    student_count = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM students
        WHERE class_name = ?
        """,
        ('4 «Д»',)
    ).fetchone()["total"]

    conn.close()

    return render_template(
        "diagnostics.html",
        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),
        diagnostics=diagnostics_list,
        student_count=student_count
    )


@app.route(
    "/diagnostics/create",
    methods=["POST"]
)
def create_diagnostic():

    if not logged_in():
        return redirect(url_for("login"))


    # --------------------------------------------------------
    # BASIC FORM DATA
    # --------------------------------------------------------

    title = request.form.get(
        "title",
        ""
    ).strip()


    subject = request.form.get(
        "subject",
        ""
    ).strip()


    objective_ids = request.form.getlist(
        "objective_ids"
    )


    # --------------------------------------------------------
    # CONTROL TYPE
    # --------------------------------------------------------

    diagnostic_type = (
        request.form.get(
            "diagnostic_type",
            "entry"
        )
        .strip()
        .lower()
    )


    if diagnostic_type not in {
        "entry",
        "intermediate",
        "exit",
    }:

        diagnostic_type = "entry"


    # --------------------------------------------------------
    # ACADEMIC YEAR
    # --------------------------------------------------------

    academic_year = (
        request.form.get(
            "academic_year",
            ""
        )
        .strip()
    )


    if not re.fullmatch(
        r"\d{4}-\d{4}",
        academic_year
    ):

        academic_year = (
            darabilim_current_academic_year()
            if "darabilim_current_academic_year"
               in globals()
            else ""
        )


    if academic_year:

        try:

            year_start, year_end = (
                academic_year.split("-")
            )


            if (
                int(year_end)
                != int(year_start) + 1
            ):

                raise ValueError


        except (
            ValueError,
            TypeError
        ):

            academic_year = ""


    if not academic_year:

        now = datetime.now()


        if now.month >= 8:

            academic_year = (
                f"{now.year}-"
                f"{now.year + 1}"
            )

        else:

            academic_year = (
                f"{now.year - 1}-"
                f"{now.year}"
            )


    # --------------------------------------------------------
    # QUESTION COUNT
    # --------------------------------------------------------

    try:

        question_count = int(
            request.form.get(
                "question_count",
                10
            )
        )

    except (
        TypeError,
        ValueError
    ):

        question_count = 10


    question_count = max(
        1,
        min(
            question_count,
            50
        )
    )


    # --------------------------------------------------------
    # VALIDATION
    # --------------------------------------------------------

    if (
        not title
        or subject not in (
            "Математика",
            "Қазақ тілі"
        )
    ):

        return redirect(
            url_for("diagnostics")
        )


    # --------------------------------------------------------
    # OBJECTIVE IDS
    # --------------------------------------------------------

    clean_ids = []


    for value in objective_ids:

        try:

            objective_id = int(
                value
            )


            if objective_id > 0:

                clean_ids.append(
                    objective_id
                )


        except (
            TypeError,
            ValueError
        ):

            pass


    clean_ids = list(
        dict.fromkeys(
            clean_ids
        )
    )


    conn = get_db()


    valid_ids = []


    if clean_ids:

        placeholders = ",".join(
            "?"
            for _ in clean_ids
        )


        query = f"""
            SELECT id

            FROM learning_objectives

            WHERE
                id IN ({placeholders})
                AND class_name = ?
                AND subject = ?
        """


        rows = conn.execute(
            query,
            (
                *clean_ids,
                '4 «Д»',
                subject
            )
        ).fetchall()


        valid_ids = [
            row["id"]
            for row in rows
        ]


    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    summary = (
        f"{len(valid_ids)} оқу мақсаты таңдалды"
        if valid_ids
        else "Оқу мақсаты таңдалмаған"
    )


    # --------------------------------------------------------
    # CREATE DIAGNOSTIC
    # --------------------------------------------------------

    cursor = conn.execute(
        """
        INSERT INTO diagnostics
        (
            title,
            class_name,
            subject,
            learning_objective,
            question_count,
            status,
            diagnostic_type,
            academic_year
        )

        VALUES (
            ?, ?, ?, ?, ?,
            ?, ?, ?
        )
        """,
        (
            title,
            '4 «Д»',
            subject,
            summary,
            question_count,
            "draft",
            diagnostic_type,
            academic_year
        )
    )


    diagnostic_id = (
        cursor.lastrowid
    )


    # --------------------------------------------------------
    # OBJECTIVE LINKS
    # --------------------------------------------------------

    for objective_id in valid_ids:

        conn.execute(
            """
            INSERT OR IGNORE INTO diagnostic_objectives
            (
                diagnostic_id,
                learning_objective_id
            )

            VALUES (?, ?)
            """,
            (
                diagnostic_id,
                objective_id
            )
        )


    conn.commit()
    conn.close()


    return redirect(
        url_for("diagnostics")
    )


@app.route(
    "/api/learning-objectives"
)
def api_learning_objectives():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "unauthorized"
        }), 401

    subject = request.args.get(
        "subject",
        ""
    ).strip()

    if subject not in (
        "Математика",
        "Қазақ тілі"
    ):

        return jsonify({
            "ok": False,
            "error": "invalid_subject",
            "items": []
        }), 400

    conn = get_db()

    rows = conn.execute(
        """
        SELECT
            MIN(id) AS id,
            objective_code,
            objective_text,
            COUNT(*) AS lesson_count
        FROM learning_objectives
        WHERE
            class_name = ?
            AND subject = ?
        GROUP BY
            objective_code,
            objective_text
        ORDER BY
            objective_code,
            objective_text
        """,
        (
            '4 «Д»',
            subject
        )
    ).fetchall()

    conn.close()

    items = [
        {
            "id": row["id"],
            "code": (
                row["objective_code"]
                or ""
            ),
            "text": row[
                "objective_text"
            ],
            "lesson_count": row[
                "lesson_count"
            ]
        }
        for row in rows
    ]

    return jsonify({
        "ok": True,
        "subject": subject,
        "count": len(items),
        "items": items
    })


# =========================================================
# КІРІС БАҚЫЛАУ — ЖЕКЕ БЕТ
# =========================================================

@app.route(
    "/diagnostics/<int:diagnostic_id>"
)
def diagnostic_detail(
    diagnostic_id
):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    diagnostic = conn.execute(
        """
        SELECT *
        FROM diagnostics
        WHERE id = ?
        """,
        (diagnostic_id,)
    ).fetchone()

    if diagnostic is None:

        conn.close()

        return (
            "Кіріс бақылау табылмады",
            404
        )

    objectives = conn.execute(
        """
        SELECT
            lo.id,
            lo.objective_code,
            lo.objective_text,
            lo.section,
            lo.lesson_topic
        FROM diagnostic_objectives d
        JOIN learning_objectives lo
            ON lo.id =
               d.learning_objective_id
        WHERE d.diagnostic_id = ?
        ORDER BY
            lo.objective_code,
            lo.id
        """,
        (diagnostic_id,)
    ).fetchall()

    questions = conn.execute(
        """
        SELECT
            q.*,
            lo.objective_code,
            lo.objective_text
        FROM diagnostic_questions q
        JOIN learning_objectives lo
            ON lo.id =
               q.learning_objective_id
        WHERE q.diagnostic_id = ?
        ORDER BY
            q.question_order,
            q.id
        """,
        (diagnostic_id,)
    ).fetchall()

    students = conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name
        FROM students
        WHERE class_name = ?
        ORDER BY full_name
        """,
        ('4 «Д»',)
    ).fetchall()

    conn.close()

    return render_template(
        "diagnostic_detail.html",
        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),
        diagnostic=diagnostic,
        objectives=objectives,
        questions=questions,
        students=students
    )


# =========================================================
# ТАПСЫРМА ҚОСУ
# =========================================================

@app.route(
    "/diagnostics/<int:diagnostic_id>/questions/add",
    methods=["POST"]
)
def add_diagnostic_question(
    diagnostic_id
):

    if not logged_in():
        return redirect(url_for("login"))

    try:
        objective_id = int(
            request.form.get(
                "learning_objective_id",
                0
            )
        )
    except (TypeError, ValueError):
        objective_id = 0

    level = request.form.get(
        "level",
        "Бекіту"
    ).strip()

    question_text = request.form.get(
        "question_text",
        ""
    ).strip()

    option_a = request.form.get(
        "option_a",
        ""
    ).strip()

    option_b = request.form.get(
        "option_b",
        ""
    ).strip()

    option_c = request.form.get(
        "option_c",
        ""
    ).strip()

    option_d = request.form.get(
        "option_d",
        ""
    ).strip()

    correct_answer = request.form.get(
        "correct_answer",
        ""
    ).strip().upper()

    explanation = request.form.get(
        "explanation",
        ""
    ).strip()

    try:
        points = int(
            request.form.get(
                "points",
                1
            )
        )
    except (TypeError, ValueError):
        points = 1

    if level not in (
        "Қолдау",
        "Бекіту",
        "Дамыту"
    ):
        level = "Бекіту"

    if correct_answer not in (
        "A",
        "B",
        "C",
        "D"
    ):
        return redirect(
            url_for(
                "diagnostic_detail",
                diagnostic_id=diagnostic_id
            )
        )

    points = max(
        1,
        min(points, 10)
    )

    if (
        not objective_id
        or not question_text
    ):
        return redirect(
            url_for(
                "diagnostic_detail",
                diagnostic_id=diagnostic_id
            )
        )

    conn = get_db()

    valid = conn.execute(
        """
        SELECT 1
        FROM diagnostic_objectives
        WHERE
            diagnostic_id = ?
            AND learning_objective_id = ?
        """,
        (
            diagnostic_id,
            objective_id
        )
    ).fetchone()

    if valid is None:

        conn.close()

        return redirect(
            url_for(
                "diagnostic_detail",
                diagnostic_id=diagnostic_id
            )
        )

    next_order = conn.execute(
        """
        SELECT
            COALESCE(
                MAX(question_order),
                0
            ) + 1 AS n
        FROM diagnostic_questions
        WHERE diagnostic_id = ?
        """,
        (diagnostic_id,)
    ).fetchone()["n"]

    conn.execute(
        """
        INSERT INTO diagnostic_questions
        (
            diagnostic_id,
            learning_objective_id,
            level,
            question_type,
            question_text,
            option_a,
            option_b,
            option_c,
            option_d,
            correct_answer,
            explanation,
            points,
            question_order
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            diagnostic_id,
            objective_id,
            level,
            "multiple_choice",
            question_text,
            option_a,
            option_b,
            option_c,
            option_d,
            correct_answer,
            explanation,
            points,
            next_order
        )
    )

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "diagnostic_detail",
            diagnostic_id=diagnostic_id
        )
    )


# =========================================================
# ТАПСЫРМАНЫ ӨШІРУ
# =========================================================

@app.route(
    "/diagnostics/<int:diagnostic_id>/questions/<int:question_id>/delete",
    methods=["POST"]
)
def delete_diagnostic_question(
    diagnostic_id,
    question_id
):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    conn.execute(
        """
        DELETE FROM diagnostic_questions
        WHERE
            id = ?
            AND diagnostic_id = ?
        """,
        (
            question_id,
            diagnostic_id
        )
    )

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "diagnostic_detail",
            diagnostic_id=diagnostic_id
        )
    )


# =========================================================
# АВТОМАТТЫ ТАПСЫРМА ҚАҢҚАСЫ
# =========================================================

@app.route(
    "/diagnostics/<int:diagnostic_id>/questions/auto-generate",
    methods=["POST"]
)
def auto_generate_diagnostic_questions(
    diagnostic_id
):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    diagnostic = conn.execute(
        """
        SELECT *
        FROM diagnostics
        WHERE id = ?
        """,
        (diagnostic_id,)
    ).fetchone()

    if diagnostic is None:

        conn.close()

        return (
            "Кіріс бақылау табылмады",
            404
        )

    objectives = conn.execute(
        """
        SELECT
            lo.id,
            lo.objective_code,
            lo.objective_text,
            lo.lesson_topic
        FROM diagnostic_objectives d
        JOIN learning_objectives lo
            ON lo.id =
               d.learning_objective_id
        WHERE d.diagnostic_id = ?
        ORDER BY lo.id
        """,
        (diagnostic_id,)
    ).fetchall()

    if not objectives:

        conn.close()

        return redirect(
            url_for(
                "diagnostic_detail",
                diagnostic_id=diagnostic_id
            )
        )

    existing = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM diagnostic_questions
        WHERE diagnostic_id = ?
        """,
        (diagnostic_id,)
    ).fetchone()["total"]

    if existing > 0:

        conn.close()

        return redirect(
            url_for(
                "diagnostic_detail",
                diagnostic_id=diagnostic_id
            )
        )

    levels = [
        ("Қолдау", 1),
        ("Бекіту", 1),
        ("Дамыту", 2)
    ]

    order_number = 1
    created = 0

    max_questions = diagnostic[
        "question_count"
    ]

    for objective in objectives:

        for level, points in levels:

            if created >= max_questions:
                break

            code_value = (
                objective["objective_code"]
                or "Оқу мақсаты"
            )

            text_value = (
                objective["objective_text"]
                or ""
            )

            topic = (
                objective["lesson_topic"]
                or ""
            )

            if level == "Қолдау":

                question_text = (
                    f"{code_value} оқу мақсаты "
                    f"бойынша негізгі дағдыны "
                    f"тексеретін тапсырма: "
                    f"{text_value}"
                )

            elif level == "Бекіту":

                question_text = (
                    f"{code_value} оқу мақсатын "
                    f"стандартты жағдайда қолдануға "
                    f"арналған тапсырма: "
                    f"{text_value}"
                )

            else:

                question_text = (
                    f"{code_value} оқу мақсаты "
                    f"бойынша білімді жаңа жағдайда "
                    f"қолданатын тапсырма: "
                    f"{text_value}"
                )

            if topic:
                question_text += (
                    f" Сабақ тақырыбы: "
                    f"{topic}."
                )

            conn.execute(
                """
                INSERT INTO diagnostic_questions
                (
                    diagnostic_id,
                    learning_objective_id,
                    level,
                    question_type,
                    question_text,
                    option_a,
                    option_b,
                    option_c,
                    option_d,
                    correct_answer,
                    explanation,
                    points,
                    question_order
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    diagnostic_id,
                    objective["id"],
                    level,
                    "draft_template",
                    question_text,
                    "",
                    "",
                    "",
                    "",
                    "A",
                    (
                        "API-сыз дайындалған "
                        "тапсырма қаңқасы."
                    ),
                    points,
                    order_number
                )
            )

            order_number += 1
            created += 1

        if created >= max_questions:
            break

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "diagnostic_detail",
            diagnostic_id=diagnostic_id
        )
    )





# =========================================================
# ОҚУШЫ — КІРІС БАҚЫЛАУДЫ БАСТАУ
# =========================================================

@app.route(
    "/diagnostics/<int:diagnostic_id>/start/<int:student_id>",
    methods=["POST"]
)
def start_diagnostic_attempt(
    diagnostic_id,
    student_id
):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    diagnostic = conn.execute(
        """
        SELECT *
        FROM diagnostics
        WHERE id = ?
        """,
        (diagnostic_id,)
    ).fetchone()

    student = conn.execute(
        """
        SELECT *
        FROM students
        WHERE id = ?
        """,
        (student_id,)
    ).fetchone()

    if diagnostic is None or student is None:
        conn.close()
        return "Бақылау немесе оқушы табылмады", 404

    questions = conn.execute(
        """
        SELECT *
        FROM diagnostic_questions
        WHERE diagnostic_id = ?
        ORDER BY question_order, id
        """,
        (diagnostic_id,)
    ).fetchall()

    if not questions:
        conn.close()

        return redirect(
            url_for(
                "diagnostic_detail",
                diagnostic_id=diagnostic_id
            )
        )

    # Аяқталмаған әрекет бар ма?
    attempt = conn.execute(
        """
        SELECT *
        FROM diagnostic_attempts
        WHERE
            diagnostic_id = ?
            AND student_id = ?
            AND status = 'started'
        ORDER BY id DESC
        LIMIT 1
        """,
        (
            diagnostic_id,
            student_id
        )
    ).fetchone()

    if attempt is None:

        max_points = sum(
            int(q["points"] or 0)
            for q in questions
        )

        cursor = conn.execute(
            """
            INSERT INTO diagnostic_attempts
            (
                diagnostic_id,
                student_id,
                status,
                total_questions,
                max_points
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                diagnostic_id,
                student_id,
                "started",
                len(questions),
                max_points
            )
        )

        attempt_id = cursor.lastrowid

        conn.commit()

    else:

        attempt_id = attempt["id"]

    conn.close()

    return redirect(
        url_for(
            "take_diagnostic",
            attempt_id=attempt_id
        )
    )


# =========================================================
# ОҚУШЫ — БАҚЫЛАУДЫ ОРЫНДАУ
# =========================================================

@app.route(
    "/diagnostic-attempts/<int:attempt_id>"
)
def take_diagnostic(attempt_id):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    attempt = conn.execute(
        """
        SELECT
            a.*,
            d.title,
            d.subject,
            s.full_name,
            s.class_name
        FROM diagnostic_attempts a
        JOIN diagnostics d
            ON d.id = a.diagnostic_id
        JOIN students s
            ON s.id = a.student_id
        WHERE a.id = ?
        """,
        (attempt_id,)
    ).fetchone()

    if attempt is None:
        conn.close()
        return "Бақылау әрекеті табылмады", 404

    questions = conn.execute(
        """
        SELECT
            q.*,
            lo.objective_code
        FROM diagnostic_questions q
        JOIN learning_objectives lo
            ON lo.id = q.learning_objective_id
        WHERE q.diagnostic_id = ?
        ORDER BY q.question_order, q.id
        """,
        (attempt["diagnostic_id"],)
    ).fetchall()

    conn.close()

    return render_template(
        "take_diagnostic.html",
        attempt=attempt,
        questions=questions
    )


# =========================================================
# ОҚУШЫ — ЖАУАПТАРДЫ ТЕКСЕРУ
# =========================================================

@app.route(
    "/diagnostic-attempts/<int:attempt_id>/submit",
    methods=["POST"]
)
def submit_diagnostic_attempt(attempt_id):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    attempt = conn.execute(
        """
        SELECT *
        FROM diagnostic_attempts
        WHERE id = ?
        """,
        (attempt_id,)
    ).fetchone()

    if attempt is None:
        conn.close()
        return "Бақылау әрекеті табылмады", 404

    if attempt["status"] == "completed":
        conn.close()

        return redirect(
            url_for(
                "diagnostic_attempt_result",
                attempt_id=attempt_id
            )
        )

    questions = conn.execute(
        """
        SELECT *
        FROM diagnostic_questions
        WHERE diagnostic_id = ?
        ORDER BY question_order, id
        """,
        (attempt["diagnostic_id"],)
    ).fetchall()

    earned_total = 0
    max_total = 0
    answered = 0

    # Ескі жауаптар болса тазалау
    conn.execute(
        """
        DELETE FROM diagnostic_answers
        WHERE attempt_id = ?
        """,
        (attempt_id,)
    )

    for question in questions:

        selected = request.form.get(
            f"question_{question['id']}",
            ""
        ).strip().upper()

        correct = (
            selected != ""
            and selected == question["correct_answer"]
        )

        points = int(
            question["points"] or 0
        )

        earned = (
            points
            if correct
            else 0
        )

        max_total += points

        if selected:
            answered += 1

        earned_total += earned

        conn.execute(
            """
            INSERT INTO diagnostic_answers
            (
                attempt_id,
                question_id,
                selected_answer,
                is_correct,
                earned_points
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                attempt_id,
                question["id"],
                selected,
                1 if correct else 0,
                earned
            )
        )

    percentage = (
        round(
            earned_total / max_total * 100,
            1
        )
        if max_total > 0
        else 0
    )

    conn.execute(
        """
        UPDATE diagnostic_attempts
        SET
            status = 'completed',
            total_questions = ?,
            answered_questions = ?,
            earned_points = ?,
            max_points = ?,
            percentage = ?,
            completed_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """,
        (
            len(questions),
            answered,
            earned_total,
            max_total,
            percentage,
            attempt_id
        )
    )

    # -----------------------------------------------------
    # ОҚУ МАҚСАТЫ БОЙЫНША НӘТИЖЕ
    # -----------------------------------------------------

    conn.execute(
        """
        DELETE FROM objective_results
        WHERE attempt_id = ?
        """,
        (attempt_id,)
    )

    objective_rows = conn.execute(
        """
        SELECT
            q.learning_objective_id,

            COUNT(q.id) AS total_questions,

            SUM(
                CASE
                    WHEN a.is_correct = 1
                    THEN 1
                    ELSE 0
                END
            ) AS correct_questions,

            COALESCE(
                SUM(a.earned_points),
                0
            ) AS earned_points,

            COALESCE(
                SUM(q.points),
                0
            ) AS max_points

        FROM diagnostic_questions q

        LEFT JOIN diagnostic_answers a
            ON a.question_id = q.id
            AND a.attempt_id = ?

        WHERE q.diagnostic_id = ?

        GROUP BY
            q.learning_objective_id
        """,
        (
            attempt_id,
            attempt["diagnostic_id"]
        )
    ).fetchall()

    for row in objective_rows:

        maximum = int(
            row["max_points"] or 0
        )

        earned = int(
            row["earned_points"] or 0
        )

        objective_percentage = (
            round(
                earned / maximum * 100,
                1
            )
            if maximum > 0
            else 0
        )

        # ---------------------------------------------
        # МЕҢГЕРУ СТАТУСЫ
        # ---------------------------------------------

        if objective_percentage >= 80:

            mastery_status = "Меңгерілген"

        elif objective_percentage >= 50:

            mastery_status = (
                "Бекітуді қажет етеді"
            )

        else:

            mastery_status = "Олқылық бар"

        conn.execute(
            """
            INSERT INTO objective_results
            (
                attempt_id,
                learning_objective_id,
                total_questions,
                correct_questions,
                earned_points,
                max_points,
                percentage,
                mastery_status
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                attempt_id,
                row["learning_objective_id"],
                row["total_questions"],
                row["correct_questions"],
                earned,
                maximum,
                objective_percentage,
                mastery_status
            )
        )

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "diagnostic_attempt_result",
            attempt_id=attempt_id
        )
    )


# =========================================================
# НӘТИЖЕ БЕТІ
# =========================================================

@app.route(
    "/diagnostic-attempts/<int:attempt_id>/result"
)
def diagnostic_attempt_result(attempt_id):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    attempt = conn.execute(
        """
        SELECT
            a.*,
            d.title,
            d.subject,
            s.full_name,
            s.class_name
        FROM diagnostic_attempts a
        JOIN diagnostics d
            ON d.id = a.diagnostic_id
        JOIN students s
            ON s.id = a.student_id
        WHERE a.id = ?
        """,
        (attempt_id,)
    ).fetchone()

    if attempt is None:
        conn.close()
        return "Нәтиже табылмады", 404

    objective_results = conn.execute(
        """
        SELECT
            r.*,
            lo.objective_code,
            lo.objective_text
        FROM objective_results r
        JOIN learning_objectives lo
            ON lo.id = r.learning_objective_id
        WHERE r.attempt_id = ?
        ORDER BY lo.objective_code
        """,
        (attempt_id,)
    ).fetchall()

    conn.close()

    return render_template(
        "diagnostic_result.html",
        attempt=attempt,
        objective_results=objective_results
    )





# =========================================================
# ОҚУ ОЛҚЫЛЫҚТАРЫ
# =========================================================

@app.route("/learning-gaps")
def learning_gaps():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()


    # -----------------------------------------------------
    # ОҚУ ОЛҚЫЛЫҚТАРЫ + ЖЕКЕ ТАПСЫРМАЛАР ПРОГРЕСІ
    # -----------------------------------------------------

    rows = conn.execute(
        """
        SELECT
            r.id,
            r.attempt_id,
            r.learning_objective_id,
            r.total_questions,
            r.correct_questions,
            r.earned_points,
            r.max_points,
            r.percentage,
            r.mastery_status,

            lo.objective_code,
            lo.objective_text,
            lo.lesson_topic,

            s.id AS student_id,
            s.full_name,
            s.class_name,

            d.id AS diagnostic_id,
            d.title AS diagnostic_title,
            d.subject,

            a.completed_at,

            COUNT(
                DISTINCT pt.id
            ) AS task_count,

            COUNT(
                DISTINCT CASE
                    WHEN pt.status = 'completed'
                    THEN pt.id
                END
            ) AS completed_task_count,

            COUNT(
                DISTINCT CASE
                    WHEN
                        pt.status = 'completed'
                        AND COALESCE(
                            pt.is_correct,
                            0
                        ) = 1
                    THEN pt.id
                END
            ) AS correct_task_count

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN students s
            ON s.id = a.student_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        JOIN learning_objectives lo
            ON lo.id =
               r.learning_objective_id

        LEFT JOIN personal_tasks pt
            ON pt.student_id = s.id
            AND pt.learning_objective_id =
                r.learning_objective_id
            AND pt.source_attempt_id =
                r.attempt_id

        WHERE
            a.status = 'completed'

            AND r.mastery_status IN (
                'Олқылық бар',
                'Бекітуді қажет етеді'
            )

        GROUP BY
            r.id,
            r.attempt_id,
            r.learning_objective_id,
            r.total_questions,
            r.correct_questions,
            r.earned_points,
            r.max_points,
            r.percentage,
            r.mastery_status,

            lo.objective_code,
            lo.objective_text,
            lo.lesson_topic,

            s.id,
            s.full_name,
            s.class_name,

            d.id,
            d.title,
            d.subject,

            a.completed_at

        ORDER BY
            r.percentage ASC,
            s.full_name ASC,
            lo.objective_code ASC
        """
    ).fetchall()


    # -----------------------------------------------------
    # TEMPLATE ҮШІН ТОЛЫҚ GAP OBJECT
    # -----------------------------------------------------

    gaps = []


    for row in rows:

        percentage = float(
            row["percentage"]
            or 0
        )


        if percentage < 50:

            gap_level = "Олқылық бар"

            recommendation = (
                "Оқу мақсатын негізгі деңгейден "
                "қайта түсіндіру және Қолдау "
                "деңгейіндегі жеке тапсырмалар беру."
            )

        else:

            gap_level = (
                "Бекітуді қажет етеді"
            )

            recommendation = (
                "Оқу мақсатын бекіту үшін "
                "қосымша жаттығулар мен "
                "қайталау тапсырмаларын орындау."
            )


        gaps.append({

            "id":
                row["id"],

            "attempt_id":
                row["attempt_id"],

            "learning_objective_id":
                row["learning_objective_id"],

            "total_questions":
                row["total_questions"],

            "correct_questions":
                row["correct_questions"],

            "earned_points":
                row["earned_points"],

            "max_points":
                row["max_points"],

            "percentage":
                percentage,

            "mastery_status":
                row["mastery_status"],

            "gap_level":
                gap_level,

            "objective_code":
                row["objective_code"],

            "objective_text":
                row["objective_text"],

            "lesson_topic":
                row["lesson_topic"]
                or "",

            "student_id":
                row["student_id"],

            "full_name":
                row["full_name"],

            "class_name":
                row["class_name"],

            "diagnostic_id":
                row["diagnostic_id"],

            "diagnostic_title":
                row["diagnostic_title"],

            "subject":
                row["subject"],

            "completed_at":
                row["completed_at"],

            "recommendation":
                recommendation,

            "task_count":
                int(
                    row["task_count"]
                    or 0
                ),

            "completed_task_count":
                int(
                    row[
                        "completed_task_count"
                    ]
                    or 0
                ),

            "correct_task_count":
                int(
                    row[
                        "correct_task_count"
                    ]
                    or 0
                )
        })


    # -----------------------------------------------------
    # ЖАЛПЫ СТАТИСТИКА
    # -----------------------------------------------------

    total_students = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM students
        WHERE class_name = ?
        """,
        ('4 «Д»',)
    ).fetchone()["total"]


    completed_attempts = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM diagnostic_attempts
        WHERE status = 'completed'
        """
    ).fetchone()["total"]


    gap_count = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM objective_results
        WHERE mastery_status =
              'Олқылық бар'
        """
    ).fetchone()["total"]


    reinforcement_count = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM objective_results
        WHERE mastery_status =
              'Бекітуді қажет етеді'
        """
    ).fetchone()["total"]


    # -----------------------------------------------------
    # ЕҢ ЖИІ КЕЗДЕСЕТІН ОЛҚЫЛЫҚТАР
    # -----------------------------------------------------

    common_gaps = conn.execute(
        """
        SELECT
            lo.objective_code,
            lo.objective_text,
            d.subject,

            COUNT(*) AS student_count,

            ROUND(
                AVG(r.percentage),
                1
            ) AS average_percentage

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        JOIN learning_objectives lo
            ON lo.id =
               r.learning_objective_id

        WHERE
            a.status = 'completed'

            AND r.mastery_status =
                'Олқылық бар'

        GROUP BY
            lo.objective_code,
            lo.objective_text,
            d.subject

        ORDER BY
            student_count DESC,
            average_percentage ASC

        LIMIT 10
        """
    ).fetchall()


    conn.close()


    return render_template(
        "learning_gaps.html",

        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),

        gaps=gaps,

        common_gaps=common_gaps,

        total_students=total_students,

        completed_attempts=
            completed_attempts,

        gap_count=gap_count,

        reinforcement_count=
            reinforcement_count
    )











# =========================================================
# ЖЕКЕ БІЛІМ БАҒЫТЫ
# =========================================================

@app.route("/learning-paths")
def learning_paths():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    # -----------------------------------------------------
    # ОҚУШЫЛАР
    # -----------------------------------------------------

    students = conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name
        FROM students
        WHERE class_name = ?
        ORDER BY full_name
        """,
        ('4 «Д»',)
    ).fetchall()


    # -----------------------------------------------------
    # СОҢҒЫ АЯҚТАЛҒАН НӘТИЖЕЛЕР
    # -----------------------------------------------------

    rows = conn.execute(
        """
        SELECT
            s.id AS student_id,
            s.full_name,
            s.class_name,

            d.subject,
            d.title AS diagnostic_title,

            a.id AS attempt_id,
            a.percentage AS diagnostic_percentage,
            a.completed_at,

            r.learning_objective_id,
            r.percentage AS objective_percentage,
            r.mastery_status,

            lo.objective_code,
            lo.objective_text

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN students s
            ON s.id = a.student_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        JOIN learning_objectives lo
            ON lo.id = r.learning_objective_id

        WHERE
            a.status = 'completed'

        ORDER BY
            s.full_name,
            d.subject,
            r.percentage ASC,
            lo.objective_code
        """
    ).fetchall()


    # -----------------------------------------------------
    # ЖЕКЕ БАҒЫТТЫ ЕСЕПТЕУ
    # -----------------------------------------------------

    paths = []

    support_count = 0
    reinforce_count = 0
    develop_count = 0

    for row in rows:

        percentage = float(
            row["objective_percentage"] or 0
        )

        if percentage < 50:

            path_level = "Қолдау"
            path_icon = "🟢"

            recommendation = (
                "Оқу мақсатын негізгі деңгейден "
                "қайта түсіндіру және қарапайым "
                "жаттығулар орындау."
            )

            support_count += 1

        elif percentage < 80:

            path_level = "Бекіту"
            path_icon = "🔵"

            recommendation = (
                "Дағдыны бекітуге арналған "
                "қосымша жаттығулар мен "
                "қайталау тапсырмаларын орындау."
            )

            reinforce_count += 1

        else:

            path_level = "Дамыту"
            path_icon = "🟣"

            recommendation = (
                "Оқу мақсатын күрделірек және "
                "жаңа жағдайда қолданатын "
                "дамыту тапсырмаларын орындау."
            )

            develop_count += 1


        paths.append({

            "student_id":
                row["student_id"],

            "full_name":
                row["full_name"],

            "class_name":
                row["class_name"],

            "subject":
                row["subject"],

            "diagnostic_title":
                row["diagnostic_title"],

            "attempt_id":
                row["attempt_id"],

            "completed_at":
                row["completed_at"],

            "learning_objective_id":
                row["learning_objective_id"],

            "objective_code":
                row["objective_code"],

            "objective_text":
                row["objective_text"],

            "percentage":
                percentage,

            "mastery_status":
                row["mastery_status"],

            "path_level":
                path_level,

            "path_icon":
                path_icon,

            "recommendation":
                recommendation
        })


    # -----------------------------------------------------
    # ЖЕКЕ ТАПСЫРМА ҚҰРУҒА ДАЙЫН БАҚЫЛАУЛАР
    # -----------------------------------------------------

    completed_attempts_for_paths = conn.execute(
        """
        SELECT
            a.id AS attempt_id,
            a.student_id,
            a.percentage,
            a.completed_at,

            s.full_name,

            d.title AS diagnostic_title,
            d.subject,

            COUNT(DISTINCT r.learning_objective_id)
                AS objective_count,

            (
                SELECT COUNT(*)
                FROM personal_tasks pt
                WHERE pt.source_attempt_id = a.id
            ) AS personal_task_count

        FROM diagnostic_attempts a

        JOIN students s
            ON s.id = a.student_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        LEFT JOIN objective_results r
            ON r.attempt_id = a.id

        WHERE a.status = 'completed'

        GROUP BY
            a.id,
            a.student_id,
            a.percentage,
            a.completed_at,
            s.full_name,
            d.title,
            d.subject

        ORDER BY
            a.completed_at DESC,
            a.id DESC
        """
    ).fetchall()


    # -----------------------------------------------------
    # ҚҰРЫЛҒАН ЖЕКЕ ТАПСЫРМАЛАР
    # -----------------------------------------------------

    personal_tasks = conn.execute(
        """
        SELECT
            pt.*,

            s.full_name,
            s.class_name,

            lo.objective_code,
            lo.objective_text

        FROM personal_tasks pt

        JOIN students s
            ON s.id = pt.student_id

        JOIN learning_objectives lo
            ON lo.id = pt.learning_objective_id

        ORDER BY
            pt.assigned_at DESC,
            pt.id DESC
        """
    ).fetchall()


    conn.close()


    return render_template(
        "learning_paths.html",

        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),

        students=students,
        paths=paths,

        personal_tasks=personal_tasks,
        completed_attempts=completed_attempts_for_paths,

        support_count=support_count,
        reinforce_count=reinforce_count,
        develop_count=develop_count
    )





# =========================================================
# API-СЫЗ ЖЕКЕ ТАПСЫРМА ГЕНЕРАТОРЫ
# =========================================================

def build_personal_task(
    subject,
    objective_code,
    objective_text,
    path_level
):

    subject = (subject or "").strip()
    objective_code = (objective_code or "").strip()
    objective_text = (objective_text or "").strip()
    path_level = (path_level or "Бекіту").strip()


    # =====================================================
    # МАТЕМАТИКА
    # =====================================================

    if subject == "Математика":

        # -------------------------------------------------
        # 4.1.1.1 — көптаңбалы сандар
        # -------------------------------------------------

        if objective_code == "4.1.1.1":

            if path_level == "Қолдау":

                return {
                    "title":
                        "Сандар қатарын жалғастыр",

                    "task_text":
                        "999 999 санынан кейін "
                        "қандай сан келеді?",

                    "correct_answer":
                        "1 000 000",

                    "explanation":
                        "999 999 санына 1 қосқанда "
                        "1 000 000 саны шығады.",

                    "max_points": 1
                }


            if path_level == "Бекіту":

                return {
                    "title":
                        "Санның орнын анықта",

                    "task_text":
                        "499 999 және 500 001 "
                        "сандарының арасында қандай "
                        "натурал сан орналасқан?",

                    "correct_answer":
                        "500 000",

                    "explanation":
                        "499 999 санынан кейін "
                        "500 000, одан кейін "
                        "500 001 орналасады.",

                    "max_points": 1
                }


            return {
                "title":
                    "Көптаңбалы санды анықта",

                "task_text":
                    "1 000 000 санынан 1-ге кем "
                    "санды және 1-ге артық санды жаз.",

                "correct_answer":
                    "999 999 және 1 000 001",

                "explanation":
                    "1 000 000 - 1 = 999 999, "
                    "1 000 000 + 1 = 1 000 001.",

                "max_points": 2
            }


        # -------------------------------------------------
        # 4.1.1.2 — оқу, жазу, салыстыру, дөңгелектеу
        # -------------------------------------------------

        if objective_code == "4.1.1.2":

            if path_level == "Қолдау":

                return {
                    "title":
                        "Сандарды салыстыр",

                    "task_text":
                        "45 670 және 45 607 "
                        "сандарын салыстыр. "
                        "Қайсысы үлкен?",

                    "correct_answer":
                        "45 670",

                    "explanation":
                        "Он мыңдық пен мыңдықтары "
                        "бірдей. Жүздіктерді "
                        "салыстырғанда 6 > 6 емес, "
                        "келесі разряд бойынша "
                        "45 670 саны үлкен.",

                    "max_points": 1
                }


            if path_level == "Бекіту":

                return {
                    "title":
                        "Санды дөңгелекте",

                    "task_text":
                        "48 764 санын жүздікке "
                        "дейін дөңгелекте.",

                    "correct_answer":
                        "48 800",

                    "explanation":
                        "Ондық разрядтағы 6 саны "
                        "5-тен үлкен болғандықтан "
                        "жүздік жоғарылайды.",

                    "max_points": 1
                }


            return {
                "title":
                    "Салыстыру және дөңгелектеу",

                "task_text":
                    "125 649 санын мыңдыққа дейін "
                    "дөңгелекте және нәтижені "
                    "125 000 санымен салыстыр.",

                "correct_answer":
                    "126 000 > 125 000",

                "explanation":
                    "Жүздік разрядтағы 6 саны "
                    "5-тен үлкен, сондықтан "
                    "125 649 ≈ 126 000.",

                "max_points": 2
            }


        # -------------------------------------------------
        # 4.1.2.8 — жазбаша қосу және азайту
        # -------------------------------------------------

        if objective_code == "4.1.2.8":

            if path_level == "Қолдау":

                return {
                    "title":
                        "Көптаңбалы сандарды қос",

                    "task_text":
                        "23 456 + 12 321 мәнін тап.",

                    "correct_answer":
                        "35 777",

                    "explanation":
                        "Сандарды разрядтары бойынша "
                        "жазып, оңнан солға қарай "
                        "қосамыз.",

                    "max_points": 1
                }


            if path_level == "Бекіту":

                return {
                    "title":
                        "Көптаңбалы сандарды азайт",

                    "task_text":
                        "54 320 - 18 745 мәнін тап.",

                    "correct_answer":
                        "35 575",

                    "explanation":
                        "Сандарды разрядтары бойынша "
                        "жазып, жазбаша азайтамыз.",

                    "max_points": 1
                }


            return {
                "title":
                    "Екі амалдан тұратын есеп",

                "task_text":
                    "Кітапханада 125 600 кітап болды. "
                    "18 750 кітап берілді, кейін "
                    "9 850 жаңа кітап келді. "
                    "Қазір неше кітап бар?",

                "correct_answer":
                    "116 700",

                "explanation":
                    "125 600 - 18 750 = 106 850. "
                    "106 850 + 9 850 = 116 700.",

                "max_points": 2
            }


    # =====================================================
    # ҚАЗАҚ ТІЛІ
    # =====================================================

    if subject == "Қазақ тілі":

        # -------------------------------------------------
        # 4.1.1.1
        # -------------------------------------------------

        if objective_code == "4.1.1.1":

            if path_level == "Қолдау":

                return {
                    "title":
                        "Негізгі ойды анықта",

                    "task_text":
                        "Мәтінді оқы: "
                        "«Асан таңертең ерте тұрып, "
                        "мектеп ауласындағы жас ағаштарға "
                        "су құйды.» "
                        "Мәтіндегі негізгі әрекетті ата.",

                    "correct_answer":
                        "Асан жас ағаштарға су құйды.",

                    "explanation":
                        "Негізгі әрекет — кейіпкердің "
                        "мәтінде атқарған басты ісі.",

                    "max_points": 1
                }


            if path_level == "Бекіту":

                return {
                    "title":
                        "Тірек сөздерді белгіле",

                    "task_text":
                        "«Балалар кітапханаға барып, "
                        "табиғат туралы кітаптарды оқыды.» "
                        "Сөйлемнен негізгі тірек сөздерді жаз.",

                    "correct_answer":
                        "балалар, кітапхана, табиғат, кітап",

                    "explanation":
                        "Тірек сөздер мәтіннің негізгі "
                        "мазмұнын білдіреді.",

                    "max_points": 1
                }


            return {
                "title":
                    "Мазмұнын қысқаша жеткіз",

                "task_text":
                    "«Оқушылар мектеп ауласына көшет "
                    "отырғызды. Олар көшеттерді суарып, "
                    "күтіп-баптауға келісті.» "
                    "Мәтіннің мазмұнын бір сөйлеммен жаз.",

                "correct_answer":
                    "Оқушылар көшет отырғызып, "
                    "оны күтіп-баптауға келісті.",

                "explanation":
                    "Қысқаша мазмұнда негізгі оқиға "
                    "мен әрекет сақталуы керек.",

                "max_points": 2
            }


        # -------------------------------------------------
        # 4.1.7.1 — үндестік
        # -------------------------------------------------

        if objective_code == "4.1.7.1":

            if path_level == "Қолдау":

                return {
                    "title":
                        "Дұрыс айтылымды анықта",

                    "task_text":
                        "«Балалар мектепке келді» "
                        "сөйлемін буын үндестігін "
                        "сақтап дауыстап оқы.",

                    "correct_answer":
                        "Буын және дыбыс үндестігі сақталады.",

                    "explanation":
                        "Сөздердегі дыбыстарды дұрыс "
                        "үндестіріп айту қажет.",

                    "max_points": 1
                }


            if path_level == "Бекіту":

                return {
                    "title":
                        "Үндестік заңын қолдан",

                    "task_text":
                        "«кітап» сөзіне көптік "
                        "жалғауын дұрыс жалға.",

                    "correct_answer":
                        "кітаптар",

                    "explanation":
                        "«Кітап» сөзінің соңғы буыны "
                        "жуан болғандықтан -тар "
                        "жалғауы жалғанады.",

                    "max_points": 1
                }


            return {
                "title":
                    "Сөздерді түрлендір",

                "task_text":
                    "«мектеп» және «бала» сөздеріне "
                    "көптік жалғауын жалғап, "
                    "үндестік заңын түсіндір.",

                "correct_answer":
                    "мектептер, балалар",

                "explanation":
                    "Жіңішке сөзге жіңішке, "
                    "жуан сөзге жуан қосымша жалғанады.",

                "max_points": 2
            }


    # =====================================================
    # ЖАЛПЫ FALLBACK
    # =====================================================

    if path_level == "Қолдау":

        instruction = (
            "Оқу мақсатын түсінуге арналған "
            "қарапайым тапсырманы орында."
        )

    elif path_level == "Бекіту":

        instruction = (
            "Оқу мақсатын бекітуге арналған "
            "жаттығуды орында."
        )

    else:

        instruction = (
            "Оқу мақсатын жаңа жағдайда "
            "қолданатын тапсырманы орында."
        )


    return {
        "title":
            f"{objective_code} — {path_level}",

        "task_text":
            f"{instruction} "
            f"Оқу мақсаты: {objective_text}",

        "correct_answer":
            "",

        "explanation":
            (
                "Бұл оқу мақсатына арналған "
                "нақты шаблон кейін толықтырылады."
            ),

        "max_points":
            1 if path_level != "Дамыту" else 2
    }





# =========================================================
# ЖЕКЕ ТАПСЫРМАЛАРДЫ АВТОМАТТЫ ҚҰРУ
# =========================================================

@app.route(
    "/learning-paths/generate/<int:attempt_id>",
    methods=["POST"]
)
def generate_personal_tasks(attempt_id):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    # -----------------------------------------------------
    # АЯҚТАЛҒАН ДИАГНОСТИКА
    # -----------------------------------------------------

    attempt = conn.execute(
        """
        SELECT
            a.*,
            d.subject,
            d.title AS diagnostic_title,
            s.full_name
        FROM diagnostic_attempts a

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        JOIN students s
            ON s.id = a.student_id

        WHERE
            a.id = ?
            AND a.status = 'completed'
        """,
        (attempt_id,)
    ).fetchone()

    if attempt is None:

        conn.close()

        return redirect(
            url_for("learning_paths")
        )


    # -----------------------------------------------------
    # ОҚУ МАҚСАТТАРЫ БОЙЫНША НӘТИЖЕ
    # -----------------------------------------------------

    results = conn.execute(
        """
        SELECT
            r.*,

            lo.objective_code,
            lo.objective_text

        FROM objective_results r

        JOIN learning_objectives lo
            ON lo.id = r.learning_objective_id

        WHERE r.attempt_id = ?

        ORDER BY
            r.percentage ASC,
            lo.objective_code
        """,
        (attempt_id,)
    ).fetchall()


    # -----------------------------------------------------
    # ӘР МАҚСАТҚА ЖЕКЕ ТАПСЫРМА
    # -----------------------------------------------------

    for result in results:

        percentage = float(
            result["percentage"] or 0
        )

        # -------------------------------------------------
        # ЖЕКЕ ТҮЗЕТУ БАҒЫТЫ
        #
        # 0–49%  -> Қолдау
        # 50–79% -> Бекіту
        # 80–100% меңгерілген мақсатқа
        # жеке түзету тапсырмасы қажет емес.
        # -------------------------------------------------

        if percentage < 50:

            path_level = "Қолдау"

        elif percentage < 80:

            path_level = "Бекіту"

        else:

            # Бұл оқу мақсаты меңгерілген.
            # Оқу олқылығын жоюға арналған
            # personal_tasks қатарына қоспаймыз.
            continue




        # -------------------------------------------------
        # ҚАЙТАЛАНАТЫН ТАПСЫРМАНЫ БОЛДЫРМАУ
        # -------------------------------------------------

        existing = conn.execute(
            """
            SELECT id
            FROM personal_tasks

            WHERE
                student_id = ?
                AND learning_objective_id = ?
                AND source_attempt_id = ?
                AND path_level = ?

            LIMIT 1
            """,
            (
                attempt["student_id"],
                result["learning_objective_id"],
                attempt_id,
                path_level
            )
        ).fetchone()

        if existing:
            continue


        # -------------------------------------------------
        # НАҚТЫ API-СЫЗ ТАПСЫРМА
        # -------------------------------------------------

        # ---------------------------------------------
        # КТЖ КОНТЕКСТІ
        # ---------------------------------------------

        ktj_context = get_ktj_context(
            conn,
            attempt["subject"],
            result["learning_objective_id"]
        )

        lesson_topics = (
            ktj_context.get(
                "lesson_topics",
                []
            )
        )

        # Бірінші тақырыпты емес,
        # оқу мақсатына ең сәйкес КТЖ тақырыбын таңдаймыз.
        current_topic = select_best_ktj_topic(
            result["objective_text"],
            lesson_topics
        )

        # ---------------------------------------------
        # ТОЛЫҚ КТЖ КОНТЕКСТІ
        # ---------------------------------------------

        secondary_topics = [
            topic
            for topic in lesson_topics
            if topic != current_topic
        ]

        ktj_context_text = (
            "Негізгі сабақ: "
            + (current_topic or "—")
        )

        if secondary_topics:

            ktj_context_text += (
                "\nҚосымша сабақтар: "
                + "; ".join(
                    secondary_topics
                )
            )

        sections = ktj_context.get(
            "sections",
            []
        )

        if sections:

            ktj_context_text += (
                "\nБөлімдер: "
                + "; ".join(
                    sections
                )
            )

        # ---------------------------------------------
        # ТОЛЫҚ КТЖ КОНТЕКСТІ
        # ---------------------------------------------

        secondary_topics = [
            topic
            for topic in lesson_topics
            if topic != current_topic
        ]

        ktj_context_text = (
            "Негізгі сабақ: "
            + (current_topic or "—")
        )

        if secondary_topics:

            ktj_context_text += (
                "\nҚосымша сабақтар: "
                + "; ".join(
                    secondary_topics
                )
            )

        sections = ktj_context.get(
            "sections",
            []
        )

        if sections:

            ktj_context_text += (
                "\nБөлімдер: "
                + "; ".join(
                    sections
                )
            )


        # ---------------------------------------------
        # НАҚТЫ ТАПСЫРМА
        # ---------------------------------------------

        # ---------------------------------------------
        # GENERATOR ROUTER
        # ---------------------------------------------

        if has_specific_task_template(
            attempt["subject"],
            result["objective_code"]
        ):

            # Қолмен дайындалған нақты шаблон
            generated = build_personal_task(
                attempt["subject"],
                result["objective_code"],
                result["objective_text"],
                path_level
            )

            generated["generator_type"] = (
                "specific_template"
            )

        else:

            # КТЖ контекстіне негізделген
            # әмбебап генератор
            generated = build_ktj_fallback_task(
                attempt["subject"],
                result["objective_code"],
                result["objective_text"],
                path_level,
                current_topic,
                lesson_topics,
                ktj_context.get(
                    "sections",
                    []
                )
            )

            generated["generator_type"] = (
                "ktj_fallback"
            )


        # ---------------------------------------------
        # КТЖ КОНТЕКСТІН ҚОСУ
        # ---------------------------------------------

        if current_topic:

            generated["ktj_topic"] = (
                current_topic
            )

        else:

            generated["ktj_topic"] = ""


        # -------------------------------------------------
        # DATABASE
        # -------------------------------------------------

        conn.execute(
            """
            INSERT INTO personal_tasks
            (
                student_id,
                learning_objective_id,
                source_attempt_id,

                subject,
                path_level,
                task_type,

                title,
                task_text,
                correct_answer,
                explanation,

                ktj_topic,
                ktj_context,

                status,
                earned_points,
                max_points
            )

            VALUES (
                ?, ?, ?,
                ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?,
                ?, ?, ?
            )
            """,
            (
                attempt["student_id"],
                result["learning_objective_id"],
                attempt_id,

                attempt["subject"],
                path_level,
                "practice",

                generated["title"],
                generated["task_text"],
                generated["correct_answer"],
                generated["explanation"],

                generated.get(
                    "ktj_topic",
                    current_topic
                ),

                ktj_context_text,

                "assigned",
                0,
                generated["max_points"]
            )
        )


    conn.commit()
    conn.close()

    return redirect(
        url_for("learning_paths")
    )





# =========================================================
# КТЖ КОНТЕКСТІ
# =========================================================

def get_ktj_context(
    conn,
    subject,
    learning_objective_id
):

    objective = conn.execute(
        """
        SELECT
            id,
            subject,
            objective_code,
            objective_text,
            lesson_topic,
            section
        FROM learning_objectives
        WHERE id = ?
        """,
        (learning_objective_id,)
    ).fetchone()

    if objective is None:

        return {
            "objective_code": "",
            "objective_text": "",
            "lesson_topics": [],
            "sections": []
        }


    # -----------------------------------------------------
    # ДӘЛ ОСЫ КОДТЫҢ КТЖ-ДАҒЫ БАРЛЫҚ САБАҚ ТАҚЫРЫПТАРЫ
    # -----------------------------------------------------

    rows = conn.execute(
        """
        SELECT
            lesson_topic,
            section
        FROM learning_objectives

        WHERE
            subject = ?
            AND objective_code = ?

        ORDER BY id
        """,
        (
            subject,
            objective["objective_code"]
        )
    ).fetchall()


    lesson_topics = []

    sections = []


    for row in rows:

        topic = (
            row["lesson_topic"]
            or ""
        ).strip()

        section = (
            row["section"]
            or ""
        ).strip()


        if (
            topic
            and topic not in lesson_topics
        ):
            lesson_topics.append(
                topic
            )


        if (
            section
            and section not in sections
        ):
            sections.append(
                section
            )


    return {

        "objective_code":
            objective["objective_code"]
            or "",

        "objective_text":
            objective["objective_text"]
            or "",

        "current_lesson_topic":
            objective["lesson_topic"]
            or "",

        "current_section":
            objective["section"]
            or "",

        "lesson_topics":
            lesson_topics,

        "sections":
            sections
    }





# =========================================================
# КТЖ — ЕҢ СӘЙКЕС САБАҚ ТАҚЫРЫБЫН ТАҢДАУ
# =========================================================

def select_best_ktj_topic(
    objective_text,
    lesson_topics
):

    import re
    from difflib import SequenceMatcher

    if not lesson_topics:
        return ""

    objective_text = (
        objective_text
        or ""
    ).strip()


    # -----------------------------------------------------
    # NORMALIZE
    # -----------------------------------------------------

    def normalize(text):

        text = (
            text
            or ""
        ).lower()

        replacements = {
            "қай": "қай",
            "азай": "азай",
            "–": " ",
            "—": " ",
            "-": " ",
            "/": " ",
        }

        for old, new in replacements.items():
            text = text.replace(
                old,
                new
            )

        text = re.sub(
            r'[^а-яәғқңөұүһі0-9\s]',
            ' ',
            text
        )

        text = re.sub(
            r'\s+',
            ' ',
            text
        )

        return text.strip()


    # -----------------------------------------------------
    # ЖАЛПЫ СӨЗДЕР
    # -----------------------------------------------------

    stop_words = {
        "және",
        "мен",
        "пен",
        "бен",
        "үшін",
        "бойынша",
        "арқылы",
        "олардың",
        "оларды",
        "қолдану",
        "анықтау",
        "табу",
        "орындау",
        "өрнектеу",
        "есептеу",
        "алгоритмдерін",
        "алгоритмін",
    }


    def get_words(text):

        result = []

        for word in normalize(
            text
        ).split():

            if len(word) < 3:
                continue

            if word in stop_words:
                continue

            result.append(
                word
            )

        return result


    # -----------------------------------------------------
    # ҚАРАПАЙЫМ ТҮБІРЛІК СӘЙКЕСТІК
    # -----------------------------------------------------

    def similar_word(a, b):

        if a == b:
            return True

        # көптаңбалы / көптаңбалы
        # дөңгелектеу / дөңгелектеу
        if (
            len(a) >= 5
            and len(b) >= 5
            and a[:5] == b[:5]
        ):
            return True

        if (
            len(a) >= 6
            and len(b) >= 6
        ):

            ratio = SequenceMatcher(
                None,
                a,
                b
            ).ratio()

            if ratio >= 0.82:
                return True

        return False


    objective_words = get_words(
        objective_text
    )

    objective_norm = normalize(
        objective_text
    )


    # -----------------------------------------------------
    # ЖАЛПЫ ТАҚЫРЫПТАРҒА PENALTY
    # -----------------------------------------------------

    generic_phrases = [
        "өткенді қайталау",
        "өткенді қорытындылау",
        "қорытындылау",
        "өзіңді тексер",
        "бақылау жұмысы",
        "қайталау",
    ]


    scored = []


    for position, topic in enumerate(
        lesson_topics
    ):

        topic_norm = normalize(
            topic
        )

        topic_words = get_words(
            topic
        )

        if not topic_words:
            continue


        # -------------------------------------------------
        # WORD OVERLAP
        # -------------------------------------------------

        matched = 0

        for topic_word in topic_words:

            if any(
                similar_word(
                    topic_word,
                    objective_word
                )
                for objective_word
                in objective_words
            ):

                matched += 1


        overlap = (
            matched
            /
            len(topic_words)
        )


        # -------------------------------------------------
        # ЕКІ ЖАҚТЫ OVERLAP
        # -------------------------------------------------

        objective_matched = 0

        for objective_word in objective_words:

            if any(
                similar_word(
                    objective_word,
                    topic_word
                )
                for topic_word
                in topic_words
            ):

                objective_matched += 1


        reverse_overlap = (
            objective_matched
            /
            len(objective_words)
            if objective_words
            else 0
        )


        # -------------------------------------------------
        # STRING SIMILARITY
        # -------------------------------------------------

        sequence = SequenceMatcher(
            None,
            topic_norm,
            objective_norm
        ).ratio()


        # -------------------------------------------------
        # GENERIC PENALTY
        # -------------------------------------------------

        penalty = 0

        for phrase in generic_phrases:

            if phrase in topic_norm:
                penalty += 35


        # -------------------------------------------------
        # FINAL SCORE
        # -------------------------------------------------

        score = (
            overlap * 65
            +
            reverse_overlap * 20
            +
            sequence * 15
            -
            penalty
        )


        # Нақты тақырыпта бірнеше мағыналы сөз болса,
        # шағын bonus
        if (
            matched >= 2
            and penalty == 0
        ):
            score += 10


        scored.append(
            {
                "topic": topic,
                "score": round(
                    score,
                    2
                ),
                "position": position
            }
        )


    if not scored:
        return lesson_topics[0]


    scored.sort(
        key=lambda item: (
            item["score"],
            -item["position"]
        ),
        reverse=True
    )


    return scored[0]["topic"]





# =========================================================
# КТЖ НЕГІЗІНДЕГІ ӘМБЕБАП ТАПСЫРМА
# =========================================================

def build_ktj_fallback_task(
    subject,
    objective_code,
    objective_text,
    path_level,
    main_topic="",
    lesson_topics=None,
    sections=None
):

    subject = (subject or "").strip()
    objective_code = (objective_code or "").strip()
    objective_text = (objective_text or "").strip()
    path_level = (path_level or "Бекіту").strip()
    main_topic = (main_topic or "").strip()

    lesson_topics = lesson_topics or []
    sections = sections or []

    topic = (
        main_topic
        or (
            lesson_topics[0]
            if lesson_topics
            else objective_code
        )
    )

    objective_lower = objective_text.lower()
    topic_lower = topic.lower()


    # =====================================================
    # МАТЕМАТИКА — ӨЛШЕМ БІРЛІКТЕРІ
    # =====================================================

    if (
        subject == "Математика"
        and (
            "өлшем" in objective_lower
            or "шама" in objective_lower
            or "гектар" in topic_lower
            or "уақыт" in topic_lower
        )
    ):

        # =================================================
        # КТЖ ТАҚЫРЫБЫНА БАСЫМДЫҚ
        # =================================================

        topic_is_area = any(
            key in topic_lower
            for key in [
                "гектар",
                "аудан",
                "ар "
            ]
        )

        topic_is_time = any(
            key in topic_lower
            for key in [
                "уақыт",
                "сағат",
                "минут"
            ]
        )

        topic_is_length = any(
            key in topic_lower
            for key in [
                "ұзындық",
                "метр",
                "километр"
            ]
        )

        topic_is_mass = any(
            key in topic_lower
            for key in [
                "масса",
                "килограмм",
                "грамм",
                "тонна"
            ]
        )

        topic_is_volume = any(
            key in topic_lower
            for key in [
                "көлем",
                "сыйымдылық",
                "литр"
            ]
        )


        # -------------------------------------------------
        # 1. УАҚЫТ
        # -------------------------------------------------

        if (
            topic_is_time
            or (
                not any([
                    topic_is_area,
                    topic_is_length,
                    topic_is_mass,
                    topic_is_volume
                ])
                and "уақыт" in objective_lower
            )
        ):

            if path_level == "Қолдау":

                return {
                    "title":
                        f"{topic} — Қолдау",

                    "task_text":
                        "1 сағатта неше минут бар?",

                    "correct_answer":
                        "60 минут",

                    "explanation":
                        "1 сағат = 60 минут.",

                    "max_points": 1
                }

            elif path_level == "Бекіту":

                return {
                    "title":
                        f"{topic} — Бекіту",

                    "task_text":
                        "2 сағат 30 минутты "
                        "минутпен өрнекте.",

                    "correct_answer":
                        "150 минут",

                    "explanation":
                        "2 сағат = 120 минут. "
                        "120 + 30 = 150 минут.",

                    "max_points": 1
                }

            else:

                return {
                    "title":
                        f"{topic} — Дамыту",

                    "task_text":
                        "Сабақ 09:15-те басталып, "
                        "10:40-та аяқталды. "
                        "Сабақ қанша уақытқа созылды?",

                    "correct_answer":
                        "1 сағат 25 минут",

                    "explanation":
                        "09:15-тен 10:15-ке дейін "
                        "1 сағат, тағы 25 минут. "
                        "Барлығы 1 сағат 25 минут.",

                    "max_points": 2
                }


        # -------------------------------------------------
        # 2. АУДАН / АР / ГЕКТАР
        # -------------------------------------------------

        if (
            topic_is_area
            or (
                not any([
                    topic_is_time,
                    topic_is_length,
                    topic_is_mass,
                    topic_is_volume
                ])
                and "аудан" in objective_lower
            )
        ):

            if path_level == "Қолдау":

                return {
                    "title":
                        f"{topic} — Қолдау",

                    "task_text":
                        "1 гектарда неше ар бар?",

                    "correct_answer":
                        "100 ар",

                    "explanation":
                        "1 гектар = 100 ар.",

                    "max_points": 1
                }

            elif path_level == "Бекіту":

                return {
                    "title":
                        f"{topic} — Бекіту",

                    "task_text":
                        "3 гектарды армен өрнекте.",

                    "correct_answer":
                        "300 ар",

                    "explanation":
                        "1 гектар = 100 ар. "
                        "3 × 100 = 300 ар.",

                    "max_points": 1
                }

            else:

                return {
                    "title":
                        f"{topic} — Дамыту",

                    "task_text":
                        "Жер телімінің ауданы 2 га 35 ар. "
                        "Оны тек армен өрнекте.",

                    "correct_answer":
                        "235 ар",

                    "explanation":
                        "2 га = 200 ар. "
                        "200 + 35 = 235 ар.",

                    "max_points": 2
                }


        # Жалпы өлшем бірліктері
        if path_level == "Қолдау":

            return {
                "title":
                    f"{topic} — Қолдау",

                "task_text":
                    "1 метрде неше сантиметр бар?",

                "correct_answer":
                    "100 сантиметр",

                "explanation":
                    "1 м = 100 см.",

                "max_points": 1
            }

        elif path_level == "Бекіту":

            return {
                "title":
                    f"{topic} — Бекіту",

                "task_text":
                    "4 м 25 см-ді сантиметрмен өрнекте.",

                "correct_answer":
                    "425 см",

                "explanation":
                    "4 м = 400 см. "
                    "400 + 25 = 425 см.",

                "max_points": 1
            }

        else:

            return {
                "title":
                    f"{topic} — Дамыту",

                "task_text":
                    "Ұзындығы 3 м 40 см және "
                    "2 м 75 см екі кесіндінің "
                    "жалпы ұзындығын тап.",

                "correct_answer":
                    "6 м 15 см",

                "explanation":
                    "3 м 40 см + 2 м 75 см = "
                    "5 м 115 см = 6 м 15 см.",

                "max_points": 2
            }


    # =====================================================
    # МАТЕМАТИКА — КӨБЕЙТУ / БӨЛУ
    # =====================================================

    if (
        subject == "Математика"
        and (
            "көбейту" in objective_lower
            or "бөлу" in objective_lower
            or "көбейту" in topic_lower
            or "бөлу" in topic_lower
        )
    ):

        if path_level == "Қолдау":

            return {
                "title":
                    f"{topic} — Қолдау",

                "task_text":
                    "240 × 3 мәнін тап.",

                "correct_answer":
                    "720",

                "explanation":
                    "24 × 3 = 72, сондықтан "
                    "240 × 3 = 720.",

                "max_points": 1
            }

        elif path_level == "Бекіту":

            return {
                "title":
                    f"{topic} — Бекіту",

                "task_text":
                    "4 824 : 6 мәнін тап.",

                "correct_answer":
                    "804",

                "explanation":
                    "4 824 санын 6-ға бөлгенде "
                    "804 шығады.",

                "max_points": 1
            }

        else:

            return {
                "title":
                    f"{topic} — Дамыту",

                "task_text":
                    "Мектепке 24 қорап дәптер әкелінді. "
                    "Әр қорапта 125 дәптер бар. "
                    "Барлығы неше дәптер әкелінді?",

                "correct_answer":
                    "3000 дәптер",

                "explanation":
                    "125 × 24 = 3000.",

                "max_points": 2
            }


    # =====================================================
    # МАТЕМАТИКА — ГЕОМЕТРИЯ
    # =====================================================

    if (
        subject == "Математика"
        and (
            "геометр" in objective_lower
            or "фигура" in objective_lower
            or "көлем" in objective_lower
            or "периметр" in objective_lower
            or "аудан" in objective_lower
            or "фигура" in topic_lower
        )
    ):

        if path_level == "Қолдау":

            return {
                "title":
                    f"{topic} — Қолдау",

                "task_text":
                    "Тіктөртбұрыштың ұзындығы 6 см, "
                    "ені 4 см. Периметрін тап.",

                "correct_answer":
                    "20 см",

                "explanation":
                    "P = (6 + 4) × 2 = 20 см.",

                "max_points": 1
            }

        elif path_level == "Бекіту":

            return {
                "title":
                    f"{topic} — Бекіту",

                "task_text":
                    "Тіктөртбұрыштың ұзындығы 8 см, "
                    "ені 5 см. Ауданын тап.",

                "correct_answer":
                    "40 см²",

                "explanation":
                    "S = 8 × 5 = 40 см².",

                "max_points": 1
            }

        else:

            return {
                "title":
                    f"{topic} — Дамыту",

                "task_text":
                    "Тікбұрышты параллелепипедтің "
                    "ұзындығы 5 см, ені 4 см, "
                    "биіктігі 3 см. Көлемін тап.",

                "correct_answer":
                    "60 см³",

                "explanation":
                    "V = 5 × 4 × 3 = 60 см³.",

                "max_points": 2
            }


    # =====================================================
    # ҚАЗАҚ ТІЛІ — ДИАЛОГ / МОНОЛОГ / СӨЙЛЕУ
    # =====================================================

    if (
        subject == "Қазақ тілі"
        and (
            "диалог" in objective_lower
            or "монолог" in objective_lower
            or "сөйлеу" in objective_lower
            or "сұхбат" in objective_lower
        )
    ):

        if path_level == "Қолдау":

            return {
                "title":
                    f"{topic} — Қолдау",

                "task_text":
                    "«Менің мектебім» тақырыбы бойынша "
                    "2 сөйлем құра.",

                "correct_answer":
                    "Мағынасы дұрыс 2 байланысты сөйлем.",

                "explanation":
                    "Сөйлемдер бір тақырыпқа қатысты "
                    "және мағыналық байланыста болуы керек.",

                "max_points": 1
            }

        elif path_level == "Бекіту":

            return {
                "title":
                    f"{topic} — Бекіту",

                "task_text":
                    "Досыңмен «Кітапхана» тақырыбында "
                    "4 репликадан тұратын қысқа диалог құра.",

                "correct_answer":
                    "Тақырыпқа сай сұрақ-жауап түріндегі диалог.",

                "explanation":
                    "Диалогте екі адамның өзара "
                    "сұрақ-жауабы болуы қажет.",

                "max_points": 1
            }

        else:

            return {
                "title":
                    f"{topic} — Дамыту",

                "task_text":
                    "«Табиғатты қорғау» тақырыбында "
                    "5–6 сөйлемнен тұратын монолог құра.",

                "correct_answer":
                    "Тақырыпты ашатын жүйелі 5–6 сөйлем.",

                "explanation":
                    "Монологтің кіріспесі, негізгі ойы "
                    "және қорытындысы болуы керек.",

                "max_points": 2
            }


    # =====================================================
    # ҚАЗАҚ ТІЛІ — ҚОСЫМША / ҮНДЕСТІК
    # =====================================================

    if (
        subject == "Қазақ тілі"
        and (
            "қосымша" in objective_lower
            or "үндест" in objective_lower
            or "буын" in objective_lower
            or "дыбыс" in objective_lower
        )
    ):

        if path_level == "Қолдау":

            return {
                "title":
                    f"{topic} — Қолдау",

                "task_text":
                    "«бала» сөзіне көптік "
                    "жалғауын дұрыс жалға.",

                "correct_answer":
                    "балалар",

                "explanation":
                    "«Бала» — жуан буынды сөз, "
                    "сондықтан -лар жалғанады.",

                "max_points": 1
            }

        elif path_level == "Бекіту":

            return {
                "title":
                    f"{topic} — Бекіту",

                "task_text":
                    "«мектеп» және «кітап» сөздеріне "
                    "көптік жалғауын жалға.",

                "correct_answer":
                    "мектептер, кітаптар",

                "explanation":
                    "Қосымша сөздің соңғы буынына "
                    "сәйкес жалғанады.",

                "max_points": 1
            }

        else:

            return {
                "title":
                    f"{topic} — Дамыту",

                "task_text":
                    "«оқушы», «мектеп», «қалам» "
                    "сөздеріне көптік жалғауын жалғап, "
                    "неліктен сол қосымшаны "
                    "таңдағаныңды түсіндір.",

                "correct_answer":
                    "оқушылар, мектептер, қаламдар",

                "explanation":
                    "Жалғау сөздің буын және "
                    "дыбыс үндестігіне сәйкес таңдалады.",

                "max_points": 2
            }


    # =====================================================
    # ҚАЗАҚ ТІЛІ — МӘТІН / НЕГІЗГІ ОЙ
    # =====================================================

    if (
        subject == "Қазақ тілі"
        and (
            "мәтін" in objective_lower
            or "мазмұн" in objective_lower
            or "негізгі" in objective_lower
            or "тірек сөз" in objective_lower
        )
    ):

        if path_level == "Қолдау":

            return {
                "title":
                    f"{topic} — Қолдау",

                "task_text":
                    "«Айжан кітапханаға барып, "
                    "табиғат туралы кітап алды.» "
                    "Сөйлемдегі негізгі әрекетті ата.",

                "correct_answer":
                    "Айжан кітап алды.",

                "explanation":
                    "Негізгі әрекет мәтіндегі "
                    "басты істі білдіреді.",

                "max_points": 1
            }

        elif path_level == "Бекіту":

            return {
                "title":
                    f"{topic} — Бекіту",

                "task_text":
                    "«Балалар аулаға ағаш отырғызып, "
                    "оны суарды.» сөйлемінен "
                    "3 тірек сөз жаз.",

                "correct_answer":
                    "балалар, ағаш, суару",

                "explanation":
                    "Тірек сөздер мәтіннің "
                    "негізгі мазмұнын білдіреді.",

                "max_points": 1
            }

        else:

            return {
                "title":
                    f"{topic} — Дамыту",

                "task_text":
                    "«Оқушылар сенбілікке шығып, "
                    "мектеп ауласын тазалады және "
                    "гүлдер отырғызды.» "
                    "Негізгі ойды бір сөйлеммен жаз.",

                "correct_answer":
                    "Оқушылар мектеп ауласын "
                    "көркейтуге атсалысты.",

                "explanation":
                    "Негізгі ой жеке әрекеттерді "
                    "бір жалпы мағынаға біріктіреді.",

                "max_points": 2
            }


    # =====================================================
    # СОҢҒЫ ӘМБЕБАП FALLBACK
    # =====================================================

    if path_level == "Қолдау":

        task_text = (
            f"«{topic}» тақырыбы бойынша "
            f"оқу мақсатын түсіндіретін "
            f"қарапайым мысал келтір."
        )

        answer = (
            "Оқу мақсатына сәйкес дұрыс "
            "қарапайым мысал."
        )

        explanation = (
            "Жауап оқу мақсатының негізгі "
            "ұғымын дұрыс көрсетуі керек."
        )

        points = 1

    elif path_level == "Бекіту":

        task_text = (
            f"«{topic}» тақырыбы бойынша "
            f"алған біліміңді қолданып "
            f"бір тапсырманы орында."
        )

        answer = (
            "Оқу мақсатына сәйкес "
            "дұрыс орындалған жауап."
        )

        explanation = (
            "Оқушы оқу мақсатын стандартты "
            "жағдайда қолдануы керек."
        )

        points = 1

    else:

        task_text = (
            f"«{topic}» тақырыбындағы білімді "
            f"жаңа жағдайда қолданатын "
            f"мысал немесе шешім ұсын."
        )

        answer = (
            "Оқу мақсатын жаңа жағдайда "
            "дұрыс қолданған жауап."
        )

        explanation = (
            "Дамыту деңгейінде білімді "
            "жаңа жағдайда қолдану бағаланады."
        )

        points = 2


    return {
        "title":
            f"{objective_code} — {topic}",

        "task_text":
            task_text,

        "correct_answer":
            answer,

        "explanation":
            explanation,

        "max_points":
            points
    }





# =========================================================
# НАҚТЫ ҚОЛМЕН ДАЙЫНДАЛҒАН ШАБЛОНДАР
# =========================================================

def has_specific_task_template(
    subject,
    objective_code
):

    registry = {

        "Математика": {
            "4.1.1.1",
            "4.1.1.2",
            "4.1.2.8"
        },

        "Қазақ тілі": {
            "4.1.1.1",
            "4.1.7.1"
        }
    }

    return (
        objective_code
        in registry.get(
            subject,
            set()
        )
    )





# =========================================================
# ЖЕКЕ ТАПСЫРМА — ЖАУАПТЫ ТЕКСЕРУ
# =========================================================

def normalize_task_answer(value):

    import re

    value = (
        value
        or ""
    ).strip().lower()

    # Қазақша тыныс белгілері мен артық бос орын
    value = value.replace("²", "2")
    value = value.replace("³", "3")

    value = re.sub(
        r'[.,;:!?]',
        ' ',
        value
    )

    value = re.sub(
        r'\s+',
        ' ',
        value
    ).strip()

    return value


def compact_task_answer(value):

    value = normalize_task_answer(
        value
    )

    # Сандық жауаптар үшін:
    # 35 777 -> 35777
    return value.replace(
        " ",
        ""
    )


def check_personal_task_answer(
    student_answer,
    correct_answer
):

    student = normalize_task_answer(
        student_answer
    )

    correct = normalize_task_answer(
        correct_answer
    )

    if not student or not correct:
        return False

    # Дәл мәтін
    if student == correct:
        return True

    # Бос орынсыз салыстыру
    if (
        compact_task_answer(student)
        ==
        compact_task_answer(correct)
    ):
        return True

    return False





# =========================================================
# ЖЕКЕ ТАПСЫРМАНЫ АШУ
# =========================================================

@app.route(
    "/personal-tasks/<int:task_id>"
)
def personal_task_detail(task_id):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    task = conn.execute(
        """
        SELECT
            pt.*,

            s.full_name,
            s.class_name,

            lo.objective_code,
            lo.objective_text

        FROM personal_tasks pt

        JOIN students s
            ON s.id = pt.student_id

        JOIN learning_objectives lo
            ON lo.id = pt.learning_objective_id

        WHERE pt.id = ?
        """,
        (task_id,)
    ).fetchone()

    conn.close()

    if task is None:
        return "Жеке тапсырма табылмады", 404

    return render_template(
        "personal_task_detail.html",
        task=task
    )


# =========================================================
# ЖЕКЕ ТАПСЫРМАҒА ЖАУАП БЕРУ
# =========================================================

@app.route(
    "/personal-tasks/<int:task_id>/submit",
    methods=["POST"]
)
def submit_personal_task(task_id):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    task = conn.execute(
        """
        SELECT *
        FROM personal_tasks
        WHERE id = ?
        """,
        (task_id,)
    ).fetchone()

    if task is None:

        conn.close()

        return "Жеке тапсырма табылмады", 404


    student_answer = (
        request.form.get(
            "student_answer",
            ""
        )
        .strip()
    )


    if not student_answer:

        conn.close()

        return redirect(
            url_for(
                "personal_task_detail",
                task_id=task_id
            )
        )


    is_correct = check_personal_task_answer(
        student_answer,
        task["correct_answer"]
    )


    earned_points = (
        int(task["max_points"] or 0)
        if is_correct
        else 0
    )


    conn.execute(
        """
        UPDATE personal_tasks

        SET
            student_answer = ?,
            is_correct = ?,
            earned_points = ?,
            attempt_count =
                COALESCE(attempt_count, 0) + 1,

            status = 'completed',

            last_answered_at =
                CURRENT_TIMESTAMP,

            completed_at =
                CURRENT_TIMESTAMP

        WHERE id = ?
        """,
        (
            student_answer,
            1 if is_correct else 0,
            earned_points,
            task_id
        )
    )

    conn.commit()
    conn.close()


    return redirect(
        url_for(
            "personal_task_detail",
            task_id=task_id
        )
    )





# =========================================================
# НӘТИЖЕЛЕР ЖӘНЕ ПРОГРЕСС
# =========================================================

@app.route("/progress")
def progress():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    # -----------------------------------------------------
    # ЖАЛПЫ СТАТИСТИКА
    # -----------------------------------------------------

    total_students = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM students
        WHERE class_name = ?
        """,
        ('4 «Д»',)
    ).fetchone()["total"]


    completed_attempts = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM diagnostic_attempts
        WHERE status = 'completed'
        """
    ).fetchone()["total"]


    average_diagnostic = conn.execute(
        """
        SELECT ROUND(
            AVG(percentage),
            1
        ) AS average
        FROM diagnostic_attempts
        WHERE status = 'completed'
        """
    ).fetchone()["average"]

    average_diagnostic = (
        average_diagnostic
        if average_diagnostic is not None
        else 0
    )


    # -----------------------------------------------------
    # ОҚУ МАҚСАТТАРЫ
    # -----------------------------------------------------

    objective_stats = conn.execute(
        """
        SELECT
            SUM(
                CASE
                    WHEN mastery_status = 'Меңгерілген'
                    THEN 1 ELSE 0
                END
            ) AS mastered,

            SUM(
                CASE
                    WHEN mastery_status =
                         'Бекітуді қажет етеді'
                    THEN 1 ELSE 0
                END
            ) AS reinforcement,

            SUM(
                CASE
                    WHEN mastery_status = 'Олқылық бар'
                    THEN 1 ELSE 0
                END
            ) AS gaps

        FROM objective_results
        """
    ).fetchone()


    mastered_count = (
        objective_stats["mastered"]
        or 0
    )

    reinforcement_count = (
        objective_stats["reinforcement"]
        or 0
    )

    gap_count = (
        objective_stats["gaps"]
        or 0
    )


    # -----------------------------------------------------
    # ЖЕКЕ ТАПСЫРМАЛАР
    # -----------------------------------------------------

    task_stats = conn.execute(
        """
        SELECT
            COUNT(*) AS total,

            SUM(
                CASE
                    WHEN status = 'completed'
                    THEN 1 ELSE 0
                END
            ) AS completed,

            SUM(
                CASE
                    WHEN status = 'completed'
                         AND is_correct = 1
                    THEN 1 ELSE 0
                END
            ) AS correct

        FROM personal_tasks
        """
    ).fetchone()


    total_tasks = (
        task_stats["total"]
        or 0
    )

    completed_tasks = (
        task_stats["completed"]
        or 0
    )

    correct_tasks = (
        task_stats["correct"]
        or 0
    )


    task_completion = (
        round(
            completed_tasks
            / total_tasks
            * 100,
            1
        )
        if total_tasks
        else 0
    )


    # -----------------------------------------------------
    # ӘР ОҚУШЫНЫҢ ПРОГРЕСІ
    # -----------------------------------------------------

    students = conn.execute(
        """
        SELECT
            s.id,
            s.full_name,
            s.class_name,

            COUNT(
                DISTINCT a.id
            ) AS diagnostic_count,

            ROUND(
                AVG(
                    CASE
                        WHEN a.status = 'completed'
                        THEN a.percentage
                    END
                ),
                1
            ) AS average_percentage,

            COUNT(
                DISTINCT pt.id
            ) AS personal_task_count,

            COUNT(
                DISTINCT CASE
                    WHEN pt.status = 'completed'
                    THEN pt.id
                END
            ) AS completed_task_count,

            COUNT(
                DISTINCT CASE
                    WHEN pt.status = 'completed'
                         AND pt.is_correct = 1
                    THEN pt.id
                END
            ) AS correct_task_count

        FROM students s

        LEFT JOIN diagnostic_attempts a
            ON a.student_id = s.id
            AND a.status = 'completed'

        LEFT JOIN personal_tasks pt
            ON pt.student_id = s.id

        WHERE s.class_name = ?

        GROUP BY
            s.id,
            s.full_name,
            s.class_name

        ORDER BY s.full_name
        """,
        ('4 «Д»',)
    ).fetchall()


    student_progress = []

    for student in students:

        task_count = (
            student["personal_task_count"]
            or 0
        )

        done_count = (
            student["completed_task_count"]
            or 0
        )

        completion_percentage = (
            round(
                done_count
                / task_count
                * 100,
                1
            )
            if task_count
            else 0
        )

        student_progress.append({
            "id":
                student["id"],

            "full_name":
                student["full_name"],

            "class_name":
                student["class_name"],

            "diagnostic_count":
                student["diagnostic_count"]
                or 0,

            "average_percentage":
                student["average_percentage"]
                or 0,

            "personal_task_count":
                task_count,

            "completed_task_count":
                done_count,

            "correct_task_count":
                student["correct_task_count"]
                or 0,

            "completion_percentage":
                completion_percentage
        })


    # -----------------------------------------------------
    # СОҢҒЫ НӘТИЖЕЛЕР
    # -----------------------------------------------------

    recent_results = conn.execute(
        """
        SELECT
            a.id AS attempt_id,
            a.percentage,
            a.earned_points,
            a.max_points,
            a.completed_at,

            s.full_name,

            d.title,
            d.subject

        FROM diagnostic_attempts a

        JOIN students s
            ON s.id = a.student_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        WHERE a.status = 'completed'

        ORDER BY
            a.completed_at DESC,
            a.id DESC

        LIMIT 10
        """
    ).fetchall()


    conn.close()


    return render_template(
        "progress.html",

        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),

        total_students=total_students,
        completed_attempts=completed_attempts,
        average_diagnostic=average_diagnostic,

        mastered_count=mastered_count,
        reinforcement_count=reinforcement_count,
        gap_count=gap_count,

        total_tasks=total_tasks,
        completed_tasks=completed_tasks,
        correct_tasks=correct_tasks,
        task_completion=task_completion,

        students=student_progress,
        recent_results=recent_results
    )





# =========================================================
# АҚЫЛДЫ САБАҚ
# =========================================================

@app.route("/smart-lessons")
def smart_lessons():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    subjects = conn.execute(
        """
        SELECT DISTINCT subject
        FROM learning_objectives
        WHERE
            subject IS NOT NULL
            AND TRIM(subject) != ''
        ORDER BY subject
        """
    ).fetchall()

    conn.close()

    return render_template(
        "smart_lessons.html",

        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),

        subjects=subjects
    )


# =========================================================
# АҚЫЛДЫ САБАҚ — КТЖ ТАҚЫРЫПТАРЫ API
# =========================================================

@app.route("/api/smart-lessons/topics")
def smart_lesson_topics_api():

    if not logged_in():
        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401

    subject = (
        request.args.get(
            "subject",
            ""
        )
        .strip()
    )

    if not subject:

        return jsonify({
            "ok": True,
            "topics": []
        })

    conn = get_db()

    rows = conn.execute(
        """
        SELECT
            lesson_topic,

            MIN(id) AS first_id,

            COUNT(*) AS objective_count

        FROM learning_objectives

        WHERE
            subject = ?
            AND lesson_topic IS NOT NULL
            AND TRIM(lesson_topic) != ''

        GROUP BY lesson_topic

        ORDER BY first_id
        """,
        (subject,)
    ).fetchall()

    conn.close()

    topics = [
        {
            "lesson_topic":
                row["lesson_topic"],

            "objective_count":
                row["objective_count"]
        }
        for row in rows
    ]

    return jsonify({
        "ok": True,
        "subject": subject,
        "topics": topics
    })


# =========================================================
# АҚЫЛДЫ САБАҚ — ОҚУ МАҚСАТТАРЫ API
# =========================================================

@app.route("/api/smart-lessons/objectives")
def smart_lesson_objectives_api():

    if not logged_in():
        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401

    subject = (
        request.args.get(
            "subject",
            ""
        )
        .strip()
    )

    lesson_topic = (
        request.args.get(
            "lesson_topic",
            ""
        )
        .strip()
    )

    if not subject or not lesson_topic:

        return jsonify({
            "ok": True,
            "objectives": []
        })

    conn = get_db()

    rows = conn.execute(
        """
        SELECT
            id,
            objective_code,
            objective_text,
            lesson_topic,
            section

        FROM learning_objectives

        WHERE
            subject = ?
            AND lesson_topic = ?

        ORDER BY id
        """,
        (
            subject,
            lesson_topic
        )
    ).fetchall()

    conn.close()


    # Бір сабақта бір код бірнеше рет кездессе,
    # интерфейске қайталамаймыз.
    unique = []
    seen = set()

    for row in rows:

        key = (
            row["objective_code"],
            row["objective_text"]
        )

        if key in seen:
            continue

        seen.add(key)

        unique.append({
            "id":
                row["id"],

            "objective_code":
                row["objective_code"],

            "objective_text":
                row["objective_text"],

            "lesson_topic":
                row["lesson_topic"],

            "section":
                row["section"] or ""
        })


    return jsonify({
        "ok": True,
        "subject": subject,
        "lesson_topic": lesson_topic,
        "objectives": unique
    })





# =========================================================
# АҚЫЛДЫ САБАҚ — ДЕҢГЕЙЛІК ТАПСЫРМАЛАР API
# =========================================================

@app.route(
    "/api/smart-lessons/generate-tasks",
    methods=["POST"]
)
def smart_lesson_generate_tasks_api():

    if not logged_in():
        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401

    data = request.get_json(
        silent=True
    ) or {}

    subject = (
        data.get(
            "subject",
            ""
        )
        or ""
    ).strip()

    lesson_topic = (
        data.get(
            "lesson_topic",
            ""
        )
        or ""
    ).strip()

    objective_ids = (
        data.get(
            "objective_ids",
            []
        )
        or []
    )

    if not subject:

        return jsonify({
            "ok": False,
            "error": "Пән таңдалмаған"
        }), 400

    if not lesson_topic:

        return jsonify({
            "ok": False,
            "error": "Сабақ тақырыбы таңдалмаған"
        }), 400

    if not objective_ids:

        return jsonify({
            "ok": False,
            "error": "Оқу мақсаты таңдалмаған"
        }), 400


    # -----------------------------------------------------
    # ID-ЛЕРДІ ҚАУІПСІЗ INTEGER-ГЕ АЙНАЛДЫРУ
    # -----------------------------------------------------

    clean_ids = []

    for value in objective_ids:

        try:

            objective_id = int(
                value
            )

        except (
            TypeError,
            ValueError
        ):

            continue

        if (
            objective_id > 0
            and objective_id not in clean_ids
        ):

            clean_ids.append(
                objective_id
            )


    if not clean_ids:

        return jsonify({
            "ok": False,
            "error": "Жарамды оқу мақсаты жоқ"
        }), 400


    conn = get_db()

    placeholders = ",".join(
        "?"
        for _ in clean_ids
    )

    rows = conn.execute(
        f"""
        SELECT
            id,
            subject,
            objective_code,
            objective_text,
            lesson_topic,
            section

        FROM learning_objectives

        WHERE
            id IN ({placeholders})
            AND subject = ?

        ORDER BY id
        """,
        tuple(
            clean_ids
            + [subject]
        )
    ).fetchall()


    if not rows:

        conn.close()

        return jsonify({
            "ok": False,
            "error": "Оқу мақсаттары табылмады"
        }), 404


    # -----------------------------------------------------
    # ӘР ОҚУ МАҚСАТЫНА 3 ДЕҢГЕЙ
    # -----------------------------------------------------

    generated_objectives = []


    for objective in rows:

        context = get_ktj_context(
            conn,
            subject,
            objective["id"]
        )

        lesson_topics = (
            context.get(
                "lesson_topics",
                []
            )
            or []
        )

        sections = (
            context.get(
                "sections",
                []
            )
            or []
        )


        # Мұғалім интерфейсте таңдаған нақты сабақ
        # бірінші басымдыққа ие.
        main_topic = lesson_topic

        if not main_topic:

            main_topic = select_best_ktj_topic(
                objective["objective_text"],
                lesson_topics
            )


        tasks = []


        for level in [
            "Қолдау",
            "Бекіту",
            "Дамыту"
        ]:

            if has_specific_task_template(
                subject,
                objective["objective_code"]
            ):

                task = build_personal_task(
                    subject,
                    objective["objective_code"],
                    objective["objective_text"],
                    level
                )

                generator_type = (
                    "specific_template"
                )

            else:

                task = build_ktj_fallback_task(
                    subject,
                    objective["objective_code"],
                    objective["objective_text"],
                    level,
                    main_topic,
                    lesson_topics,
                    sections
                )

                generator_type = (
                    "ktj_fallback"
                )


            tasks.append({

                "level":
                    level,

                "title":
                    task.get(
                        "title",
                        ""
                    ),

                "task_text":
                    task.get(
                        "task_text",
                        ""
                    ),

                "correct_answer":
                    task.get(
                        "correct_answer",
                        ""
                    ),

                "explanation":
                    task.get(
                        "explanation",
                        ""
                    ),

                "max_points":
                    task.get(
                        "max_points",
                        1
                    ),

                "generator_type":
                    generator_type
            })


        generated_objectives.append({

            "objective_id":
                objective["id"],

            "objective_code":
                objective["objective_code"],

            "objective_text":
                objective["objective_text"],

            "lesson_topic":
                main_topic,

            "section":
                objective["section"] or "",

            "tasks":
                tasks
        })


    conn.close()


    return jsonify({

        "ok":
            True,

        "subject":
            subject,

        "lesson_topic":
            lesson_topic,

        "objectives":
            generated_objectives
    })





# =========================================================
# АҚЫЛДЫ САБАҚ — ТОЛЫҚ САБАҚ ЖОСПАРЫ
# =========================================================

def build_smart_lesson_plan(
    subject,
    lesson_topic,
    objectives,
    duration=45
):

    subject = (
        subject
        or ""
    ).strip()

    lesson_topic = (
        lesson_topic
        or ""
    ).strip()

    try:
        duration = int(
            duration
        )
    except (
        TypeError,
        ValueError
    ):
        duration = 45

    if duration not in (
        40,
        45
    ):
        duration = 45


    # -----------------------------------------------------
    # УАҚЫТТЫ БӨЛУ
    # -----------------------------------------------------

    if duration == 40:

        start_minutes = 5
        explanation_minutes = 10
        practice_minutes = 18
        reflection_minutes = 7

    else:

        start_minutes = 5
        explanation_minutes = 10
        practice_minutes = 20
        reflection_minutes = 10


    # -----------------------------------------------------
    # ОҚУ МАҚСАТТАРЫ
    # -----------------------------------------------------

    objective_codes = []

    objective_texts = []


    for objective in objectives:

        code = (
            objective.get(
                "objective_code",
                ""
            )
            or ""
        ).strip()

        text = (
            objective.get(
                "objective_text",
                ""
            )
            or ""
        ).strip()


        if (
            code
            and code not in objective_codes
        ):
            objective_codes.append(
                code
            )


        if (
            text
            and text not in objective_texts
        ):
            objective_texts.append(
                text
            )


    # -----------------------------------------------------
    # САБАҚ МАҚСАТЫ
    # -----------------------------------------------------

    if len(objective_texts) == 1:

        lesson_goal = (
            "Оқушылар "
            + objective_texts[0]
            .rstrip(
                ".;"
            )
            + " дағдысын қалыптастырады."
        )

    else:

        lesson_goal = (
            "Оқушылар сабақ барысында "
            "таңдалған оқу мақсаттарына сәйкес "
            "негізгі білім мен дағдыларды "
            "меңгереді және қолданады."
        )


    # -----------------------------------------------------
    # САБАҚТЫҢ БАСЫ
    # -----------------------------------------------------

    lesson_start = {
        "title":
            "Сабақтың басы",

        "minutes":
            start_minutes,

        "teacher_action":
            (
                "Оқушылармен амандасады, "
                "сабаққа назарын аударады. "
                f"«{lesson_topic}» тақырыбына "
                "байланысты алдыңғы білімді "
                "еске түсіретін қысқа сұрақтар қояды."
            ),

        "student_action":
            (
                "Сұрақтарға жауап береді, "
                "алдыңғы білімін еске түсіреді "
                "және сабақ тақырыбын анықтайды."
            ),

        "assessment":
            "Ауызша мадақтау және жедел кері байланыс."
    }


    # -----------------------------------------------------
    # ЖАҢА БІЛІМ
    # -----------------------------------------------------

    if subject == "Математика":

        explanation = (
            f"«{lesson_topic}» тақырыбының "
            "негізгі ережесін түсіндіреді. "
            "Математикалық мысалдарды кезең-кезеңімен "
            "орындап көрсетеді. Оқушылармен бірге "
            "шешу алгоритмін талдайды."
        )

        student_explanation = (
            "Мысалдарды талдайды, ережені анықтайды, "
            "шешу қадамдарын түсіндіреді және "
            "мұғаліммен бірге үлгі тапсырманы орындайды."
        )

    elif subject == "Қазақ тілі":

        explanation = (
            f"«{lesson_topic}» тақырыбы бойынша "
            "тілдік ұғымды немесе ережені "
            "мысалдар арқылы түсіндіреді. "
            "Сөз, сөйлем немесе мәтін үлгілерін "
            "бірге талдайды."
        )

        student_explanation = (
            "Берілген тілдік материалды талдайды, "
            "ережені анықтайды, мысал келтіреді "
            "және қорытынды жасайды."
        )

    else:

        explanation = (
            f"«{lesson_topic}» тақырыбы бойынша "
            "негізгі ұғымдарды мысалдар арқылы "
            "түсіндіреді."
        )

        student_explanation = (
            "Жаңа ақпаратты талдайды және "
            "мысалдар арқылы түсіндіреді."
        )


    new_knowledge = {
        "title":
            "Жаңа білім",

        "minutes":
            explanation_minutes,

        "teacher_action":
            explanation,

        "student_action":
            student_explanation,

        "assessment":
            (
                "Сұрақ-жауап, түсіндірме және "
                "қысқа қалыптастырушы бағалау."
            )
    }


    # -----------------------------------------------------
    # БАҒАЛАУ КРИТЕРИЙЛЕРІ
    # -----------------------------------------------------

    criteria = []

    descriptors = []


    for index, objective in enumerate(
        objectives,
        start=1
    ):

        code = (
            objective.get(
                "objective_code",
                ""
            )
            or ""
        )

        text = (
            objective.get(
                "objective_text",
                ""
            )
            or ""
        ).rstrip(
            ".;"
        )


        criteria.append({
            "number":
                index,

            "objective_code":
                code,

            "text":
                (
                    f"{text} бойынша "
                    "тапсырмаларды дұрыс орындайды."
                )
        })


        descriptors.append({
            "objective_code":
                code,

            "items": [
                (
                    "тапсырманың шартын "
                    "дұрыс түсінеді"
                ),
                (
                    "тиісті ереже немесе "
                    "алгоритмді қолданады"
                ),
                (
                    "жауабын дұрыс анықтайды"
                )
            ]
        })


    # -----------------------------------------------------
    # РЕФЛЕКСИЯ
    # -----------------------------------------------------

    reflection = {
        "title":
            "Бағалау және рефлексия",

        "minutes":
            reflection_minutes,

        "teacher_action":
            (
                "Оқу мақсатына жету деңгейін "
                "қорытындылайды. Оқушыларға "
                "қысқа рефлексия жүргізеді."
            ),

        "student_action":
            (
                "Өз жұмысын бағалайды және "
                "«Бүгін білдім...», "
                "«Маған қиын болды...», "
                "«Келесіде жақсартамын...» "
                "сөйлемдерін аяқтайды."
            ),

        "assessment":
            (
                "Өзін-өзі бағалау және "
                "мұғалімнің қорытынды кері байланысы."
            )
    }


    # -----------------------------------------------------
    # ҮЙ ТАПСЫРМАСЫ
    # -----------------------------------------------------

    if subject == "Математика":

        homework = (
            f"«{lesson_topic}» тақырыбы бойынша "
            "сабақта орындалған тапсырмаларға "
            "ұқсас 2 есеп орындау және "
            "бір есептің шешу жолын түсіндіру."
        )

    elif subject == "Қазақ тілі":

        homework = (
            f"«{lesson_topic}» тақырыбы бойынша "
            "ережені қайталау және "
            "3 мысал жазу."
        )

    else:

        homework = (
            f"«{lesson_topic}» тақырыбын қайталау "
            "және қысқа тапсырма орындау."
        )


    return {

        "subject":
            subject,

        "class_name":
            '4 «Д»',

        "lesson_topic":
            lesson_topic,

        "duration":
            duration,

        "objective_codes":
            objective_codes,

        "objective_texts":
            objective_texts,

        "lesson_goal":
            lesson_goal,

        "stages": {
            "start":
                lesson_start,

            "new_knowledge":
                new_knowledge,

            "practice": {
                "title":
                    "Деңгейлік тапсырмалар",

                "minutes":
                    practice_minutes
            },

            "reflection":
                reflection
        },

        "criteria":
            criteria,

        "descriptors":
            descriptors,

        "homework":
            homework
    }





# =========================================================
# АҚЫЛДЫ САБАҚ — ТОЛЫҚ ЖОСПАР API
# =========================================================

@app.route(
    "/api/smart-lessons/generate-plan",
    methods=["POST"]
)
def smart_lesson_generate_plan_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    data = request.get_json(
        silent=True
    ) or {}


    subject = (
        data.get(
            "subject",
            ""
        )
        or ""
    ).strip()


    lesson_topic = (
        data.get(
            "lesson_topic",
            ""
        )
        or ""
    ).strip()


    duration = data.get(
        "duration",
        45
    )


    objective_ids = (
        data.get(
            "objective_ids",
            []
        )
        or []
    )


    if not subject:

        return jsonify({
            "ok": False,
            "error": "Пән таңдалмаған"
        }), 400


    if not lesson_topic:

        return jsonify({
            "ok": False,
            "error": "Сабақ тақырыбы таңдалмаған"
        }), 400


    # -----------------------------------------------------
    # OBJECTIVE IDS
    # -----------------------------------------------------

    clean_ids = []

    for value in objective_ids:

        try:

            objective_id = int(
                value
            )

        except (
            TypeError,
            ValueError
        ):

            continue


        if (
            objective_id > 0
            and objective_id not in clean_ids
        ):

            clean_ids.append(
                objective_id
            )


    if not clean_ids:

        return jsonify({
            "ok": False,
            "error": "Оқу мақсаты таңдалмаған"
        }), 400


    conn = get_db()


    placeholders = ",".join(
        "?"
        for _ in clean_ids
    )


    rows = conn.execute(
        f"""
        SELECT
            id,
            subject,
            objective_code,
            objective_text,
            lesson_topic,
            section

        FROM learning_objectives

        WHERE
            id IN ({placeholders})
            AND subject = ?

        ORDER BY id
        """,
        tuple(
            clean_ids
            + [subject]
        )
    ).fetchall()


    if not rows:

        conn.close()

        return jsonify({
            "ok": False,
            "error": "Оқу мақсаттары табылмады"
        }), 404


    objectives = [
        {
            "id":
                row["id"],

            "objective_code":
                row["objective_code"],

            "objective_text":
                row["objective_text"],

            "section":
                row["section"] or ""
        }
        for row in rows
    ]


    # -----------------------------------------------------
    # САБАҚ ЖОСПАРЫ
    # -----------------------------------------------------

    plan = build_smart_lesson_plan(
        subject,
        lesson_topic,
        objectives,
        duration
    )


    # -----------------------------------------------------
    # ДЕҢГЕЙЛІК ТАПСЫРМАЛАР
    # -----------------------------------------------------

    generated_tasks = []


    for objective in rows:

        context = get_ktj_context(
            conn,
            subject,
            objective["id"]
        )


        lesson_topics = (
            context.get(
                "lesson_topics",
                []
            )
            or []
        )


        sections = (
            context.get(
                "sections",
                []
            )
            or []
        )


        tasks = []


        for level in [
            "Қолдау",
            "Бекіту",
            "Дамыту"
        ]:

            if has_specific_task_template(
                subject,
                objective["objective_code"]
            ):

                task = build_personal_task(
                    subject,
                    objective["objective_code"],
                    objective["objective_text"],
                    level
                )

                generator_type = (
                    "specific_template"
                )

            else:

                task = build_ktj_fallback_task(
                    subject,
                    objective["objective_code"],
                    objective["objective_text"],
                    level,
                    lesson_topic,
                    lesson_topics,
                    sections
                )

                generator_type = (
                    "ktj_fallback"
                )


            tasks.append({

                "level":
                    level,

                "title":
                    task.get(
                        "title",
                        ""
                    ),

                "task_text":
                    task.get(
                        "task_text",
                        ""
                    ),

                "correct_answer":
                    task.get(
                        "correct_answer",
                        ""
                    ),

                "explanation":
                    task.get(
                        "explanation",
                        ""
                    ),

                "max_points":
                    task.get(
                        "max_points",
                        1
                    ),

                "generator_type":
                    generator_type
            })


        generated_tasks.append({

            "objective_id":
                objective["id"],

            "objective_code":
                objective["objective_code"],

            "objective_text":
                objective["objective_text"],

            "tasks":
                tasks
        })


    conn.close()


    plan[
        "level_tasks"
    ] = generated_tasks


    return jsonify({
        "ok": True,
        "plan": plan
    })





# =========================================================
# АҚЫЛДЫ САБАҚ — САҚТАУ
# =========================================================

@app.route(
    "/api/smart-lessons/save",
    methods=["POST"]
)
def save_smart_lesson_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    data = request.get_json(
        silent=True
    ) or {}


    plan = (
        data.get(
            "plan",
            {}
        )
        or {}
    )


    subject = (
        plan.get(
            "subject",
            ""
        )
        or ""
    ).strip()


    lesson_topic = (
        plan.get(
            "lesson_topic",
            ""
        )
        or ""
    ).strip()


    if not subject:

        return jsonify({
            "ok": False,
            "error": "Пән көрсетілмеген"
        }), 400


    if not lesson_topic:

        return jsonify({
            "ok": False,
            "error": "Сабақ тақырыбы көрсетілмеген"
        }), 400


    try:

        duration = int(
            plan.get(
                "duration",
                45
            )
        )

    except (
        TypeError,
        ValueError
    ):

        duration = 45


    objective_codes = (
        plan.get(
            "objective_codes",
            []
        )
        or []
    )


    objective_texts = (
        plan.get(
            "objective_texts",
            []
        )
        or []
    )


    lesson_goal = (
        plan.get(
            "lesson_goal",
            ""
        )
        or ""
    )


    # JSON ішінде Unicode қазақша
    # күйінде сақталады.
    plan_json = json.dumps(
        plan,
        ensure_ascii=False
    )


    conn = get_db()


    cursor = conn.execute(
        """
        INSERT INTO smart_lessons
        (
            subject,
            class_name,
            lesson_topic,
            duration,

            objective_codes,
            objective_texts,

            lesson_goal,
            plan_json,

            status,
            created_at,
            updated_at
        )

        VALUES (
            ?, ?, ?, ?,
            ?, ?,
            ?, ?,
            'saved',
            CURRENT_TIMESTAMP,
            CURRENT_TIMESTAMP
        )
        """,
        (
            subject,
            plan.get(
                "class_name",
                '4 «Д»'
            ),
            lesson_topic,
            duration,

            json.dumps(
                objective_codes,
                ensure_ascii=False
            ),

            json.dumps(
                objective_texts,
                ensure_ascii=False
            ),

            lesson_goal,
            plan_json
        )
    )


    lesson_id = (
        cursor.lastrowid
    )


    conn.commit()
    conn.close()


    return jsonify({
        "ok": True,
        "lesson_id": lesson_id,
        "message": "Сабақ сақталды"
    })


# =========================================================
# МЕНІҢ САБАҚТАРЫМ
# =========================================================

@app.route("/smart-lessons/saved")
def saved_smart_lessons():

    if not logged_in():
        return redirect(
            url_for("login")
        )


    conn = get_db()


    rows = conn.execute(
        """
        SELECT
            id,
            subject,
            class_name,
            lesson_topic,
            duration,
            objective_codes,
            lesson_goal,
            status,
            created_at,
            updated_at

        FROM smart_lessons

        ORDER BY
            created_at DESC,
            id DESC
        """
    ).fetchall()


    lessons = []


    for row in rows:

        try:

            codes = json.loads(
                row["objective_codes"]
                or "[]"
            )

        except Exception:

            codes = []


        lessons.append({

            "id":
                row["id"],

            "subject":
                row["subject"],

            "class_name":
                row["class_name"],

            "lesson_topic":
                row["lesson_topic"],

            "duration":
                row["duration"],

            "objective_codes":
                codes,

            "lesson_goal":
                row["lesson_goal"],

            "status":
                row["status"],

            "created_at":
                row["created_at"],

            "updated_at":
                row["updated_at"]
        })


    conn.close()


    return render_template(
        "saved_smart_lessons.html",

        teacher_name=session.get(
            "teacher_name",
            "Мұғалім"
        ),

        lessons=lessons
    )


# =========================================================
# САҚТАЛҒАН САБАҚТЫ АШУ
# =========================================================

@app.route(
    "/smart-lessons/saved/<int:lesson_id>"
)
def saved_smart_lesson_detail(
    lesson_id
):

    if not logged_in():
        return redirect(
            url_for("login")
        )


    conn = get_db()


    row = conn.execute(
        """
        SELECT *
        FROM smart_lessons
        WHERE id = ?
        """,
        (
            lesson_id,
        )
    ).fetchone()


    conn.close()


    if row is None:

        return (
            "Сақталған сабақ табылмады",
            404
        )


    try:

        plan = json.loads(
            row["plan_json"]
            or "{}"
        )

    except Exception:

        plan = {}


    return render_template(
        "saved_smart_lesson_detail.html",

        lesson=row,
        plan=plan
    )





# =========================================================
# САҚТАЛҒАН АҚЫЛДЫ САБАҚ — WORD EXPORT
# =========================================================

@app.route(
    "/smart-lessons/saved/<int:lesson_id>/word"
)
def export_saved_smart_lesson_word(
    lesson_id
):

    if not logged_in():
        return redirect(
            url_for("login")
        )


    # -----------------------------------------------------
    # IMPORT
    # -----------------------------------------------------

    from docx import Document

    from docx.shared import (
        Pt,
        Cm
    )

    from docx.enum.text import (
        WD_ALIGN_PARAGRAPH
    )

    from docx.enum.table import (
        WD_TABLE_ALIGNMENT,
        WD_CELL_VERTICAL_ALIGNMENT
    )

    from docx.oxml import OxmlElement

    from docx.oxml.ns import qn

    import tempfile
    import re


    # -----------------------------------------------------
    # DATABASE
    # -----------------------------------------------------

    conn = get_db()

    row = conn.execute(
        """
        SELECT *
        FROM smart_lessons
        WHERE id = ?
        """,
        (
            lesson_id,
        )
    ).fetchone()

    conn.close()


    if row is None:

        return (
            "Сақталған сабақ табылмады",
            404
        )


    try:

        plan = json.loads(
            row["plan_json"]
            or "{}"
        )

    except Exception:

        plan = {}


    if not plan:

        return (
            "Сабақ жоспарының деректері жоқ",
            400
        )


    # -----------------------------------------------------
    # DOCUMENT
    # -----------------------------------------------------

    document = Document()


    section = document.sections[0]

    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)


    # -----------------------------------------------------
    # NORMAL STYLE
    # -----------------------------------------------------

    normal = document.styles["Normal"]

    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    normal._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman"
    )


    # -----------------------------------------------------
    # HELPERS
    # -----------------------------------------------------

    def set_cell_text(
        cell,
        text,
        bold=False,
        size=10
    ):

        cell.text = ""

        paragraph = (
            cell.paragraphs[0]
        )

        run = paragraph.add_run(
            str(
                text
                if text is not None
                else ""
            )
        )

        run.bold = bold
        run.font.name = (
            "Times New Roman"
        )
        run.font.size = Pt(size)

        run._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            "Times New Roman"
        )

        cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )


    def shade_cell(
        cell,
        fill
    ):

        tcPr = (
            cell._tc.get_or_add_tcPr()
        )

        shd = OxmlElement(
            "w:shd"
        )

        shd.set(
            qn("w:fill"),
            fill
        )

        tcPr.append(
            shd
        )


    def add_heading_text(
        text,
        size=13
    ):

        p = document.add_paragraph()

        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(5)

        run = p.add_run(
            text
        )

        run.bold = True
        run.font.name = (
            "Times New Roman"
        )
        run.font.size = Pt(size)

        run._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            "Times New Roman"
        )

        return p


    def safe_filename(
        value
    ):

        value = (
            value
            or "sabak"
        )

        value = re.sub(
            r'[\\/:*?"<>|]',
            "_",
            value
        )

        value = re.sub(
            r"\s+",
            "_",
            value
        )

        return value[:80]


    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        "ҚЫСҚА МЕРЗІМДІ ЖОСПАР"
    )

    run.bold = True
    run.font.name = (
        "Times New Roman"
    )
    run.font.size = Pt(14)

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman"
    )


    # -----------------------------------------------------
    # BASIC INFORMATION
    # -----------------------------------------------------

    info = document.add_table(
        rows=5,
        cols=2
    )

    info.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    info.style = (
        "Table Grid"
    )


    basic_rows = [
        (
            "Пән",
            plan.get(
                "subject",
                ""
            )
        ),
        (
            "Сынып",
            plan.get(
                "class_name",
                '4 «Д»'
            )
        ),
        (
            "Сабақ тақырыбы",
            plan.get(
                "lesson_topic",
                ""
            )
        ),
        (
            "Сабақ ұзақтығы",
            (
                str(
                    plan.get(
                        "duration",
                        45
                    )
                )
                + " минут"
            )
        ),
        (
            "Сабақ мақсаты",
            plan.get(
                "lesson_goal",
                ""
            )
        )
    ]


    for i, (
        label,
        value
    ) in enumerate(
        basic_rows
    ):

        set_cell_text(
            info.cell(i, 0),
            label,
            bold=True
        )

        shade_cell(
            info.cell(i, 0),
            "EAF2F8"
        )

        set_cell_text(
            info.cell(i, 1),
            value
        )


    # -----------------------------------------------------
    # OBJECTIVES
    # -----------------------------------------------------

    add_heading_text(
        "Оқу мақсаттары"
    )


    codes = (
        plan.get(
            "objective_codes",
            []
        )
        or []
    )

    texts = (
        plan.get(
            "objective_texts",
            []
        )
        or []
    )


    objective_table = (
        document.add_table(
            rows=1,
            cols=2
        )
    )

    objective_table.style = (
        "Table Grid"
    )

    objective_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )


    set_cell_text(
        objective_table.cell(
            0,
            0
        ),
        "Код",
        bold=True
    )

    set_cell_text(
        objective_table.cell(
            0,
            1
        ),
        "Оқу мақсаты",
        bold=True
    )

    shade_cell(
        objective_table.cell(
            0,
            0
        ),
        "D9EAF7"
    )

    shade_cell(
        objective_table.cell(
            0,
            1
        ),
        "D9EAF7"
    )


    for index, code_value in enumerate(
        codes
    ):

        cells = (
            objective_table.add_row()
            .cells
        )

        set_cell_text(
            cells[0],
            code_value,
            bold=True
        )

        text_value = (
            texts[index]
            if index < len(texts)
            else ""
        )

        set_cell_text(
            cells[1],
            text_value
        )


    # -----------------------------------------------------
    # LESSON STAGES
    # -----------------------------------------------------

    add_heading_text(
        "Сабақ барысы"
    )


    stages_table = (
        document.add_table(
            rows=1,
            cols=5
        )
    )

    stages_table.style = (
        "Table Grid"
    )

    stages_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )


    headers = [
        "Сабақ кезеңі",
        "Уақыты",
        "Мұғалім әрекеті",
        "Оқушы әрекеті",
        "Бағалау"
    ]


    for i, header in enumerate(
        headers
    ):

        set_cell_text(
            stages_table.cell(
                0,
                i
            ),
            header,
            bold=True
        )

        shade_cell(
            stages_table.cell(
                0,
                i
            ),
            "D9EAF7"
        )


    stages = (
        plan.get(
            "stages",
            {}
        )
        or {}
    )


    for key in [
        "start",
        "new_knowledge",
        "practice",
        "reflection"
    ]:

        stage = (
            stages.get(
                key,
                {}
            )
            or {}
        )

        if not stage:
            continue


        cells = (
            stages_table.add_row()
            .cells
        )


        set_cell_text(
            cells[0],
            stage.get(
                "title",
                ""
            ),
            bold=True
        )

        set_cell_text(
            cells[1],
            (
                str(
                    stage.get(
                        "minutes",
                        ""
                    )
                )
                + " мин"
            )
        )

        set_cell_text(
            cells[2],
            stage.get(
                "teacher_action",
                ""
            )
        )

        set_cell_text(
            cells[3],
            stage.get(
                "student_action",
                ""
            )
        )

        set_cell_text(
            cells[4],
            stage.get(
                "assessment",
                ""
            )
        )


    # -----------------------------------------------------
    # LEVEL TASKS
    # -----------------------------------------------------

    add_heading_text(
        "Деңгейлік тапсырмалар"
    )


    for objective in (
        plan.get(
            "level_tasks",
            []
        )
        or []
    ):

        p = document.add_paragraph()

        run = p.add_run(
            str(
                objective.get(
                    "objective_code",
                    ""
                )
            )
        )

        run.bold = True
        run.font.name = (
            "Times New Roman"
        )
        run.font.size = Pt(11)


        task_table = (
            document.add_table(
                rows=1,
                cols=5
            )
        )

        task_table.style = (
            "Table Grid"
        )

        task_headers = [
            "Деңгей",
            "Тапсырма",
            "Дұрыс жауап",
            "Түсіндірме",
            "Балл"
        ]


        for i, header in enumerate(
            task_headers
        ):

            set_cell_text(
                task_table.cell(
                    0,
                    i
                ),
                header,
                bold=True
            )

            shade_cell(
                task_table.cell(
                    0,
                    i
                ),
                "E2F0D9"
            )


        for task in (
            objective.get(
                "tasks",
                []
            )
            or []
        ):

            cells = (
                task_table.add_row()
                .cells
            )

            set_cell_text(
                cells[0],
                task.get(
                    "level",
                    ""
                ),
                bold=True
            )

            set_cell_text(
                cells[1],
                task.get(
                    "task_text",
                    ""
                )
            )

            set_cell_text(
                cells[2],
                task.get(
                    "correct_answer",
                    ""
                )
            )

            set_cell_text(
                cells[3],
                task.get(
                    "explanation",
                    ""
                )
            )

            set_cell_text(
                cells[4],
                task.get(
                    "max_points",
                    ""
                )
            )


    # -----------------------------------------------------
    # CRITERIA
    # -----------------------------------------------------

    add_heading_text(
        "Бағалау критерийлері"
    )


    criteria_table = (
        document.add_table(
            rows=1,
            cols=2
        )
    )

    criteria_table.style = (
        "Table Grid"
    )


    set_cell_text(
        criteria_table.cell(
            0,
            0
        ),
        "Оқу мақсаты",
        bold=True
    )

    set_cell_text(
        criteria_table.cell(
            0,
            1
        ),
        "Бағалау критерийі",
        bold=True
    )


    for item in (
        plan.get(
            "criteria",
            []
        )
        or []
    ):

        cells = (
            criteria_table.add_row()
            .cells
        )

        set_cell_text(
            cells[0],
            item.get(
                "objective_code",
                ""
            ),
            bold=True
        )

        set_cell_text(
            cells[1],
            item.get(
                "text",
                ""
            )
        )


    # -----------------------------------------------------
    # DESCRIPTORS
    # -----------------------------------------------------

    add_heading_text(
        "Дескрипторлар"
    )


    for descriptor in (
        plan.get(
            "descriptors",
            []
        )
        or []
    ):

        p = document.add_paragraph()

        run = p.add_run(
            str(
                descriptor.get(
                    "objective_code",
                    ""
                )
            )
        )

        run.bold = True
        run.font.name = (
            "Times New Roman"
        )
        run.font.size = Pt(11)


        for item in (
            descriptor.get(
                "items",
                []
            )
            or []
        ):

            paragraph = (
                document.add_paragraph(
                    style="List Bullet"
                )
            )

            paragraph.add_run(
                str(item)
            )


    # -----------------------------------------------------
    # REFLECTION
    # -----------------------------------------------------

    add_heading_text(
        "Рефлексия"
    )


    reflection = (
        stages.get(
            "reflection",
            {}
        )
        or {}
    )


    reflection_p = (
        document.add_paragraph()
    )

    reflection_p.add_run(
        str(
            reflection.get(
                "student_action",
                ""
            )
        )
    )


    # -----------------------------------------------------
    # HOMEWORK
    # -----------------------------------------------------

    add_heading_text(
        "Үй тапсырмасы"
    )


    homework_p = (
        document.add_paragraph()
    )

    homework_p.add_run(
        str(
            plan.get(
                "homework",
                ""
            )
        )
    )


    # -----------------------------------------------------
    # SAVE TEMP
    # -----------------------------------------------------

    temp_dir = tempfile.gettempdir()

    filename = (
        safe_filename(
            plan.get(
                "lesson_topic",
                "sabak"
            )
        )
        + ".docx"
    )

    output_path = os.path.join(
        temp_dir,
        filename
    )


    document.save(
        output_path
    )


    return send_file(
        output_path,

        as_attachment=True,

        download_name=filename,

        mimetype=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )



# =========================================================
# RUN
# =========================================================


# ============================================================
# FINAL PLATFORM MODULES
# ============================================================

@app.route(
    "/settings",
    methods=["GET", "POST"],
    endpoint="final_settings"
)
def final_settings():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    if request.method == "POST":

        values = {
            "teacher_name":
                request.form.get(
                    "teacher_name",
                    ""
                ).strip()
                or "Мұғалім",

            "school_name":
                request.form.get(
                    "school_name",
                    ""
                ).strip(),

            "city_name":
                request.form.get(
                    "city_name",
                    ""
                ).strip(),

            "default_class":
                request.form.get(
                    "default_class",
                    ""
                ).strip(),

            "default_subject":
                request.form.get(
                    "default_subject",
                    "Математика"
                ).strip()
        }

        for key, value in values.items():

            conn.execute(
                """
                INSERT INTO app_settings (
                    setting_key,
                    setting_value,
                    updated_at
                )
                VALUES (?, ?, CURRENT_TIMESTAMP)

                ON CONFLICT(setting_key)
                DO UPDATE SET

                    setting_value =
                        excluded.setting_value,

                    updated_at =
                        CURRENT_TIMESTAMP
                """,
                (
                    key,
                    value
                )
            )

        conn.commit()

        session["teacher_name"] = (
            values["teacher_name"]
        )

        conn.close()

        return redirect(
            url_for(
                "final_settings",
                saved=1
            )
        )

    rows = conn.execute(
        """
        SELECT
            setting_key,
            setting_value
        FROM app_settings
        """
    ).fetchall()

    conn.close()

    current_settings = {
        row["setting_key"]:
            row["setting_value"]
        for row in rows
    }

    return render_template(
        "settings.html",

        teacher_name=session.get(
            "teacher_name",
            current_settings.get(
                "teacher_name",
                "Мұғалім"
            )
        ),

        settings=current_settings,

        saved=(
            request.args.get("saved")
            == "1"
        )
    )


@app.route(
    "/students/<int:student_id>/profile"
)
def student_profile(student_id):

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    student = conn.execute(
        """
        SELECT *
        FROM students
        WHERE id = ?
        """,
        (student_id,)
    ).fetchone()

    if student is None:

        conn.close()

        return redirect(
            url_for("students")
        )

    results = conn.execute(
        """
        SELECT
            r.*,
            lo.objective_code,
            lo.objective_text

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN learning_objectives lo
            ON lo.id =
               r.learning_objective_id

        WHERE a.student_id = ?

        ORDER BY r.created_at DESC
        """,
        (student_id,)
    ).fetchall()

    tasks = conn.execute(
        """
        SELECT *
        FROM personal_tasks

        WHERE student_id = ?

        ORDER BY id DESC
        """,
        (student_id,)
    ).fetchall()

    attempts = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM diagnostic_attempts
        WHERE student_id = ?
        """,
        (student_id,)
    ).fetchone()["total"]

    objective_count = len(results)

    completed = sum(
        1
        for task in tasks
        if task["status"] == "completed"
    )

    stats = {
        "attempts":
            attempts,

        "objectives":
            objective_count,

        "tasks":
            len(tasks),

        "completed":
            completed
    }

    conn.close()

    return render_template(
        "student_profile.html",

        student=student,
        results=results,
        tasks=tasks,
        stats=stats
    )


@app.route("/analytics")
def analytics():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    students_count = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM students
        """
    ).fetchone()["total"]

    attempts_count = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM diagnostic_attempts
        WHERE status = 'completed'
        """
    ).fetchone()["total"]

    gap_count = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM objective_results
        WHERE mastery_status = 'Олқылық бар'
        """
    ).fetchone()["total"]

    task_count = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM personal_tasks
        """
    ).fetchone()["total"]

    objectives = conn.execute(
        """
        SELECT
            lo.objective_code,
            d.subject,

            COUNT(
                DISTINCT a.student_id
            ) AS student_count,

            ROUND(
                AVG(r.percentage),
                1
            ) AS average_percentage,

            SUM(
                CASE
                    WHEN r.mastery_status =
                         'Олқылық бар'
                    THEN 1
                    ELSE 0
                END
            ) AS gap_count

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        JOIN learning_objectives lo
            ON lo.id =
               r.learning_objective_id

        WHERE a.status = 'completed'

        GROUP BY
            lo.objective_code,
            d.subject

        ORDER BY
            average_percentage ASC
        """
    ).fetchall()

    conn.close()

    return render_template(
        "analytics.html",

        stats={
            "students":
                students_count,

            "attempts":
                attempts_count,

            "gaps":
                gap_count,

            "tasks":
                task_count
        },

        objectives=objectives
    )


@app.route("/reports")
def reports():

    if not logged_in():
        return redirect(url_for("login"))

    return render_template(
        "reports.html"
    )


@app.route("/notifications")
def notifications():

    if not logged_in():
        return redirect(url_for("login"))

    conn = get_db()

    rows = conn.execute(
        """
        SELECT *
        FROM notifications
        ORDER BY created_at DESC, id DESC
        LIMIT 100
        """
    ).fetchall()

    conn.close()

    return render_template(
        "notifications.html",
        notifications=rows
    )




# ============================================================
# PUBLIC DIAGNOSTIC SHARE LINKS
# ============================================================

@app.route(
    "/diagnostics/<int:diagnostic_id>/share",
    methods=["POST"]
)
def share_diagnostic(diagnostic_id):

    if not logged_in():
        return redirect(
            url_for("login")
        )

    conn = get_db()

    diagnostic = conn.execute(
        """
        SELECT *
        FROM diagnostics
        WHERE id = ?
        """,
        (diagnostic_id,)
    ).fetchone()

    if diagnostic is None:

        conn.close()

        return redirect(
            url_for("diagnostics")
        )


    existing = conn.execute(
        """
        SELECT *
        FROM diagnostic_share_links

        WHERE
            diagnostic_id = ?
            AND is_active = 1

        ORDER BY id DESC
        LIMIT 1
        """,
        (diagnostic_id,)
    ).fetchone()


    if existing:

        token = existing["share_token"]

    else:

        token = secrets.token_urlsafe(24)

        conn.execute(
            """
            INSERT INTO diagnostic_share_links (
                diagnostic_id,
                share_token,
                is_active
            )
            VALUES (?, ?, 1)
            """,
            (
                diagnostic_id,
                token
            )
        )

        conn.commit()


    conn.close()

    return redirect(
        url_for(
            "diagnostic_detail",
            diagnostic_id=diagnostic_id,
            share=token
        )
    )


@app.route(
    "/d/<token>",
    methods=["GET"]
)
def public_diagnostic(token):

    conn = get_db()

    link = conn.execute(
        """
        SELECT *
        FROM diagnostic_share_links

        WHERE
            share_token = ?
            AND is_active = 1
        """,
        (token,)
    ).fetchone()


    if link is None:

        conn.close()

        return (
            "Сілтеме жарамсыз немесе өшірілген.",
            404
        )


    diagnostic = conn.execute(
        """
        SELECT *
        FROM diagnostics
        WHERE id = ?
        """,
        (
            link["diagnostic_id"],
        )
    ).fetchone()


    if diagnostic is None:

        conn.close()

        return (
            "Кіріс бақылау табылмады.",
            404
        )


    students = conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name

        FROM students

        WHERE class_name = ?

        ORDER BY full_name
        """,
        (
            diagnostic["class_name"],
        )
    ).fetchall()


    conn.close()


    return render_template(
        "public_diagnostic.html",

        diagnostic=diagnostic,

        students=students,

        token=token
    )


@app.route(
    "/d/<token>/start",
    methods=["POST"]
)
def public_diagnostic_start(token):

    conn = get_db()

    link = conn.execute(
        """
        SELECT *
        FROM diagnostic_share_links

        WHERE
            share_token = ?
            AND is_active = 1
        """,
        (token,)
    ).fetchone()


    if link is None:

        conn.close()

        return (
            "Сілтеме жарамсыз немесе өшірілген.",
            404
        )


    student_id = request.form.get(
        "student_id",
        type=int
    )


    if not student_id:

        conn.close()

        return redirect(
            url_for(
                "public_diagnostic",
                token=token
            )
        )


    diagnostic_id = (
        link["diagnostic_id"]
    )


    student = conn.execute(
        """
        SELECT *
        FROM students
        WHERE id = ?
        """,
        (student_id,)
    ).fetchone()


    diagnostic = conn.execute(
        """
        SELECT *
        FROM diagnostics
        WHERE id = ?
        """,
        (diagnostic_id,)
    ).fetchone()


    if (
        student is None
        or diagnostic is None
        or student["class_name"]
           != diagnostic["class_name"]
    ):

        conn.close()

        return (
            "Оқушы немесе бақылау дерегі қате.",
            400
        )


    questions = conn.execute(
        """
        SELECT *
        FROM diagnostic_questions

        WHERE diagnostic_id = ?

        ORDER BY
            question_order,
            id
        """,
        (diagnostic_id,)
    ).fetchall()


    if not questions:

        conn.close()

        return (
            "Бұл бақылауда тапсырмалар жоқ.",
            400
        )


    # --------------------------------------------------------
    # Бұрын басталған, аяқталмаған attempt бар ма?
    # --------------------------------------------------------

    existing_attempt = conn.execute(
        """
        SELECT *
        FROM diagnostic_attempts

        WHERE
            diagnostic_id = ?
            AND student_id = ?
            AND status = 'started'

        ORDER BY id DESC
        LIMIT 1
        """,
        (
            diagnostic_id,
            student_id
        )
    ).fetchone()


    if existing_attempt:

        attempt_id = (
            existing_attempt["id"]
        )

        conn.close()

        return redirect(
            url_for(
                "take_diagnostic",
                attempt_id=attempt_id
            )
        )


    max_points = sum(
        int(
            question["points"]
            or 0
        )
        for question in questions
    )


    cursor = conn.execute(
        """
        INSERT INTO diagnostic_attempts (

            diagnostic_id,
            student_id,
            status,
            total_questions,
            answered_questions,
            earned_points,
            max_points,
            percentage,
            started_at

        )
        VALUES (
            ?,
            ?,
            'started',
            ?,
            0,
            0,
            ?,
            0,
            CURRENT_TIMESTAMP
        )
        """,
        (
            diagnostic_id,
            student_id,
            len(questions),
            max_points
        )
    )


    attempt_id = (
        cursor.lastrowid
    )


    conn.commit()
    conn.close()


    return redirect(
        url_for(
            "take_diagnostic",
            attempt_id=attempt_id
        )
    )


@app.route(
    "/diagnostics/<int:diagnostic_id>/share/disable",
    methods=["POST"]
)
def disable_diagnostic_share(diagnostic_id):

    if not logged_in():
        return redirect(
            url_for("login")
        )

    conn = get_db()

    conn.execute(
        """
        UPDATE diagnostic_share_links

        SET is_active = 0

        WHERE diagnostic_id = ?
        """,
        (diagnostic_id,)
    )

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "diagnostic_detail",
            diagnostic_id=diagnostic_id
        )
    )




# ============================================================
# STUDENT PUBLIC DIAGNOSTIC PORTAL
# ============================================================

@app.route(
    "/student-test/<token>",
    methods=["GET"]
)
def student_diagnostic_portal(token):

    conn = get_db()

    portal = conn.execute(
        """
        SELECT *
        FROM student_portal_links

        WHERE
            portal_token = ?
            AND is_active = 1
        """,
        (token,)
    ).fetchone()


    if portal is None:

        conn.close()

        return (
            "Сілтеме жарамсыз немесе өшірілген.",
            404
        )


    classes = conn.execute(
        """
        SELECT DISTINCT class_name

        FROM students

        WHERE
            class_name IS NOT NULL
            AND TRIM(class_name) != ''

        ORDER BY class_name
        """
    ).fetchall()


    diagnostics = conn.execute(
        """
        SELECT
            id,
            title,
            class_name,
            subject,
            question_count,
            status

        FROM diagnostics

        WHERE question_count > 0

        ORDER BY
            subject,
            title
        """
    ).fetchall()


    subjects = sorted(
        set(
            row["subject"]
            for row in diagnostics
            if row["subject"]
        )
    )


    conn.close()


    return render_template(
        "student_diagnostic_portal.html",

        token=token,

        classes=[
            row["class_name"]
            for row in classes
        ],

        subjects=subjects,

        diagnostics=diagnostics,

        error=None
    )


@app.route(
    "/student-test/<token>/start",
    methods=["POST"]
)
def student_portal_start(token):

    conn = get_db()


    portal = conn.execute(
        """
        SELECT *
        FROM student_portal_links

        WHERE
            portal_token = ?
            AND is_active = 1
        """,
        (token,)
    ).fetchone()


    if portal is None:

        conn.close()

        return (
            "Сілтеме жарамсыз немесе өшірілген.",
            404
        )


    full_name = request.form.get(
        "full_name",
        ""
    ).strip()


    class_name = request.form.get(
        "class_name",
        ""
    ).strip()


    subject = request.form.get(
        "subject",
        ""
    ).strip()


    diagnostic_id = request.form.get(
        "diagnostic_id",
        type=int
    )


    if (
        not full_name
        or not class_name
        or not subject
        or not diagnostic_id
    ):

        conn.close()

        return redirect(
            url_for(
                "student_diagnostic_portal",
                token=token
            )
        )


    diagnostic = conn.execute(
        """
        SELECT *
        FROM diagnostics

        WHERE
            id = ?
            AND subject = ?
            AND class_name = ?
        """,
        (
            diagnostic_id,
            subject,
            class_name
        )
    ).fetchone()


    if diagnostic is None:

        conn.close()

        return (
            "Таңдалған кіріс бақылау "
            "сынып пен пәнге сәйкес емес.",
            400
        )


    # --------------------------------------------------------
    # STUDENT LOOKUP
    # --------------------------------------------------------

    student = conn.execute(
        """
        SELECT *
        FROM students

        WHERE
            LOWER(TRIM(full_name))
                = LOWER(TRIM(?))

            AND class_name = ?

        ORDER BY id
        LIMIT 1
        """,
        (
            full_name,
            class_name
        )
    ).fetchone()


    # --------------------------------------------------------
    # ЖАҢА ОҚУШЫ БОЛСА — АВТОМАТТЫ ТІРКЕУ
    # --------------------------------------------------------

    if student is None:

        # Students table-дің нақты міндетті
        # өрістерін бұзбай ең негізгі мәндерді енгіземіз.

        columns = conn.execute(
            """
            PRAGMA table_info(students)
            """
        ).fetchall()

        column_names = {
            column["name"]
            for column in columns
        }


        insert_columns = [
            "full_name",
            "class_name"
        ]

        insert_values = [
            full_name,
            class_name
        ]


        # Кейбір жобаларда subject бағаны болуы мүмкін.
        if "subject" in column_names:

            insert_columns.append(
                "subject"
            )

            insert_values.append(
                subject
            )


        placeholders = ", ".join(
            "?"
            for _ in insert_columns
        )

        sql = (
            "INSERT INTO students ("
            + ", ".join(insert_columns)
            + ") VALUES ("
            + placeholders
            + ")"
        )


        cursor = conn.execute(
            sql,
            tuple(insert_values)
        )

        student_id = (
            cursor.lastrowid
        )

        conn.commit()

    else:

        student_id = (
            student["id"]
        )


    # --------------------------------------------------------
    # QUESTIONS
    # --------------------------------------------------------

    questions = conn.execute(
        """
        SELECT *
        FROM diagnostic_questions

        WHERE diagnostic_id = ?

        ORDER BY
            question_order,
            id
        """,
        (diagnostic_id,)
    ).fetchall()


    if not questions:

        conn.close()

        return (
            "Бұл кіріс бақылауда "
            "тапсырмалар жоқ.",
            400
        )


    # --------------------------------------------------------
    # БҰРЫН БАСТАЛҒАН ATTEMPT
    # --------------------------------------------------------

    existing_attempt = conn.execute(
        """
        SELECT *
        FROM diagnostic_attempts

        WHERE
            diagnostic_id = ?
            AND student_id = ?
            AND status = 'started'

        ORDER BY id DESC
        LIMIT 1
        """,
        (
            diagnostic_id,
            student_id
        )
    ).fetchone()


    if existing_attempt:

        attempt_id = (
            existing_attempt["id"]
        )

        conn.close()

        return redirect(
            url_for(
                "take_diagnostic",
                attempt_id=attempt_id
            )
        )


    max_points = sum(
        int(
            question["points"]
            or 0
        )
        for question in questions
    )


    cursor = conn.execute(
        """
        INSERT INTO diagnostic_attempts (

            diagnostic_id,
            student_id,
            status,
            total_questions,
            answered_questions,
            earned_points,
            max_points,
            percentage,
            started_at

        )

        VALUES (
            ?,
            ?,
            'started',
            ?,
            0,
            0,
            ?,
            0,
            CURRENT_TIMESTAMP
        )
        """,
        (
            diagnostic_id,
            student_id,
            len(questions),
            max_points
        )
    )


    attempt_id = (
        cursor.lastrowid
    )


    conn.commit()
    conn.close()


    return redirect(
        url_for(
            "take_diagnostic",
            attempt_id=attempt_id
        )
    )




@app.route(
    "/student-test-link"
)
def student_portal_link():

    if not logged_in():

        return redirect(
            url_for("login")
        )


    conn = get_db()


    portal = conn.execute(
        """
        SELECT *
        FROM student_portal_links

        WHERE is_active = 1

        ORDER BY id DESC
        LIMIT 1
        """
    ).fetchone()


    if portal is None:

        import secrets

        token = (
            secrets.token_urlsafe(24)
        )

        conn.execute(
            """
            INSERT INTO student_portal_links (
                portal_token,
                is_active
            )
            VALUES (?, 1)
            """,
            (token,)
        )

        conn.commit()

    else:

        token = (
            portal["portal_token"]
        )


    conn.close()


    return redirect(
        url_for(
            "student_diagnostic_portal",
            token=token
        )
    )




# ============================================================
# CLEAR PROGRESS RESULTS
# ============================================================

@app.route(
    "/progress/clear",
    methods=["POST"]
)
def clear_progress_results():

    if not logged_in():
        return redirect(
            url_for("login")
        )

    conn = get_db()

    try:

        # ----------------------------------------------------
        # 1. Жеке тапсырмалар
        # ----------------------------------------------------

        conn.execute(
            """
            DELETE FROM personal_tasks
            """
        )


        # ----------------------------------------------------
        # 2. Оқу мақсаты бойынша нәтижелер
        # ----------------------------------------------------

        conn.execute(
            """
            DELETE FROM objective_results
            """
        )


        # ----------------------------------------------------
        # 3. Оқушы жауаптары
        # ----------------------------------------------------

        conn.execute(
            """
            DELETE FROM diagnostic_answers
            """
        )


        # ----------------------------------------------------
        # 4. Диагностика әрекеттері / нәтижелері
        # ----------------------------------------------------

        conn.execute(
            """
            DELETE FROM diagnostic_attempts
            """
        )


        # ----------------------------------------------------
        # diagnostic_results деген ескі/қосымша кесте
        # бар болса ғана тазалаймыз.
        # ----------------------------------------------------

        table_exists = conn.execute(
            """
            SELECT name
            FROM sqlite_master
            WHERE
                type = 'table'
                AND name = 'diagnostic_results'
            """
        ).fetchone()


        if table_exists:

            conn.execute(
                """
                DELETE FROM diagnostic_results
                """
            )


        conn.commit()

    except Exception:

        conn.rollback()
        conn.close()
        raise


    conn.close()


    return redirect(
        url_for(
            "progress",
            cleared=1
        )
    )




# ============================================================
# STUDENT PORTAL LINK HELPER
# ============================================================

def get_student_portal_token():

    conn = get_db()

    portal = conn.execute(
        """
        SELECT portal_token

        FROM student_portal_links

        WHERE is_active = 1

        ORDER BY id DESC
        LIMIT 1
        """
    ).fetchone()


    if portal is None:

        import secrets

        token = secrets.token_urlsafe(24)

        conn.execute(
            """
            INSERT INTO student_portal_links (
                portal_token,
                is_active
            )
            VALUES (?, 1)
            """,
            (token,)
        )

        conn.commit()

    else:

        token = portal["portal_token"]


    conn.close()

    return token


@app.context_processor
def inject_student_portal_link():

    try:

        token = (
            get_student_portal_token()
        )

        return {
            "student_portal_token":
                token
        }

    except Exception:

        return {
            "student_portal_token":
                None
        }




# ============================================================
# DARABILIM_SMART_ASSESSMENT_V3
# Smart Lesson 10-point assessment
# ============================================================

@app.route(
    "/api/smart-lessons/assessment/students"
)
def smart_lesson_assessment_students_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    requested_class = (
        request.args.get(
            "class_name",
            ""
        )
        .strip()
    )


    conn = get_db()


    rows = (
        _da_assessment_students_for_class_v1(
            conn,
            requested_class
        )
    )


    students = [
        {
            "id":
                row["id"],

            "full_name":
                row["full_name"],

            "class_name":
                row["class_name"]
        }

        for row in rows
    ]


    conn.close()


    return jsonify({

        "ok":
            True,

        "requested_class":
            requested_class,

        "grade":
            _da_assessment_grade_number_v1(
                requested_class
            ),

        "count":
            len(students),

        "students":
            students
    })


@app.route(
    "/api/smart-lessons/assessment/save",
    methods=["POST"]
)
def smart_lesson_assessment_save_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    data = (
        request.get_json(
            silent=True
        )
        or {}
    )


    try:

        student_id = int(
            data.get(
                "student_id"
            )
        )

        score = int(
            data.get(
                "score",
                0
            )
        )

    except (
        TypeError,
        ValueError
    ):

        return jsonify({
            "ok": False,
            "error": "Оқушы немесе балл қате."
        }), 400


    if score < 0 or score > 10:

        return jsonify({
            "ok": False,
            "error": "Балл 0-10 аралығында болуы керек."
        }), 400


    subject = (
        str(
            data.get(
                "subject",
                ""
            )
        )
        .strip()
    )


    lesson_topic = (
        str(
            data.get(
                "lesson_topic",
                ""
            )
        )
        .strip()
    )


    class_name = (
        str(
            data.get(
                "class_name",
                ""
            )
        )
        .strip()
    )


    teacher_comment = (
        str(
            data.get(
                "teacher_comment",
                ""
            )
        )
        .strip()
    )


    objective_ids = (
        data.get(
            "objective_ids"
        )
        or []
    )


    activity_type = (
        str(
            data.get(
                "activity_type",
                "lesson_total"
            )
        )
        .strip()
    )


    activity_title = (
        str(
            data.get(
                "activity_title",
                "Сабақтағы жалпы бағалау"
            )
        )
        .strip()
    )


    task_text = (
        str(
            data.get(
                "task_text",
                ""
            )
        )
        .strip()
    )


    task_level = (
        str(
            data.get(
                "task_level",
                ""
            )
        )
        .strip()
    )


    assessment_date = (
        str(
            data.get(
                "assessment_date",
                ""
            )
        )
        .strip()
    )


    if not assessment_date:

        assessment_date = (
            datetime.now()
            .strftime(
                "%Y-%m-%d"
            )
        )


    conn = get_db()


    student = conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name

        FROM students

        WHERE id = ?
        """,
        (
            student_id,
        )
    ).fetchone()


    if student is None:

        conn.close()

        return jsonify({
            "ok": False,
            "error": "Оқушы табылмады."
        }), 404


    # --------------------------------------------------------
    # Осы оқушы + пән + тақырып + күн бойынша
    # бар бағалауды табамыз.
    # --------------------------------------------------------

    assessment = conn.execute(
        """
        SELECT id

        FROM lesson_assessments

        WHERE
            student_id = ?
            AND COALESCE(subject, '') = ?
            AND COALESCE(lesson_topic, '') = ?
            AND COALESCE(assessment_date, '') = ?

        ORDER BY id DESC
        LIMIT 1
        """,
        (
            student_id,
            subject,
            lesson_topic,
            assessment_date
        )
    ).fetchone()


    if assessment:

        assessment_id = (
            assessment["id"]
        )


        conn.execute(
            """
            UPDATE lesson_assessments

            SET
                class_name = ?,
                total_score = ?,
                max_score = 10,
                teacher_comment = ?,
                updated_at = CURRENT_TIMESTAMP

            WHERE id = ?
            """,
            (
                class_name
                or student["class_name"],

                score,

                teacher_comment,

                assessment_id
            )
        )


    else:

        cursor = conn.execute(
            """
            INSERT INTO lesson_assessments (
                student_id,
                subject,
                lesson_topic,
                class_name,
                total_score,
                max_score,
                teacher_comment,
                assessment_date
            )

            VALUES (
                ?, ?, ?, ?, ?, 10, ?, ?
            )
            """,
            (
                student_id,
                subject,
                lesson_topic,

                class_name
                or student["class_name"],

                score,
                teacher_comment,
                assessment_date
            )
        )


        assessment_id = (
            cursor.lastrowid
        )


    # --------------------------------------------------------
    # Тапсырма / әрекет көзі
    #
    # lesson_total болса бір жалпы item ұстаймыз.
    # Кейін топтық/жұптық/жеке/бекіту/деңгейлік
    # тапсырмалар осы API арқылы бөлек сақталады.
    # --------------------------------------------------------

    objective_id = None
    objective_code = ""


    if objective_ids:

        try:

            objective_id = int(
                objective_ids[0]
            )

        except (
            TypeError,
            ValueError,
            IndexError
        ):

            objective_id = None


    if objective_id:

        objective = conn.execute(
            """
            SELECT
                id,
                objective_code

            FROM learning_objectives

            WHERE id = ?
            """,
            (
                objective_id,
            )
        ).fetchone()


        if objective:

            objective_code = (
                objective[
                    "objective_code"
                ]
                or ""
            )


    existing_item = conn.execute(
        """
        SELECT id

        FROM lesson_assessment_items

        WHERE
            assessment_id = ?
            AND activity_type = ?
            AND COALESCE(activity_title, '') = ?

        ORDER BY id DESC
        LIMIT 1
        """,
        (
            assessment_id,
            activity_type,
            activity_title
        )
    ).fetchone()


    if existing_item:

        conn.execute(
            """
            UPDATE lesson_assessment_items

            SET
                learning_objective_id = ?,
                objective_code = ?,
                task_text = ?,
                task_level = ?,
                earned_points = ?,
                max_points = 10,
                teacher_comment = ?,
                updated_at = CURRENT_TIMESTAMP

            WHERE id = ?
            """,
            (
                objective_id,
                objective_code,
                task_text,
                task_level,
                score,
                teacher_comment,
                existing_item["id"]
            )
        )


    else:

        conn.execute(
            """
            INSERT INTO lesson_assessment_items (
                assessment_id,
                learning_objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                task_level,
                earned_points,
                max_points,
                teacher_comment
            )

            VALUES (
                ?, ?, ?, ?, ?, ?, ?, ?, 10, ?
            )
            """,
            (
                assessment_id,
                objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                task_level,
                score,
                teacher_comment
            )
        )


    conn.commit()
    conn.close()


    return jsonify({
        "ok": True,

        "assessment_id":
            assessment_id,

        "student_id":
            student_id,

        "score":
            score,

        "max_score":
            10
    })


@app.route(
    "/api/smart-lessons/assessment/results"
)
def smart_lesson_assessment_results_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    subject = (
        request.args.get(
            "subject",
            ""
        )
        .strip()
    )


    lesson_topic = (
        request.args.get(
            "lesson_topic",
            ""
        )
        .strip()
    )


    conn = get_db()


    rows = conn.execute(
        """
        SELECT
            la.id,
            la.student_id,
            la.total_score,
            la.max_score,
            la.teacher_comment,
            la.assessment_date,

            s.full_name,
            s.class_name

        FROM lesson_assessments la

        JOIN students s
            ON s.id = la.student_id

        WHERE
            COALESCE(la.subject, '') = ?
            AND COALESCE(
                la.lesson_topic,
                ''
            ) = ?

        ORDER BY
            s.full_name,
            la.id DESC
        """,
        (
            subject,
            lesson_topic
        )
    ).fetchall()


    # Бір оқушыға соңғы нәтиже.
    latest = {}

    for row in rows:

        student_id = (
            row["student_id"]
        )

        if student_id in latest:
            continue


        latest[student_id] = {
            "assessment_id":
                row["id"],

            "student_id":
                student_id,

            "full_name":
                row["full_name"],

            "class_name":
                row["class_name"],

            "score":
                row["total_score"],

            "max_score":
                row["max_score"],

            "teacher_comment":
                row["teacher_comment"]
                or "",

            "assessment_date":
                row["assessment_date"]
                or ""
        }


    conn.close()


    return jsonify({
        "ok": True,
        "results":
            list(
                latest.values()
            )
    })




# ============================================================
# DARABILIM_TASK_ASSESSMENT_V1
# Task -> points -> automatic total /10
# ============================================================

@app.route(
    "/api/smart-lessons/task-assessment/save",
    methods=["POST"]
)
def smart_lesson_task_assessment_save_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    data = (
        request.get_json(
            silent=True
        )
        or {}
    )


    # --------------------------------------------------------
    # BASIC DATA
    # --------------------------------------------------------

    try:

        student_id = int(
            data.get(
                "student_id"
            )
        )

    except (
        TypeError,
        ValueError
    ):

        return jsonify({
            "ok": False,
            "error": "Оқушы қате."
        }), 400


    subject = (
        str(
            data.get(
                "subject",
                ""
            )
        )
        .strip()
    )


    lesson_topic = (
        str(
            data.get(
                "lesson_topic",
                ""
            )
        )
        .strip()
    )


    class_name = (
        str(
            data.get(
                "class_name",
                ""
            )
        )
        .strip()
    )


    lesson_key = (
        str(
            data.get(
                "lesson_key",
                ""
            )
        )
        .strip()
    )


    assessment_date = (
        str(
            data.get(
                "assessment_date",
                ""
            )
        )
        .strip()
    )


    if not assessment_date:

        assessment_date = (
            datetime.now()
            .strftime(
                "%Y-%m-%d"
            )
        )


    objective_ids = (
        data.get(
            "objective_ids"
        )
        or []
    )


    task = (
        data.get(
            "task"
        )
        or {}
    )


    task_key = (
        str(
            task.get(
                "task_key",
                ""
            )
        )
        .strip()
    )


    activity_type = (
        str(
            task.get(
                "activity_type",
                ""
            )
        )
        .strip()
    )


    activity_title = (
        str(
            task.get(
                "activity_title",
                ""
            )
        )
        .strip()
    )


    stage_group = (
        str(
            task.get(
                "stage_group",
                "middle"
            )
        )
        .strip()
    )


    stage_key = (
        str(
            task.get(
                "stage_key",
                ""
            )
        )
        .strip()
    )


    task_text = (
        str(
            task.get(
                "task_text",
                ""
            )
        )
        .strip()
    )


    task_level = (
        str(
            task.get(
                "task_level",
                ""
            )
        )
        .strip()
    )


    try:

        earned_points = int(
            task.get(
                "earned_points",
                0
            )
        )

        max_points = int(
            task.get(
                "max_points",
                0
            )
        )

    except (
        TypeError,
        ValueError
    ):

        return jsonify({
            "ok": False,
            "error": "Тапсырма баллы қате."
        }), 400


    is_counted = (
        1
        if task.get(
            "is_counted",
            True
        )
        else 0
    )


    if not subject:

        return jsonify({
            "ok": False,
            "error": "Пән таңдалмаған."
        }), 400


    if not lesson_topic:

        return jsonify({
            "ok": False,
            "error": "Сабақ тақырыбы таңдалмаған."
        }), 400


    if not task_key:

        return jsonify({
            "ok": False,
            "error": "Тапсырма кілті жоқ."
        }), 400


    if max_points < 1:

        return jsonify({
            "ok": False,
            "error": "Максималды балл кемінде 1 болуы керек."
        }), 400


    if earned_points < 0:

        return jsonify({
            "ok": False,
            "error": "Балл 0-ден төмен болмауы керек."
        }), 400


    if earned_points > max_points:

        return jsonify({
            "ok": False,
            "error": "Алынған балл максималды баллдан аспауы керек."
        }), 400


    # --------------------------------------------------------
    # DATABASE
    # --------------------------------------------------------

    conn = get_db()


    student = conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name

        FROM students

        WHERE id = ?
        """,
        (
            student_id,
        )
    ).fetchone()


    if student is None:

        conn.close()

        return jsonify({
            "ok": False,
            "error": "Оқушы табылмады."
        }), 404


    # --------------------------------------------------------
    # LESSON KEY
    # --------------------------------------------------------

    if not lesson_key:

        lesson_key = (
            subject
            + "|"
            + lesson_topic
            + "|"
            + assessment_date
        )


    import json as _json


    objective_ids_clean = []


    for value in objective_ids:

        try:

            objective_ids_clean.append(
                int(value)
            )

        except (
            TypeError,
            ValueError
        ):

            pass


    objective_json = (
        _json.dumps(
            objective_ids_clean,
            ensure_ascii=False
        )
    )


    # --------------------------------------------------------
    # FIND / CREATE ASSESSMENT
    # --------------------------------------------------------

    assessment = conn.execute(
        """
        SELECT id

        FROM lesson_assessments

        WHERE
            student_id = ?
            AND COALESCE(
                lesson_key,
                ''
            ) = ?

        ORDER BY id DESC
        LIMIT 1
        """,
        (
            student_id,
            lesson_key
        )
    ).fetchone()


    if assessment:

        assessment_id = (
            assessment["id"]
        )


        conn.execute(
            """
            UPDATE lesson_assessments

            SET
                subject = ?,
                lesson_topic = ?,
                class_name = ?,
                objective_ids_json = ?,
                calculation_mode = 'task_sum',
                assessment_date = ?,
                updated_at = CURRENT_TIMESTAMP

            WHERE id = ?
            """,
            (
                subject,
                lesson_topic,

                class_name
                or student["class_name"],

                objective_json,
                assessment_date,
                assessment_id
            )
        )


    else:

        cursor = conn.execute(
            """
            INSERT INTO lesson_assessments (
                student_id,
                subject,
                lesson_topic,
                class_name,
                total_score,
                max_score,
                assessment_date,
                lesson_key,
                objective_ids_json,
                calculation_mode
            )

            VALUES (
                ?, ?, ?, ?, 0, 10, ?, ?, ?, 'task_sum'
            )
            """,
            (
                student_id,
                subject,
                lesson_topic,

                class_name
                or student["class_name"],

                assessment_date,
                lesson_key,
                objective_json
            )
        )


        assessment_id = (
            cursor.lastrowid
        )


    # --------------------------------------------------------
    # OBJECTIVE
    # --------------------------------------------------------

    objective_id = None
    objective_code = ""


    task_objective_id = (
        task.get(
            "learning_objective_id"
        )
    )


    try:

        if task_objective_id:

            objective_id = int(
                task_objective_id
            )

    except (
        TypeError,
        ValueError
    ):

        objective_id = None


    if (
        objective_id is None
        and objective_ids_clean
    ):

        objective_id = (
            objective_ids_clean[0]
        )


    if objective_id:

        objective = conn.execute(
            """
            SELECT
                id,
                objective_code

            FROM learning_objectives

            WHERE id = ?
            """,
            (
                objective_id,
            )
        ).fetchone()


        if objective:

            objective_code = (
                objective[
                    "objective_code"
                ]
                or ""
            )


    # --------------------------------------------------------
    # UPSERT TASK RESULT
    # --------------------------------------------------------

    existing = conn.execute(
        """
        SELECT id

        FROM lesson_assessment_items

        WHERE
            assessment_id = ?
            AND COALESCE(
                task_key,
                ''
            ) = ?

        ORDER BY id DESC
        LIMIT 1
        """,
        (
            assessment_id,
            task_key
        )
    ).fetchone()


    if existing:

        conn.execute(
            """
            UPDATE lesson_assessment_items

            SET
                learning_objective_id = ?,
                objective_code = ?,
                activity_type = ?,
                activity_title = ?,
                task_text = ?,
                task_level = ?,
                earned_points = ?,
                max_points = ?,
                stage_group = ?,
                stage_key = ?,
                is_counted = ?,
                updated_at = CURRENT_TIMESTAMP

            WHERE id = ?
            """,
            (
                objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                task_level,
                earned_points,
                max_points,
                stage_group,
                stage_key,
                is_counted,
                existing["id"]
            )
        )


    else:

        conn.execute(
            """
            INSERT INTO lesson_assessment_items (
                assessment_id,
                learning_objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                task_level,
                earned_points,
                max_points,
                task_key,
                stage_group,
                stage_key,
                is_counted
            )

            VALUES (
                ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
            )
            """,
            (
                assessment_id,
                objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                task_level,
                earned_points,
                max_points,
                task_key,
                stage_group,
                stage_key,
                is_counted
            )
        )


    # --------------------------------------------------------
    # RECALCULATE
    # --------------------------------------------------------

    totals = conn.execute(
        """
        SELECT
            COALESCE(
                SUM(earned_points),
                0
            ) AS earned,

            COALESCE(
                SUM(max_points),
                0
            ) AS maximum

        FROM lesson_assessment_items

        WHERE
            assessment_id = ?
            AND COALESCE(
                is_counted,
                1
            ) = 1
        """,
        (
            assessment_id,
        )
    ).fetchone()


    raw_earned = int(
        totals["earned"]
        or 0
    )


    raw_max = int(
        totals["maximum"]
        or 0
    )


    if raw_max > 0:

        normalized_score = round(
            raw_earned
            / raw_max
            * 10
        )

    else:

        normalized_score = 0


    normalized_score = max(
        0,
        min(
            10,
            normalized_score
        )
    )


    conn.execute(
        """
        UPDATE lesson_assessments

        SET
            total_score = ?,
            max_score = 10,
            updated_at = CURRENT_TIMESTAMP

        WHERE id = ?
        """,
        (
            normalized_score,
            assessment_id
        )
    )


    conn.commit()


    # --------------------------------------------------------
    # RETURN ALL TASK RESULTS
    # --------------------------------------------------------

    items = conn.execute(
        """
        SELECT
            id,
            learning_objective_id,
            objective_code,
            activity_type,
            activity_title,
            task_text,
            task_level,
            earned_points,
            max_points,
            task_key,
            stage_group,
            stage_key,
            is_counted

        FROM lesson_assessment_items

        WHERE assessment_id = ?

        ORDER BY
            sort_order,
            id
        """,
        (
            assessment_id,
        )
    ).fetchall()


    conn.close()


    return jsonify({
        "ok": True,

        "assessment_id":
            assessment_id,

        "student_id":
            student_id,

        "raw_earned":
            raw_earned,

        "raw_max":
            raw_max,

        "score":
            normalized_score,

        "max_score":
            10,

        "items": [
            dict(row)
            for row in items
        ]
    })


@app.route(
    "/api/smart-lessons/task-assessment/results"
)
def smart_lesson_task_assessment_results_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    subject = (
        request.args.get(
            "subject",
            ""
        )
        .strip()
    )


    lesson_topic = (
        request.args.get(
            "lesson_topic",
            ""
        )
        .strip()
    )


    lesson_key = (
        request.args.get(
            "lesson_key",
            ""
        )
        .strip()
    )


    conn = get_db()


    params = []
    where = []


    if lesson_key:

        where.append(
            "COALESCE(la.lesson_key, '') = ?"
        )

        params.append(
            lesson_key
        )

    else:

        where.append(
            "COALESCE(la.subject, '') = ?"
        )

        params.append(
            subject
        )

        where.append(
            "COALESCE(la.lesson_topic, '') = ?"
        )

        params.append(
            lesson_topic
        )


    sql = """
        SELECT
            la.id,
            la.student_id,
            la.subject,
            la.lesson_topic,
            la.class_name,
            la.total_score,
            la.max_score,
            la.assessment_date,
            la.lesson_key,

            s.full_name

        FROM lesson_assessments la

        JOIN students s
            ON s.id = la.student_id

        WHERE
    """ + " AND ".join(where) + """

        ORDER BY
            s.full_name,
            la.id DESC
    """


    rows = conn.execute(
        sql,
        params
    ).fetchall()


    results = []


    for row in rows:

        items = conn.execute(
            """
            SELECT
                learning_objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                task_level,
                earned_points,
                max_points,
                task_key,
                stage_group,
                stage_key,
                is_counted

            FROM lesson_assessment_items

            WHERE assessment_id = ?

            ORDER BY
                sort_order,
                id
            """,
            (
                row["id"],
            )
        ).fetchall()


        results.append({
            "assessment_id":
                row["id"],

            "student_id":
                row["student_id"],

            "full_name":
                row["full_name"],

            "class_name":
                row["class_name"],

            "subject":
                row["subject"],

            "lesson_topic":
                row["lesson_topic"],

            "score":
                row["total_score"],

            "max_score":
                row["max_score"],

            "assessment_date":
                row["assessment_date"],

            "lesson_key":
                row["lesson_key"],

            "items": [
                dict(item)
                for item in items
            ]
        })


    conn.close()


    return jsonify({
        "ok": True,
        "results": results
    })




# ============================================================
# DARABILIM_DESCRIPTOR_ASSESSMENT_V1
# Descriptor -> task -> objective -> student
# ============================================================


@app.route(
    "/api/smart-lessons/descriptor-assessment/save",
    methods=["POST"]
)
def descriptor_assessment_save_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    data = (
        request.get_json(
            silent=True
        )
        or {}
    )


    # --------------------------------------------------------
    # STUDENT
    # --------------------------------------------------------

    try:

        student_id = int(
            data.get(
                "student_id"
            )
        )

    except (
        TypeError,
        ValueError
    ):

        return jsonify({
            "ok": False,
            "error": "Оқушы қате."
        }), 400


    subject = (
        str(
            data.get(
                "subject",
                ""
            )
        )
        .strip()
    )


    lesson_topic = (
        str(
            data.get(
                "lesson_topic",
                ""
            )
        )
        .strip()
    )


    lesson_key = (
        str(
            data.get(
                "lesson_key",
                ""
            )
        )
        .strip()
    )


    assessment_date = (
        str(
            data.get(
                "assessment_date",
                ""
            )
        )
        .strip()
    )


    if not assessment_date:

        assessment_date = (
            datetime.now()
            .strftime(
                "%Y-%m-%d"
            )
        )


    if not lesson_key:

        lesson_key = (
            subject
            + "|"
            + lesson_topic
            + "|"
            + assessment_date
        )


    task = (
        data.get(
            "task"
        )
        or {}
    )


    descriptor = (
        data.get(
            "descriptor"
        )
        or {}
    )


    task_key = (
        str(
            task.get(
                "task_key",
                ""
            )
        )
        .strip()
    )


    activity_type = (
        str(
            task.get(
                "activity_type",
                ""
            )
        )
        .strip()
    )


    activity_title = (
        str(
            task.get(
                "activity_title",
                ""
            )
        )
        .strip()
    )


    task_text = (
        str(
            task.get(
                "task_text",
                ""
            )
        )
        .strip()
    )


    descriptor_key = (
        str(
            descriptor.get(
                "descriptor_key",
                ""
            )
        )
        .strip()
    )


    descriptor_text = (
        str(
            descriptor.get(
                "descriptor_text",
                ""
            )
        )
        .strip()
    )


    completed = (
        1
        if descriptor.get(
            "is_completed",
            False
        )
        else 0
    )


    if not subject:

        return jsonify({
            "ok": False,
            "error": "Пән таңдалмаған."
        }), 400


    if not lesson_topic:

        return jsonify({
            "ok": False,
            "error": "Сабақ тақырыбы таңдалмаған."
        }), 400


    if not task_key:

        return jsonify({
            "ok": False,
            "error": "Тапсырма анықталмаған."
        }), 400


    if not descriptor_key:

        return jsonify({
            "ok": False,
            "error": "Дескриптор анықталмаған."
        }), 400


    if not descriptor_text:

        return jsonify({
            "ok": False,
            "error": "Дескриптор мәтіні жоқ."
        }), 400


    conn = get_db()


    student = conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name

        FROM students

        WHERE id = ?
        """,
        (
            student_id,
        )
    ).fetchone()


    if student is None:

        conn.close()

        return jsonify({
            "ok": False,
            "error": "Оқушы табылмады."
        }), 404


    # --------------------------------------------------------
    # ASSESSMENT
    # --------------------------------------------------------

    assessment = conn.execute(
        """
        SELECT id

        FROM lesson_assessments

        WHERE
            student_id = ?
            AND COALESCE(
                lesson_key,
                ''
            ) = ?

        ORDER BY id DESC
        LIMIT 1
        """,
        (
            student_id,
            lesson_key
        )
    ).fetchone()


    if assessment:

        assessment_id = (
            assessment["id"]
        )


    else:

        cursor = conn.execute(
            """
            INSERT INTO lesson_assessments (
                student_id,
                subject,
                lesson_topic,
                class_name,
                total_score,
                max_score,
                assessment_date,
                lesson_key,
                calculation_mode
            )

            VALUES (
                ?, ?, ?, ?, 0, 10, ?, ?, 'descriptor_sum'
            )
            """,
            (
                student_id,
                subject,
                lesson_topic,
                student["class_name"],
                assessment_date,
                lesson_key
            )
        )


        assessment_id = (
            cursor.lastrowid
        )


    # --------------------------------------------------------
    # OBJECTIVE
    # --------------------------------------------------------

    objective_id = None
    objective_code = ""


    try:

        value = task.get(
            "learning_objective_id"
        )

        if value:

            objective_id = int(
                value
            )

    except (
        TypeError,
        ValueError
    ):

        objective_id = None


    if objective_id:

        objective = conn.execute(
            """
            SELECT
                objective_code

            FROM learning_objectives

            WHERE id = ?
            """,
            (
                objective_id,
            )
        ).fetchone()


        if objective:

            objective_code = (
                objective[
                    "objective_code"
                ]
                or ""
            )


    # --------------------------------------------------------
    # ASSESSMENT ITEM
    # --------------------------------------------------------

    item = conn.execute(
        """
        SELECT id

        FROM lesson_assessment_items

        WHERE
            assessment_id = ?
            AND COALESCE(
                task_key,
                ''
            ) = ?

        ORDER BY id DESC
        LIMIT 1
        """,
        (
            assessment_id,
            task_key
        )
    ).fetchone()


    if item:

        item_id = (
            item["id"]
        )


        conn.execute(
            """
            UPDATE lesson_assessment_items

            SET
                learning_objective_id = ?,
                objective_code = ?,
                activity_type = ?,
                activity_title = ?,
                task_text = ?,
                updated_at = CURRENT_TIMESTAMP

            WHERE id = ?
            """,
            (
                objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                item_id
            )
        )


    else:

        cursor = conn.execute(
            """
            INSERT INTO lesson_assessment_items (
                assessment_id,
                learning_objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                earned_points,
                max_points,
                task_key,
                is_counted
            )

            VALUES (
                ?, ?, ?, ?, ?, ?, 0, 0, ?, 1
            )
            """,
            (
                assessment_id,
                objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                task_key
            )
        )


        item_id = (
            cursor.lastrowid
        )


    # --------------------------------------------------------
    # DESCRIPTOR UPSERT
    # --------------------------------------------------------

    existing = conn.execute(
        """
        SELECT id

        FROM lesson_assessment_descriptors

        WHERE
            assessment_id = ?
            AND student_id = ?
            AND task_key = ?
            AND descriptor_key = ?

        LIMIT 1
        """,
        (
            assessment_id,
            student_id,
            task_key,
            descriptor_key
        )
    ).fetchone()


    if existing:

        conn.execute(
            """
            UPDATE lesson_assessment_descriptors

            SET
                assessment_item_id = ?,
                learning_objective_id = ?,
                objective_code = ?,
                activity_type = ?,
                activity_title = ?,
                task_text = ?,
                descriptor_text = ?,
                earned_point = ?,
                max_point = 1,
                is_completed = ?,
                updated_at = CURRENT_TIMESTAMP

            WHERE id = ?
            """,
            (
                item_id,
                objective_id,
                objective_code,
                activity_type,
                activity_title,
                task_text,
                descriptor_text,
                completed,
                completed,
                existing["id"]
            )
        )


    else:

        conn.execute(
            """
            INSERT INTO lesson_assessment_descriptors (
                assessment_id,
                assessment_item_id,
                student_id,
                learning_objective_id,
                objective_code,
                task_key,
                activity_type,
                activity_title,
                task_text,
                descriptor_key,
                descriptor_text,
                earned_point,
                max_point,
                is_completed
            )

            VALUES (
                ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?
            )
            """,
            (
                assessment_id,
                item_id,
                student_id,
                objective_id,
                objective_code,
                task_key,
                activity_type,
                activity_title,
                task_text,
                descriptor_key,
                descriptor_text,
                completed,
                completed
            )
        )


    # --------------------------------------------------------
    # RECALCULATE EACH TASK
    # --------------------------------------------------------

    task_total = conn.execute(
        """
        SELECT
            COALESCE(
                SUM(earned_point),
                0
            ) AS earned,

            COALESCE(
                SUM(max_point),
                0
            ) AS maximum

        FROM lesson_assessment_descriptors

        WHERE
            assessment_id = ?
            AND task_key = ?
        """,
        (
            assessment_id,
            task_key
        )
    ).fetchone()


    task_earned = int(
        task_total["earned"]
        or 0
    )


    task_max = int(
        task_total["maximum"]
        or 0
    )


    conn.execute(
        """
        UPDATE lesson_assessment_items

        SET
            earned_points = ?,
            max_points = ?,
            updated_at = CURRENT_TIMESTAMP

        WHERE id = ?
        """,
        (
            task_earned,
            task_max,
            item_id
        )
    )


    # --------------------------------------------------------
    # DYNAMIC DESCRIPTOR TOTAL
    #
    # 1 completed descriptor = 1 point
    # maximum = actual descriptor count
    # --------------------------------------------------------

    lesson_earned, lesson_max = (
        _da_dynamic_descriptor_recalculate_v1(
            conn,
            assessment_id
        )
    )


    score_10 = lesson_earned


    conn.commit()


    descriptors = conn.execute(
        """
        SELECT
            id,
            learning_objective_id,
            objective_code,
            task_key,
            activity_type,
            activity_title,
            descriptor_key,
            descriptor_text,
            earned_point,
            max_point,
            is_completed

        FROM lesson_assessment_descriptors

        WHERE assessment_id = ?

        ORDER BY
            task_key,
            sort_order,
            id
        """,
        (
            assessment_id,
        )
    ).fetchall()


    conn.close()


    return jsonify({

        "ok":
            True,

        "assessment_id":
            assessment_id,

        "student_id":
            student_id,

        "task_key":
            task_key,

        "task_earned":
            task_earned,

        "task_max":
            task_max,

        "lesson_earned":
            lesson_earned,

        "lesson_max":
            lesson_max,

        "score":
            score_10,

        "max_score": lesson_max,

        "descriptors": [
            dict(row)
            for row in descriptors
        ]

    })


@app.route(
    "/api/smart-lessons/descriptor-assessment/results"
)
def descriptor_assessment_results_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    lesson_key = (
        request.args.get(
            "lesson_key",
            ""
        )
        .strip()
    )


    conn = get_db()


    assessments = conn.execute(
        """
        SELECT
            la.id,
            la.student_id,
            la.subject,
            la.lesson_topic,
            la.total_score,
            la.max_score,
            la.lesson_key,

            s.full_name,
            s.class_name

        FROM lesson_assessments la

        JOIN students s
            ON s.id = la.student_id

        WHERE
            COALESCE(
                la.lesson_key,
                ''
            ) = ?

        ORDER BY
            s.full_name,
            la.id DESC
        """,
        (
            lesson_key,
        )
    ).fetchall()


    results = []


    for assessment in assessments:

        descriptors = conn.execute(
            """
            SELECT
                learning_objective_id,
                objective_code,
                task_key,
                activity_type,
                activity_title,
                task_text,
                descriptor_key,
                descriptor_text,
                earned_point,
                max_point,
                is_completed

            FROM lesson_assessment_descriptors

            WHERE assessment_id = ?

            ORDER BY
                task_key,
                sort_order,
                id
            """,
            (
                assessment["id"],
            )
        ).fetchall()


        results.append({

            "assessment_id":
                assessment["id"],

            "student_id":
                assessment["student_id"],

            "full_name":
                assessment["full_name"],

            "class_name":
                assessment["class_name"],

            "score":
                assessment["total_score"],

            "max_score":
                assessment["max_score"],

            "descriptors": [
                dict(row)
                for row in descriptors
            ]

        })


    conn.close()


    return jsonify({
        "ok": True,
        "results": results
    })




# ============================================================
# DARABILIM_DESCRIPTOR_ANALYTICS_V1
# Descriptor evidence -> objective analytics
# ============================================================


def rebuild_descriptor_objective_progress():

    conn = get_db()


    conn.execute(
        """
        DELETE FROM descriptor_objective_progress
        """
    )


    rows = conn.execute(
        """
        SELECT
            d.student_id,
            d.learning_objective_id,
            d.objective_code,

            COUNT(*) AS total_descriptors,

            SUM(
                CASE
                    WHEN d.is_completed = 1
                    THEN 1
                    ELSE 0
                END
            ) AS completed_descriptors,

            COUNT(
                DISTINCT d.assessment_id
            ) AS evidence_count,

            MIN(d.created_at)
                AS first_evidence_at,

            MAX(d.updated_at)
                AS last_evidence_at,

            lo.subject,

            s.class_name

        FROM lesson_assessment_descriptors d

        JOIN students s
            ON s.id = d.student_id

        LEFT JOIN learning_objectives lo
            ON lo.id = d.learning_objective_id

        WHERE
            d.learning_objective_id
            IS NOT NULL

        GROUP BY
            d.student_id,
            d.learning_objective_id,
            d.objective_code,
            lo.subject,
            s.class_name
        """
    ).fetchall()


    for row in rows:

        total = int(
            row["total_descriptors"]
            or 0
        )

        completed = int(
            row["completed_descriptors"]
            or 0
        )

        evidence = int(
            row["evidence_count"]
            or 0
        )


        percentage = (
            completed
            / total
            * 100
            if total > 0
            else 0
        )


        if evidence < 2:

            status = (
                "Дерек жеткіліксіз"
            )

        elif percentage >= 80:

            status = (
                "Меңгерді"
            )

        elif percentage >= 60:

            status = (
                "Бекіту қажет"
            )

        else:

            status = (
                "Олқылық бар"
            )


        conn.execute(
            """
            INSERT INTO descriptor_objective_progress (
                student_id,
                learning_objective_id,
                objective_code,
                subject,
                class_name,
                completed_descriptors,
                total_descriptors,
                percentage,
                evidence_count,
                status,
                first_evidence_at,
                last_evidence_at
            )

            VALUES (
                ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?,
                ?, ?
            )
            """,
            (
                row["student_id"],
                row["learning_objective_id"],
                row["objective_code"],
                row["subject"],
                row["class_name"],
                completed,
                total,
                round(
                    percentage,
                    2
                ),
                evidence,
                status,
                row["first_evidence_at"],
                row["last_evidence_at"]
            )
        )


    conn.commit()
    conn.close()


@app.route(
    "/api/descriptor-progress"
)
def descriptor_progress_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    student_id = (
        request.args.get(
            "student_id",
            ""
        )
        .strip()
    )


    subject = (
        request.args.get(
            "subject",
            ""
        )
        .strip()
    )


    rebuild_descriptor_objective_progress()


    conn = get_db()


    where = []
    params = []


    if student_id:

        try:

            student_id_int = int(
                student_id
            )

        except ValueError:

            conn.close()

            return jsonify({
                "ok": False,
                "error": "student_id қате."
            }), 400


        where.append(
            "dop.student_id = ?"
        )

        params.append(
            student_id_int
        )


    if subject:

        where.append(
            "COALESCE(dop.subject, '') = ?"
        )

        params.append(
            subject
        )


    where_sql = ""

    if where:

        where_sql = (
            " WHERE "
            + " AND ".join(
                where
            )
        )


    rows = conn.execute(
        """
        SELECT
            dop.*,

            s.full_name,

            lo.objective_text,
            lo.lesson_topic,
            lo.section

        FROM descriptor_objective_progress dop

        JOIN students s
            ON s.id = dop.student_id

        LEFT JOIN learning_objectives lo
            ON lo.id = dop.learning_objective_id

        """
        + where_sql
        +
        """
        ORDER BY
            s.full_name,
            dop.objective_code
        """,
        params
    ).fetchall()


    conn.close()


    return jsonify({
        "ok": True,

        "results": [
            dict(row)
            for row in rows
        ]
    })


@app.route(
    "/api/descriptor-learning-gaps"
)
def descriptor_learning_gaps_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    rebuild_descriptor_objective_progress()


    conn = get_db()


    rows = conn.execute(
        """
        SELECT
            dop.*,

            s.full_name,

            lo.objective_text,
            lo.lesson_topic

        FROM descriptor_objective_progress dop

        JOIN students s
            ON s.id = dop.student_id

        LEFT JOIN learning_objectives lo
            ON lo.id = dop.learning_objective_id

        WHERE
            dop.evidence_count >= 2

            AND

            dop.status IN (
                'Олқылық бар',
                'Бекіту қажет'
            )

        ORDER BY
            dop.percentage ASC,
            s.full_name
        """
    ).fetchall()


    conn.close()


    return jsonify({
        "ok": True,

        "gaps": [
            dict(row)
            for row in rows
        ]
    })




# ============================================================
# DARABILIM_ANNUAL_PROGRESS_V1
# Entry -> Year evidence -> Exit
# ============================================================


def darabilim_current_academic_year():

    now = datetime.now()

    if now.month >= 8:

        return (
            f"{now.year}-"
            f"{now.year + 1}"
        )

    return (
        f"{now.year - 1}-"
        f"{now.year}"
    )


def darabilim_table_columns(
    conn,
    table
):

    return {
        row["name"]
        for row in conn.execute(
            f"PRAGMA table_info({table})"
        ).fetchall()
    }


def rebuild_annual_objective_progress(
    academic_year=None
):

    if not academic_year:

        academic_year = (
            darabilim_current_academic_year()
        )


    # --------------------------------------------------------
    # Descriptor analytics fresh
    # --------------------------------------------------------

    try:
        rebuild_descriptor_objective_progress()
    except Exception:
        pass


    conn = get_db()


    # --------------------------------------------------------
    # Derived annual table only
    # --------------------------------------------------------

    conn.execute(
        """
        DELETE FROM annual_objective_progress
        WHERE academic_year = ?
        """,
        (
            academic_year,
        )
    )


    records = {}


    # ========================================================
    # 1. YEAR-LONG SMART LESSON DESCRIPTOR EVIDENCE
    # ========================================================

    descriptor_rows = conn.execute(
        """
        SELECT
            dop.student_id,
            dop.learning_objective_id,
            dop.objective_code,
            dop.subject,
            dop.class_name,
            dop.percentage,
            dop.evidence_count,
            dop.completed_descriptors,
            dop.total_descriptors

        FROM descriptor_objective_progress dop
        """
    ).fetchall()


    for row in descriptor_rows:

        key = (
            row["student_id"],
            row["learning_objective_id"]
        )


        records[key] = {

            "student_id":
                row["student_id"],

            "learning_objective_id":
                row["learning_objective_id"],

            "objective_code":
                row["objective_code"],

            "subject":
                row["subject"],

            "class_name":
                row["class_name"],

            "entry_percentage":
                None,

            "entry_diagnostic_id":
                None,

            "entry_attempt_id":
                None,

            "year_percentage":
                row["percentage"],

            "year_evidence_count":
                row["evidence_count"] or 0,

            "year_completed":
                row["completed_descriptors"] or 0,

            "year_total":
                row["total_descriptors"] or 0,

            "exit_percentage":
                None,

            "exit_diagnostic_id":
                None,

            "exit_attempt_id":
                None,
        }


    # ========================================================
    # 2. ENTRY / INTERMEDIATE / EXIT DIAGNOSTIC RESULTS
    #
    # IMPORTANT:
    #
    # objective_results has NO student_id.
    #
    # objective_results.attempt_id
    #       ↓
    # diagnostic_attempts.id
    #       ↓
    # diagnostic_attempts.student_id
    #       ↓
    # diagnostics.diagnostic_type / academic_year
    # ========================================================

    result_rows = conn.execute(
        """
        SELECT
            r.attempt_id,
            r.learning_objective_id,
            r.percentage,

            a.student_id,
            a.diagnostic_id,
            a.status AS attempt_status,
            a.completed_at,

            d.diagnostic_type,
            d.academic_year,
            d.subject,

            s.class_name,

            lo.objective_code

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        JOIN students s
            ON s.id = a.student_id

        LEFT JOIN learning_objectives lo
            ON lo.id = r.learning_objective_id

        WHERE
            a.status = 'completed'

            AND

            d.academic_year = ?

        ORDER BY
            a.completed_at,
            a.id
        """,
        (
            academic_year,
        )
    ).fetchall()


    for row in result_rows:

        student_id = row["student_id"]

        objective_id = (
            row["learning_objective_id"]
        )


        if (
            student_id is None
            or
            objective_id is None
        ):
            continue


        key = (
            student_id,
            objective_id
        )


        if key not in records:

            records[key] = {

                "student_id":
                    student_id,

                "learning_objective_id":
                    objective_id,

                "objective_code":
                    row["objective_code"],

                "subject":
                    row["subject"],

                "class_name":
                    row["class_name"],

                "entry_percentage":
                    None,

                "entry_diagnostic_id":
                    None,

                "entry_attempt_id":
                    None,

                "year_percentage":
                    None,

                "year_evidence_count":
                    0,

                "year_completed":
                    0,

                "year_total":
                    0,

                "exit_percentage":
                    None,

                "exit_diagnostic_id":
                    None,

                "exit_attempt_id":
                    None,
            }


        diagnostic_type = (
            str(
                row["diagnostic_type"]
                or ""
            )
            .strip()
            .lower()
        )


        percentage = (
            float(row["percentage"])
            if row["percentage"] is not None
            else None
        )


        # ----------------------------------------------------
        # ENTRY
        #
        # ORDER BY completed_at means if more than one entry
        # exists, latest completed entry becomes active.
        # ----------------------------------------------------

        if diagnostic_type == "entry":

            records[key][
                "entry_percentage"
            ] = percentage

            records[key][
                "entry_diagnostic_id"
            ] = row["diagnostic_id"]

            records[key][
                "entry_attempt_id"
            ] = row["attempt_id"]


        # ----------------------------------------------------
        # INTERMEDIATE
        #
        # Intermediate is evidence, but it does NOT replace
        # descriptor year_percentage in V2.
        #
        # If no descriptor evidence exists yet, latest
        # intermediate result can serve as current percentage.
        # ----------------------------------------------------

        elif diagnostic_type == "intermediate":

            if (
                records[key][
                    "year_percentage"
                ]
                is None
            ):

                records[key][
                    "year_percentage"
                ] = percentage


        # ----------------------------------------------------
        # EXIT
        # ----------------------------------------------------

        elif diagnostic_type == "exit":

            records[key][
                "exit_percentage"
            ] = percentage

            records[key][
                "exit_diagnostic_id"
            ] = row["diagnostic_id"]

            records[key][
                "exit_attempt_id"
            ] = row["attempt_id"]


    # ========================================================
    # 3. SAVE DERIVED ANNUAL VIEW
    # ========================================================

    for record in records.values():

        entry = (
            float(
                record["entry_percentage"]
            )
            if record["entry_percentage"]
            is not None
            else None
        )


        current = (
            float(
                record["year_percentage"]
            )
            if record["year_percentage"]
            is not None
            else None
        )


        exit_value = (
            float(
                record["exit_percentage"]
            )
            if record["exit_percentage"]
            is not None
            else None
        )


        entry_to_current = (
            round(
                current - entry,
                2
            )
            if (
                entry is not None
                and
                current is not None
            )
            else None
        )


        entry_to_exit = (
            round(
                exit_value - entry,
                2
            )
            if (
                entry is not None
                and
                exit_value is not None
            )
            else None
        )


        # ----------------------------------------------------
        # STATUS
        # ----------------------------------------------------

        if (
            entry is not None
            and
            exit_value is not None
        ):

            if entry_to_exit > 0:

                status = "Өсу бар"

            elif entry_to_exit < 0:

                status = "Төмендеу"

            else:

                status = "Өзгеріс жоқ"


        elif (
            entry is not None
            and
            current is not None
        ):

            if entry_to_current > 0:

                status = (
                    "Жылдық прогресс жүріп жатыр"
                )

            elif entry_to_current < 0:

                status = (
                    "Қосымша қолдау қажет"
                )

            else:

                status = (
                    "Өзгеріс жоқ"
                )


        elif current is not None:

            status = (
                "Жыл бойғы нәтиже бар"
            )


        elif entry is not None:

            status = (
                "Кіріс бақылау бар"
            )


        elif exit_value is not None:

            status = (
                "Шығыс бақылау бар"
            )


        else:

            status = (
                "Дерек жеткіліксіз"
            )


        conn.execute(
            """
            INSERT INTO annual_objective_progress
            (
                student_id,
                learning_objective_id,
                objective_code,
                subject,
                class_name,
                academic_year,

                entry_percentage,
                entry_diagnostic_id,
                entry_attempt_id,

                year_percentage,
                year_evidence_count,
                year_completed_descriptors,
                year_total_descriptors,

                exit_percentage,
                exit_diagnostic_id,
                exit_attempt_id,

                entry_to_current_change,
                entry_to_exit_change,
                progress_status
            )

            VALUES
            (
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?,
                ?, ?, ?
            )
            """,
            (
                record["student_id"],
                record[
                    "learning_objective_id"
                ],
                record["objective_code"],
                record["subject"],
                record["class_name"],
                academic_year,

                entry,
                record[
                    "entry_diagnostic_id"
                ],
                record[
                    "entry_attempt_id"
                ],

                current,
                record[
                    "year_evidence_count"
                ],
                record[
                    "year_completed"
                ],
                record[
                    "year_total"
                ],

                exit_value,
                record[
                    "exit_diagnostic_id"
                ],
                record[
                    "exit_attempt_id"
                ],

                entry_to_current,
                entry_to_exit,
                status
            )
        )


    conn.commit()
    conn.close()


@app.route(
    "/api/annual-progress"
)
def annual_progress_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    academic_year = (
        request.args.get(
            "academic_year",
            ""
        )
        .strip()
    )


    if not academic_year:

        academic_year = (
            darabilim_current_academic_year()
        )


    rebuild_annual_objective_progress(
        academic_year
    )


    student_id = (
        request.args.get(
            "student_id",
            ""
        )
        .strip()
    )


    subject = (
        request.args.get(
            "subject",
            ""
        )
        .strip()
    )


    conn = get_db()


    where = [
        "aop.academic_year = ?"
    ]

    params = [
        academic_year
    ]


    if student_id:

        try:

            student_id_int = int(
                student_id
            )

        except ValueError:

            conn.close()

            return jsonify({
                "ok": False,
                "error": "student_id қате."
            }), 400


        where.append(
            "aop.student_id = ?"
        )

        params.append(
            student_id_int
        )


    if subject:

        where.append(
            "COALESCE(aop.subject, '') = ?"
        )

        params.append(
            subject
        )


    rows = conn.execute(
        """
        SELECT
            aop.*,

            s.full_name,

            lo.objective_text,
            lo.lesson_topic,
            lo.section

        FROM annual_objective_progress aop

        JOIN students s
            ON s.id = aop.student_id

        LEFT JOIN learning_objectives lo
            ON lo.id =
               aop.learning_objective_id

        WHERE
        """
        + " AND ".join(
            where
        )
        +
        """
        ORDER BY
            s.full_name,
            aop.objective_code
        """,
        params
    ).fetchall()


    conn.close()


    return jsonify({

        "ok":
            True,

        "academic_year":
            academic_year,

        "results": [
            dict(row)
            for row in rows
        ]

    })




# ============================================================
# DARABILIM_DIAGNOSTIC_TYPES_V1
# Entry / Intermediate / Exit controls
# ============================================================


def darabilim_diagnostic_type_info(
    diagnostic_type
):

    types = {

        "entry": {
            "code": "entry",
            "label": "Кіріс бақылау",
            "icon": "🟠",
        },

        "intermediate": {
            "code": "intermediate",
            "label": "Аралық бақылау",
            "icon": "🔵",
        },

        "exit": {
            "code": "exit",
            "label": "Шығыс бақылау",
            "icon": "🟢",
        },

    }


    return types.get(
        diagnostic_type,
        types["entry"]
    )


def darabilim_validate_diagnostic_type(
    value
):

    value = (
        str(
            value or ""
        )
        .strip()
        .lower()
    )


    if value not in {
        "entry",
        "intermediate",
        "exit",
    }:

        return None


    return value


def darabilim_validate_academic_year(
    value
):

    value = (
        str(
            value or ""
        )
        .strip()
    )


    if not re.fullmatch(
        r"\d{4}-\d{4}",
        value
    ):

        return None


    first, second = (
        value.split("-")
    )


    if int(second) != int(first) + 1:

        return None


    return value


@app.route(
    "/api/diagnostic-types"
)
def diagnostic_types_api():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    return jsonify({

        "ok": True,

        "types": [

            {
                "code": "entry",
                "label": "Кіріс бақылау",
                "icon": "🟠",
            },

            {
                "code": "intermediate",
                "label": "Аралық бақылау",
                "icon": "🔵",
            },

            {
                "code": "exit",
                "label": "Шығыс бақылау",
                "icon": "🟢",
            },

        ]

    })





# ============================================================
# DARABILIM_SMART_LESSON_STORAGE_V4
# ============================================================

import json as _da_json
import mimetypes as _da_mimetypes
import uuid as _da_uuid
from pathlib import Path as _DaPath

from flask import (
    jsonify as _da_jsonify,
    request as _da_request,
    send_file as _da_send_file,
)

from werkzeug.utils import (
    secure_filename as _da_secure_filename,
)


_DA_SMART_UPLOAD_ROOT = (
    _DaPath(__file__).resolve().parent
    / "uploads"
    / "smart_lessons"
)

_DA_SMART_UPLOAD_ROOT.mkdir(
    parents=True,
    exist_ok=True
)


_DA_SMART_MAX_FILE_SIZE = (
    250 * 1024 * 1024
)


_DA_SMART_ALLOWED_EXTENSIONS = {

    # documents
    ".pdf",
    ".doc",
    ".docx",
    ".ppt",
    ".pptx",
    ".xls",
    ".xlsx",
    ".csv",
    ".txt",
    ".rtf",

    # images
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".webp",
    ".svg",

    # video
    ".mp4",
    ".webm",
    ".mov",
    ".m4v",

    # audio
    ".mp3",
    ".wav",
    ".m4a",
    ".ogg",

    # archives / misc teacher resources
    ".zip",
}


_DA_SMART_ACTIVITY_MAP = {

    "organization": (
        "start",
        "Ұйымдастыру кезеңі"
    ),

    "psychology": (
        "start",
        "Психологиялық ахуал"
    ),

    "grouping": (
        "start",
        "Топқа бөлу"
    ),

    "explanation": (
        "middle",
        "Жаңа сабақ түсіндірмесі"
    ),

    "video": (
        "middle",
        "Бейнесабақ"
    ),

    "group": (
        "middle",
        "Топтық жұмыс"
    ),

    "pair": (
        "middle",
        "Жұптық жұмыс"
    ),

    "individual": (
        "middle",
        "Жеке жұмыс"
    ),

    "consolidation": (
        "middle",
        "Бекіту тапсырмасы"
    ),

    "feedback": (
        "end",
        "Кері байланыс"
    ),

    "homework": (
        "end",
        "Үй тапсырмасы"
    ),
}


def _da_smart_safe_text(
    value,
    max_length=10000
):

    value = str(
        value or ""
    ).strip()

    return value[
        :max_length
    ]


def _da_smart_lesson_key(
    payload
):

    explicit = _da_smart_safe_text(
        payload.get(
            "lesson_key"
        ),
        300
    )

    if explicit:
        return explicit


    subject = _da_smart_safe_text(
        payload.get(
            "subject"
        ),
        120
    )

    class_name = _da_smart_safe_text(
        payload.get(
            "class_name"
        ),
        80
    )

    topic = _da_smart_safe_text(
        payload.get(
            "topic"
        ),
        300
    )


    if not (
        subject
        or
        class_name
        or
        topic
    ):

        return ""


    return (
        subject
        + "||"
        + class_name
        + "||"
        + topic
    )


def _da_smart_validate_activity(
    activity_key
):

    activity_key = (
        _da_smart_safe_text(
            activity_key,
            80
        )
    )


    if (
        activity_key
        not in
        _DA_SMART_ACTIVITY_MAP
    ):

        return None


    stage_group, title = (
        _DA_SMART_ACTIVITY_MAP[
            activity_key
        ]
    )


    return {
        "activity_key":
            activity_key,

        "stage_group":
            stage_group,

        "activity_title":
            title,
    }


def _da_smart_material_row(
    row
):

    return {

        "id":
            row["id"],

        "lesson_key":
            row["lesson_key"],

        "stage_group":
            row["stage_group"],

        "activity_key":
            row["activity_key"],

        "original_name":
            row["original_name"],

        "stored_name":
            row["stored_name"],

        "file_size":
            row["file_size"],

        "mime_type":
            row["mime_type"],

        "download_url":
            (
                "/api/smart-lessons/"
                "materials/"
                + str(row["id"])
                + "/download"
            ),
    }


@app.route(
    "/api/smart-lessons/stage-content",
    methods=["GET"]
)
def da_smart_stage_content_get():

    lesson_key = (
        _da_smart_safe_text(
            _da_request.args.get(
                "lesson_key"
            ),
            300
        )
    )


    if not lesson_key:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "lesson_key required"
            }
        ), 400


    conn = get_db()


    rows = conn.execute(
        """
        SELECT *
        FROM smart_lesson_stage_content

        WHERE lesson_key = ?

        ORDER BY
            CASE stage_group
                WHEN 'start' THEN 1
                WHEN 'middle' THEN 2
                WHEN 'end' THEN 3
                ELSE 4
            END,
            id
        """,
        (
            lesson_key,
        )
    ).fetchall()


    materials = conn.execute(
        """
        SELECT *
        FROM smart_lesson_materials

        WHERE lesson_key = ?

        ORDER BY id
        """,
        (
            lesson_key,
        )
    ).fetchall()


    conn.close()


    content = {}


    for row in rows:

        content[
            row["activity_key"]
        ] = {

            "id":
                row["id"],

            "lesson_key":
                row["lesson_key"],

            "stage_group":
                row["stage_group"],

            "activity_key":
                row["activity_key"],

            "activity_title":
                row["activity_title"],

            "content_text":
                row["content_text"]
                or "",

            "resource_url":
                row["resource_url"]
                or "",

            "stage_minutes":
                row["stage_minutes"]
                or 0,
        }


    material_map = {}


    for row in materials:

        key = (
            row["activity_key"]
            or ""
        )

        material_map.setdefault(
            key,
            []
        ).append(
            _da_smart_material_row(
                row
            )
        )


    return _da_jsonify(
        {
            "ok": True,
            "lesson_key":
                lesson_key,
            "content":
                content,
            "materials":
                material_map,
        }
    )


@app.route(
    "/api/smart-lessons/stage-content/save",
    methods=["POST"]
)
def da_smart_stage_content_save():

    payload = (
        _da_request.get_json(
            silent=True
        )
        or {}
    )


    lesson_key = (
        _da_smart_lesson_key(
            payload
        )
    )


    if not lesson_key:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "Lesson context required"
            }
        ), 400


    activity = (
        _da_smart_validate_activity(
            payload.get(
                "activity_key"
            )
        )
    )


    if activity is None:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "Invalid activity"
            }
        ), 400


    subject = _da_smart_safe_text(
        payload.get(
            "subject"
        ),
        120
    )

    class_name = _da_smart_safe_text(
        payload.get(
            "class_name"
        ),
        80
    )

    topic = _da_smart_safe_text(
        payload.get(
            "topic"
        ),
        300
    )

    content_text = (
        _da_smart_safe_text(
            payload.get(
                "content_text"
            ),
            30000
        )
    )

    resource_url = (
        _da_smart_safe_text(
            payload.get(
                "resource_url"
            ),
            2000
        )
    )


    try:

        stage_minutes = int(
            payload.get(
                "stage_minutes"
            )
            or 0
        )

    except Exception:

        stage_minutes = 0


    stage_minutes = max(
        0,
        min(
            stage_minutes,
            180
        )
    )


    conn = get_db()


    existing = conn.execute(
        """
        SELECT id
        FROM smart_lesson_stage_content

        WHERE
            lesson_key = ?
            AND stage_group = ?
            AND activity_key = ?

        LIMIT 1
        """,
        (
            lesson_key,
            activity[
                "stage_group"
            ],
            activity[
                "activity_key"
            ],
        )
    ).fetchone()


    if existing:

        conn.execute(
            """
            UPDATE
                smart_lesson_stage_content

            SET
                subject = ?,
                class_name = ?,
                topic = ?,
                activity_title = ?,
                content_text = ?,
                resource_url = ?,
                stage_minutes = ?,
                updated_at =
                    CURRENT_TIMESTAMP

            WHERE id = ?
            """,
            (
                subject,
                class_name,
                topic,

                activity[
                    "activity_title"
                ],

                content_text,
                resource_url,
                stage_minutes,

                existing["id"],
            )
        )


        row_id = (
            existing["id"]
        )


    else:

        cursor = conn.execute(
            """
            INSERT INTO
                smart_lesson_stage_content
            (
                lesson_key,
                subject,
                class_name,
                topic,

                stage_group,
                activity_key,
                activity_title,

                content_text,
                resource_url,

                stage_minutes
            )

            VALUES (
                ?, ?, ?, ?,
                ?, ?, ?,
                ?, ?,
                ?
            )
            """,
            (
                lesson_key,
                subject,
                class_name,
                topic,

                activity[
                    "stage_group"
                ],

                activity[
                    "activity_key"
                ],

                activity[
                    "activity_title"
                ],

                content_text,
                resource_url,
                stage_minutes,
            )
        )


        row_id = (
            cursor.lastrowid
        )


    conn.commit()
    conn.close()


    return _da_jsonify(
        {
            "ok": True,
            "id":
                row_id,
            "lesson_key":
                lesson_key,
            "activity_key":
                activity[
                    "activity_key"
                ],
        }
    )


@app.route(
    "/api/smart-lessons/materials/upload",
    methods=["POST"]
)
def da_smart_material_upload():

    lesson_key = (
        _da_smart_safe_text(
            _da_request.form.get(
                "lesson_key"
            ),
            300
        )
    )


    if not lesson_key:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "lesson_key required"
            }
        ), 400


    activity = (
        _da_smart_validate_activity(
            _da_request.form.get(
                "activity_key"
            )
        )
    )


    if activity is None:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "Invalid activity"
            }
        ), 400


    files = (
        _da_request.files.getlist(
            "files"
        )
    )


    files = [
        item
        for item in files
        if (
            item
            and
            item.filename
        )
    ]


    if not files:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "No files"
            }
        ), 400


    lesson_folder_name = (
        _da_uuid.uuid5(
            _da_uuid.NAMESPACE_URL,
            lesson_key
        ).hex
    )


    lesson_folder = (
        _DA_SMART_UPLOAD_ROOT
        / lesson_folder_name
        / activity[
            "activity_key"
        ]
    )


    lesson_folder.mkdir(
        parents=True,
        exist_ok=True
    )


    conn = get_db()

    saved = []


    try:

        for file in files:

            original_name = (
                _da_secure_filename(
                    file.filename
                )
            )


            if not original_name:

                continue


            extension = (
                _DaPath(
                    original_name
                ).suffix.lower()
            )


            if (
                extension
                not in
                _DA_SMART_ALLOWED_EXTENSIONS
            ):

                continue


            file.stream.seek(
                0,
                os.SEEK_END
            )

            size = (
                file.stream.tell()
            )

            file.stream.seek(
                0
            )


            if (
                size
                >
                _DA_SMART_MAX_FILE_SIZE
            ):

                continue


            stored_name = (
                _da_uuid.uuid4().hex
                + extension
            )


            target = (
                lesson_folder
                / stored_name
            )


            file.save(
                str(target)
            )


            mime_type = (
                file.mimetype
                or
                _da_mimetypes.guess_type(
                    original_name
                )[0]
                or
                "application/octet-stream"
            )


            cursor = conn.execute(
                """
                INSERT INTO
                    smart_lesson_materials
                (
                    lesson_key,
                    stage_group,
                    activity_key,

                    original_name,
                    stored_name,
                    file_path,
                    file_size,
                    mime_type,

                    created_at
                )

                VALUES (
                    ?, ?, ?,
                    ?, ?, ?, ?, ?,
                    CURRENT_TIMESTAMP
                )
                """,
                (
                    lesson_key,

                    activity[
                        "stage_group"
                    ],

                    activity[
                        "activity_key"
                    ],

                    original_name,
                    stored_name,
                    str(target),
                    size,
                    mime_type,
                )
            )


            saved.append(
                {
                    "id":
                        cursor.lastrowid,

                    "original_name":
                        original_name,

                    "file_size":
                        size,

                    "mime_type":
                        mime_type,

                    "download_url":
                        (
                            "/api/smart-lessons/"
                            "materials/"
                            + str(
                                cursor.lastrowid
                            )
                            + "/download"
                        ),
                }
            )


        conn.commit()


    except Exception:

        conn.rollback()
        conn.close()

        raise


    conn.close()


    if not saved:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "No supported files saved"
            }
        ), 400


    return _da_jsonify(
        {
            "ok": True,
            "lesson_key":
                lesson_key,
            "activity_key":
                activity[
                    "activity_key"
                ],
            "files":
                saved,
        }
    )


@app.route(
    "/api/smart-lessons/materials/<int:material_id>/download",
    methods=["GET"]
)
def da_smart_material_download(
    material_id
):

    conn = get_db()


    row = conn.execute(
        """
        SELECT *
        FROM smart_lesson_materials

        WHERE id = ?

        LIMIT 1
        """,
        (
            material_id,
        )
    ).fetchone()


    conn.close()


    if row is None:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "Material not found"
            }
        ), 404


    path = _DaPath(
        row["file_path"]
        or ""
    )


    try:

        resolved = (
            path.resolve()
        )

        upload_root = (
            _DA_SMART_UPLOAD_ROOT
            .resolve()
        )


        if (
            upload_root
            not in resolved.parents
        ):

            return _da_jsonify(
                {
                    "ok": False,
                    "error":
                        "Invalid file path"
                }
            ), 403


    except Exception:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "Invalid file path"
            }
        ), 403


    if not resolved.exists():

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "File missing"
            }
        ), 404


    return _da_send_file(
        str(resolved),

        as_attachment=True,

        download_name=(
            row["original_name"]
            or
            resolved.name
        ),

        mimetype=(
            row["mime_type"]
            or None
        ),
    )


@app.route(
    "/api/smart-lessons/materials/<int:material_id>/delete",
    methods=["POST"]
)
def da_smart_material_delete(
    material_id
):

    conn = get_db()


    row = conn.execute(
        """
        SELECT *
        FROM smart_lesson_materials

        WHERE id = ?

        LIMIT 1
        """,
        (
            material_id,
        )
    ).fetchone()


    if row is None:

        conn.close()

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "Material not found"
            }
        ), 404


    path = _DaPath(
        row["file_path"]
        or ""
    )


    conn.execute(
        """
        DELETE FROM
            smart_lesson_materials

        WHERE id = ?
        """,
        (
            material_id,
        )
    )


    conn.commit()
    conn.close()


    try:

        resolved = path.resolve()

        upload_root = (
            _DA_SMART_UPLOAD_ROOT
            .resolve()
        )


        if (
            upload_root
            in resolved.parents
            and
            resolved.exists()
        ):

            resolved.unlink()

    except Exception:

        pass


    return _da_jsonify(
        {
            "ok": True,
            "deleted_id":
                material_id,
        }
    )





# ============================================================
# DARABILIM_TASK_DESCRIPTOR_AUTHORING_FINAL_V1
# ============================================================

from flask import request as _td_request
from flask import jsonify as _td_jsonify


_TD_ALLOWED_ACTIVITIES = {
    "group": ("middle", "Топтық жұмыс"),
    "pair": ("middle", "Жұптық жұмыс"),
    "individual": ("middle", "Жеке жұмыс"),
    "consolidation": ("middle", "Бекіту тапсырмасы"),
    "homework": ("end", "Үй тапсырмасы"),
}


def _td_text(value, limit=30000):
    return str(value or "").strip()[:limit]


@app.route(
    "/api/smart-lessons/teacher-task",
    methods=["GET"]
)
def td_teacher_task_get():

    lesson_key = _td_text(
        _td_request.args.get("lesson_key"),
        300
    )

    activity_key = _td_text(
        _td_request.args.get("activity_key"),
        80
    )

    if not lesson_key or activity_key not in _TD_ALLOWED_ACTIVITIES:
        return _td_jsonify({
            "ok": False,
            "error": "Invalid lesson/activity"
        }), 400

    conn = get_db()

    task = conn.execute(
        """
        SELECT *
        FROM smart_lesson_teacher_tasks
        WHERE lesson_key = ?
          AND activity_key = ?
        LIMIT 1
        """,
        (lesson_key, activity_key)
    ).fetchone()

    if task is None:
        conn.close()

        return _td_jsonify({
            "ok": True,
            "task": None,
            "descriptors": []
        })

    descriptors = conn.execute(
        """
        SELECT *
        FROM smart_lesson_teacher_descriptors
        WHERE teacher_task_id = ?
        ORDER BY descriptor_order, id
        """,
        (task["id"],)
    ).fetchall()

    conn.close()

    return _td_jsonify({
        "ok": True,

        "task": {
            "id": task["id"],
            "lesson_key": task["lesson_key"],
            "learning_objective_id":
                task["learning_objective_id"],
            "objective_code":
                task["objective_code"],
            "activity_key":
                task["activity_key"],
            "activity_title":
                task["activity_title"],
            "task_text":
                task["task_text"] or "",
        },

        "descriptors": [
            {
                "id": row["id"],
                "descriptor_order":
                    row["descriptor_order"],
                "descriptor_text":
                    row["descriptor_text"],
                "max_point":
                    row["max_point"],
            }
            for row in descriptors
        ]
    })


@app.route(
    "/api/smart-lessons/teacher-task/save",
    methods=["POST"]
)
def td_teacher_task_save():

    payload = _td_request.get_json(
        silent=True
    ) or {}

    lesson_key = _td_text(
        payload.get("lesson_key"),
        300
    )

    activity_key = _td_text(
        payload.get("activity_key"),
        80
    )

    if (
        not lesson_key
        or
        activity_key not in _TD_ALLOWED_ACTIVITIES
    ):
        return _td_jsonify({
            "ok": False,
            "error": "Invalid lesson/activity"
        }), 400

    stage_group, activity_title = (
        _TD_ALLOWED_ACTIVITIES[
            activity_key
        ]
    )

    task_text = _td_text(
        payload.get("task_text"),
        30000
    )

    subject = _td_text(
        payload.get("subject"),
        120
    )

    class_name = _td_text(
        payload.get("class_name"),
        80
    )

    topic = _td_text(
        payload.get("topic"),
        300
    )

    objective_code = _td_text(
        payload.get("objective_code"),
        100
    )

    objective_id = payload.get(
        "learning_objective_id"
    )

    try:
        objective_id = (
            int(objective_id)
            if objective_id not in (
                None,
                ""
            )
            else None
        )
    except Exception:
        objective_id = None

    descriptors = payload.get(
        "descriptors"
    ) or []

    cleaned_descriptors = []

    for item in descriptors:

        if isinstance(item, dict):
            text = _td_text(
                item.get("descriptor_text"),
                3000
            )
        else:
            text = _td_text(
                item,
                3000
            )

        if text:
            cleaned_descriptors.append(
                text
            )

    conn = get_db()

    existing = conn.execute(
        """
        SELECT id
        FROM smart_lesson_teacher_tasks
        WHERE lesson_key = ?
          AND activity_key = ?
        LIMIT 1
        """,
        (
            lesson_key,
            activity_key
        )
    ).fetchone()

    if existing:

        task_id = existing["id"]

        conn.execute(
            """
            UPDATE smart_lesson_teacher_tasks
            SET
                subject = ?,
                class_name = ?,
                topic = ?,
                learning_objective_id = ?,
                objective_code = ?,
                stage_group = ?,
                activity_title = ?,
                task_text = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (
                subject,
                class_name,
                topic,
                objective_id,
                objective_code,
                stage_group,
                activity_title,
                task_text,
                task_id,
            )
        )

    else:

        cursor = conn.execute(
            """
            INSERT INTO smart_lesson_teacher_tasks
            (
                lesson_key,
                subject,
                class_name,
                topic,
                learning_objective_id,
                objective_code,
                stage_group,
                activity_key,
                activity_title,
                task_text
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                lesson_key,
                subject,
                class_name,
                topic,
                objective_id,
                objective_code,
                stage_group,
                activity_key,
                activity_title,
                task_text,
            )
        )

        task_id = cursor.lastrowid

    # Мұғалімнің қазіргі дескрипторлары —
    # осы тапсырманың canonical тізімі.
    conn.execute(
        """
        DELETE FROM smart_lesson_teacher_descriptors
        WHERE teacher_task_id = ?
        """,
        (task_id,)
    )

    for index, text in enumerate(
        cleaned_descriptors,
        start=1
    ):

        conn.execute(
            """
            INSERT INTO
            smart_lesson_teacher_descriptors
            (
                teacher_task_id,
                lesson_key,
                learning_objective_id,
                objective_code,
                activity_key,
                descriptor_order,
                descriptor_text,
                max_point
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, 1)
            """,
            (
                task_id,
                lesson_key,
                objective_id,
                objective_code,
                activity_key,
                index,
                text,
            )
        )

    conn.commit()
    conn.close()

    return _td_jsonify({
        "ok": True,
        "task_id": task_id,
        "descriptor_count":
            len(cleaned_descriptors)
    })


@app.route(
    "/api/smart-lessons/teacher-tasks",
    methods=["GET"]
)
def td_teacher_tasks_all():

    lesson_key = _td_text(
        _td_request.args.get(
            "lesson_key"
        ),
        300
    )

    if not lesson_key:
        return _td_jsonify({
            "ok": False,
            "error": "lesson_key required"
        }), 400

    conn = get_db()

    tasks = conn.execute(
        """
        SELECT *
        FROM smart_lesson_teacher_tasks
        WHERE lesson_key = ?
        ORDER BY
            CASE activity_key
                WHEN 'group' THEN 1
                WHEN 'pair' THEN 2
                WHEN 'individual' THEN 3
                WHEN 'consolidation' THEN 4
                WHEN 'homework' THEN 5
                ELSE 6
            END
        """,
        (lesson_key,)
    ).fetchall()

    results = []

    for task in tasks:

        descriptors = conn.execute(
            """
            SELECT *
            FROM smart_lesson_teacher_descriptors
            WHERE teacher_task_id = ?
            ORDER BY descriptor_order, id
            """,
            (task["id"],)
        ).fetchall()

        results.append({
            "id":
                task["id"],

            "lesson_key":
                task["lesson_key"],

            "activity_key":
                task["activity_key"],

            "activity_title":
                task["activity_title"],

            "task_text":
                task["task_text"] or "",

            "learning_objective_id":
                task["learning_objective_id"],

            "objective_code":
                task["objective_code"],

            "descriptors": [
                {
                    "id":
                        row["id"],

                    "descriptor_order":
                        row["descriptor_order"],

                    "descriptor_text":
                        row["descriptor_text"],

                    "max_point":
                        row["max_point"],
                }
                for row in descriptors
            ]
        })

    conn.close()

    return _td_jsonify({
        "ok": True,
        "lesson_key": lesson_key,
        "tasks": results
    })






# ============================================================
# DARABILIM_SMART_MATERIAL_PREVIEW_FINAL_V1
# ============================================================

@app.route(
    "/api/smart-lessons/materials/<int:material_id>/preview",
    methods=["GET"]
)
def da_smart_material_preview_final_v1(
    material_id
):

    conn = get_db()

    row = conn.execute(
        """
        SELECT *
        FROM smart_lesson_materials
        WHERE id = ?
        LIMIT 1
        """,
        (
            material_id,
        )
    ).fetchone()

    conn.close()


    if row is None:

        return _da_jsonify(
            {
                "ok": False,
                "error": "Material not found"
            }
        ), 404


    path = _DaPath(
        row["file_path"]
        or ""
    )


    try:

        resolved = path.resolve()

        upload_root = (
            _DA_SMART_UPLOAD_ROOT
            .resolve()
        )

        if (
            upload_root
            not in resolved.parents
        ):

            return _da_jsonify(
                {
                    "ok": False,
                    "error": "Invalid file path"
                }
            ), 403

    except Exception:

        return _da_jsonify(
            {
                "ok": False,
                "error": "Invalid file path"
            }
        ), 403


    if not resolved.exists():

        return _da_jsonify(
            {
                "ok": False,
                "error": "File missing"
            }
        ), 404


    return _da_send_file(
        str(resolved),

        # Preview үшін attachment ЕМЕС.
        as_attachment=False,

        download_name=(
            row["original_name"]
            or resolved.name
        ),

        mimetype=(
            row["mime_type"]
            or
            _da_mimetypes.guess_type(
                str(resolved)
            )[0]
            or
            "application/octet-stream"
        ),
    )




# =====================================================================
# DARABILIM_POWERPOINT_PDF_PREVIEW_FINAL_V1
# =====================================================================

import subprocess as _da_ppt_subprocess
import shutil as _da_ppt_shutil
import tempfile as _da_ppt_tempfile
import threading as _da_ppt_threading
from pathlib import Path as _DaPptPath


_DA_PPT_CONVERT_LOCK = _da_ppt_threading.Lock()


def _da_ppt_office_binary_v1():

    return (
        _da_ppt_shutil.which(
            "libreoffice"
        )
        or
        _da_ppt_shutil.which(
            "soffice"
        )
    )


def _da_ppt_safe_material_v1(
    material_id
):

    conn = get_db()

    row = conn.execute(
        """
        SELECT *
        FROM smart_lesson_materials
        WHERE id = ?
        LIMIT 1
        """,
        (
            material_id,
        )
    ).fetchone()

    conn.close()


    if row is None:

        return None, None, (
            _da_jsonify(
                {
                    "ok": False,
                    "error":
                        "Material not found"
                }
            ),
            404
        )


    source = _DaPptPath(
        row["file_path"]
        or ""
    )


    try:

        resolved = source.resolve()

        upload_root = (
            _DA_SMART_UPLOAD_ROOT
            .resolve()
        )


        if (
            upload_root
            not in resolved.parents
        ):

            return None, None, (
                _da_jsonify(
                    {
                        "ok": False,
                        "error":
                            "Invalid file path"
                    }
                ),
                403
            )


    except Exception:

        return None, None, (
            _da_jsonify(
                {
                    "ok": False,
                    "error":
                        "Invalid file path"
                }
            ),
            403
        )


    if not resolved.exists():

        return None, None, (
            _da_jsonify(
                {
                    "ok": False,
                    "error":
                        "File missing"
                }
            ),
            404
        )


    if (
        resolved.suffix.lower()
        not in {
            ".ppt",
            ".pptx"
        }
    ):

        return None, None, (
            _da_jsonify(
                {
                    "ok": False,
                    "error":
                        "Not a PowerPoint file"
                }
            ),
            400
        )


    return row, resolved, None


def _da_ppt_pdf_path_v1(
    source
):

    preview_dir = (
        source.parent
        / "_preview"
    )

    preview_dir.mkdir(
        parents=True,
        exist_ok=True
    )


    return (
        preview_dir
        /
        (
            source.stem
            + ".pdf"
        )
    )


def _da_ppt_convert_pdf_v1(
    source,
    target_pdf
):

    office = (
        _da_ppt_office_binary_v1()
    )


    if not office:

        raise RuntimeError(
            "LibreOffice is not installed"
        )


    target_pdf.parent.mkdir(
        parents=True,
        exist_ok=True
    )


    # Егер PDF бар және PPTX-тен жаңа болса,
    # қайта конвертация қажет емес.

    if (
        target_pdf.exists()
        and
        target_pdf.stat().st_size > 0
        and
        target_pdf.stat().st_mtime
        >=
        source.stat().st_mtime
    ):

        return target_pdf


    with _DA_PPT_CONVERT_LOCK:

        # Lock күткен кезде басқа request
        # PDF жасап қойған болуы мүмкін.

        if (
            target_pdf.exists()
            and
            target_pdf.stat().st_size > 0
            and
            target_pdf.stat().st_mtime
            >=
            source.stat().st_mtime
        ):

            return target_pdf


        with _da_ppt_tempfile.TemporaryDirectory(
            prefix="darabilim_ppt_"
        ) as temp_dir:

            temp_path = _DaPptPath(
                temp_dir
            )


            command = [
                office,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temp_path),
                str(source),
            ]


            result = (
                _da_ppt_subprocess.run(
                    command,
                    stdout=
                        _da_ppt_subprocess.PIPE,
                    stderr=
                        _da_ppt_subprocess.PIPE,
                    text=True,
                    timeout=180
                )
            )


            generated = (
                temp_path
                /
                (
                    source.stem
                    + ".pdf"
                )
            )


            if (
                result.returncode != 0
                or
                not generated.exists()
                or
                generated.stat().st_size <= 0
            ):

                raise RuntimeError(
                    "PowerPoint conversion failed. "
                    + (
                        result.stderr.strip()
                        or
                        result.stdout.strip()
                        or
                        "Unknown LibreOffice error"
                    )
                )


            temp_target = (
                target_pdf.parent
                /
                (
                    target_pdf.name
                    + ".tmp"
                )
            )


            _da_ppt_shutil.copy2(
                generated,
                temp_target
            )


            temp_target.replace(
                target_pdf
            )


    return target_pdf


@app.route(
    "/api/smart-lessons/materials/"
    "<int:material_id>/powerpoint-preview",
    methods=["GET"]
)
def da_smart_powerpoint_preview_final_v1(
    material_id
):

    row, source, error = (
        _da_ppt_safe_material_v1(
            material_id
        )
    )


    if error is not None:

        return error


    target_pdf = (
        _da_ppt_pdf_path_v1(
            source
        )
    )


    try:

        pdf_path = (
            _da_ppt_convert_pdf_v1(
                source,
                target_pdf
            )
        )


    except _da_ppt_subprocess.TimeoutExpired:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "PowerPoint conversion timeout"
            }
        ), 504


    except Exception as exc:

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "PowerPoint preview error",
                "detail":
                    str(exc)
            }
        ), 500


    return _da_send_file(
        str(pdf_path),

        as_attachment=False,

        download_name=(
            (
                row["original_name"]
                or source.name
            )
            + ".pdf"
        ),

        mimetype=
            "application/pdf",

        conditional=True
    )





# =====================================================================
# DARABILIM_MATERIAL_DELETE_FINAL_V2
# =====================================================================

@app.route(
    "/api/smart-lessons/materials/"
    "<int:material_id>/delete-final",
    methods=["POST"]
)
def da_smart_material_delete_final_v2(
    material_id
):

    conn = get_db()

    row = conn.execute(
        """
        SELECT *
        FROM smart_lesson_materials
        WHERE id = ?
        LIMIT 1
        """,
        (
            material_id,
        )
    ).fetchone()


    if row is None:

        conn.close()

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "Material not found"
            }
        ), 404


    source = _DaPath(
        row["file_path"]
        or ""
    )


    # ---------------------------------------------------------
    # PATH SAFETY
    # ---------------------------------------------------------

    try:

        resolved = source.resolve()

        upload_root = (
            _DA_SMART_UPLOAD_ROOT
            .resolve()
        )


        if (
            upload_root
            not in resolved.parents
        ):

            conn.close()

            return _da_jsonify(
                {
                    "ok": False,
                    "error":
                        "Invalid file path"
                }
            ), 403


    except Exception:

        conn.close()

        return _da_jsonify(
            {
                "ok": False,
                "error":
                    "Invalid file path"
            }
        ), 403


    original_name = (
        row["original_name"]
        or resolved.name
    )


    lesson_key = (
        row["lesson_key"]
        or ""
    )


    activity_key = (
        row["activity_key"]
        or ""
    )


    deleted_original = False
    deleted_preview = False


    # ---------------------------------------------------------
    # DELETE DATABASE ROW FIRST INSIDE TRANSACTION
    # ---------------------------------------------------------

    try:

        conn.execute(
            """
            DELETE FROM smart_lesson_materials
            WHERE id = ?
            """,
            (
                material_id,
            )
        )

        conn.commit()


    except Exception:

        conn.rollback()
        conn.close()

        raise


    conn.close()


    # ---------------------------------------------------------
    # DELETE ORIGINAL FILE
    # ---------------------------------------------------------

    try:

        if (
            resolved.exists()
            and
            resolved.is_file()
        ):

            resolved.unlink()

            deleted_original = True

    except Exception as exc:

        # DB row already removed.
        # Return warning instead of restoring stale DB record.
        return _da_jsonify(
            {
                "ok": True,
                "warning":
                    "Database row deleted, "
                    "but physical file could not be removed.",

                "detail":
                    str(exc),

                "id":
                    material_id,

                "original_name":
                    original_name,

                "lesson_key":
                    lesson_key,

                "activity_key":
                    activity_key,
            }
        )


    # ---------------------------------------------------------
    # DELETE GENERATED POWERPOINT PDF PREVIEW
    # ---------------------------------------------------------

    try:

        preview_pdf = (
            resolved.parent
            / "_preview"
            / (
                resolved.stem
                + ".pdf"
            )
        )


        if (
            preview_pdf.exists()
            and
            preview_pdf.is_file()
        ):

            preview_pdf.unlink()

            deleted_preview = True


        preview_dir = (
            resolved.parent
            / "_preview"
        )


        if (
            preview_dir.exists()
            and
            preview_dir.is_dir()
            and
            not any(
                preview_dir.iterdir()
            )
        ):

            preview_dir.rmdir()


    except Exception:

        pass


    # ---------------------------------------------------------
    # CLEAN EMPTY ACTIVITY DIRECTORY
    # ---------------------------------------------------------

    try:

        parent = resolved.parent

        if (
            parent.exists()
            and
            parent.is_dir()
            and
            parent != upload_root
            and
            not any(
                parent.iterdir()
            )
        ):

            parent.rmdir()

    except Exception:

        pass


    return _da_jsonify(
        {
            "ok": True,

            "id":
                material_id,

            "original_name":
                original_name,

            "lesson_key":
                lesson_key,

            "activity_key":
                activity_key,

            "deleted_original":
                deleted_original,

            "deleted_preview":
                deleted_preview,
        }
    )





# =====================================================================
# DARABILIM_DYNAMIC_DESCRIPTOR_SCORE_FINAL_V1
#
# 1 descriptor = 1 point
#
# lesson_assessments.total_score =
#     completed descriptor count
#
# lesson_assessments.max_score =
#     actual descriptor count
# =====================================================================


def _da_dynamic_descriptor_recalculate_v1(
    conn,
    assessment_id
):

    # ---------------------------------------------------------
    # EACH TASK
    # ---------------------------------------------------------

    task_rows = conn.execute(
        """
        SELECT DISTINCT
            assessment_item_id,
            task_key

        FROM lesson_assessment_descriptors

        WHERE assessment_id = ?
        """,
        (
            assessment_id,
        )
    ).fetchall()


    for task_row in task_rows:

        task_key = (
            task_row["task_key"]
            or ""
        )


        item_id = (
            task_row["assessment_item_id"]
        )


        totals = conn.execute(
            """
            SELECT
                COALESCE(
                    SUM(
                        CASE
                            WHEN COALESCE(is_completed, 0) = 1
                            THEN 1
                            ELSE 0
                        END
                    ),
                    0
                ) AS earned,

                COUNT(*) AS maximum

            FROM lesson_assessment_descriptors

            WHERE
                assessment_id = ?
                AND task_key = ?
            """,
            (
                assessment_id,
                task_key
            )
        ).fetchone()


        earned = int(
            totals["earned"]
            or 0
        )


        maximum = int(
            totals["maximum"]
            or 0
        )


        if item_id:

            conn.execute(
                """
                UPDATE lesson_assessment_items

                SET
                    earned_points = ?,
                    max_points = ?,
                    updated_at = CURRENT_TIMESTAMP

                WHERE id = ?
                """,
                (
                    earned,
                    maximum,
                    item_id
                )
            )


    # ---------------------------------------------------------
    # WHOLE LESSON
    # ---------------------------------------------------------
    #
    # DARABILIM_BLOOM_ASSESSMENT_EXCLUDE_FINAL_V2
    #
    # Bloom деңгейлік тапсырмалары интерфейсте және
    # descriptor кестесінде сақталады, бірақ негізгі
    # сабақ бағасына қосылмайды.
    #
    # ТЕК:
    #     individual_level_work
    #
    # есептен шығарылады.
    #
    # individual_work және басқа тапсырмалар
    # бұрынғыдай бағаланады.
    # ---------------------------------------------------------

    totals = conn.execute(
        """
        SELECT
            COALESCE(
                SUM(
                    CASE
                        WHEN COALESCE(is_completed, 0) = 1
                        THEN 1
                        ELSE 0
                    END
                ),
                0
            ) AS earned,

            COUNT(*) AS maximum

        FROM lesson_assessment_descriptors

        WHERE
            assessment_id = ?
            AND COALESCE(task_key, '') <> 'individual_level_work'
        """,
        (
            assessment_id,
        )
    ).fetchone()


    lesson_earned = int(
        totals["earned"]
        or 0
    )


    lesson_max = int(
        totals["maximum"]
        or 0
    )


    conn.execute(
        """
        UPDATE lesson_assessments

        SET
            total_score = ?,
            max_score = ?,
            calculation_mode = 'descriptor_count',
            updated_at = CURRENT_TIMESTAMP

        WHERE id = ?
        """,
        (
            lesson_earned,
            lesson_max,
            assessment_id
        )
    )


    return (
        lesson_earned,
        lesson_max
    )





# =====================================================================
# DARABILIM_ASSESSMENT_CLASS_FILTER_FINAL_V1
# =====================================================================


def _da_assessment_grade_number_v1(
    value
):

    """
    Examples:

    4-сынып -> 4
    4 сынып -> 4
    4 А -> 4
    4 "А" -> 4
    4А -> 4
    """

    import re as _da_re

    text = str(
        value
        or ""
    ).strip()


    match = _da_re.match(
        r"^\s*(\d{1,2})",
        text
    )


    if not match:

        return None


    try:

        grade = int(
            match.group(1)
        )

    except (
        TypeError,
        ValueError
    ):

        return None


    if (
        grade < 1
        or
        grade > 12
    ):

        return None


    return grade


def _da_assessment_students_for_class_v1(
    conn,
    requested_class
):

    requested_class = str(
        requested_class
        or ""
    ).strip()


    grade = (
        _da_assessment_grade_number_v1(
            requested_class
        )
    )


    rows = conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name

        FROM students

        ORDER BY
            class_name,
            full_name
        """
    ).fetchall()


    # No class selected:
    # preserve old fallback.
    if not requested_class:

        return rows


    # First exact class match.
    exact = [
        row
        for row in rows
        if str(
            row["class_name"]
            or ""
        ).strip().casefold()
        ==
        requested_class.casefold()
    ]


    if exact:

        return exact


    # "4-сынып" means every class whose grade is 4:
    # 4 А, 4 Ә, 4 Б, ...
    if grade is not None:

        grade_rows = [
            row
            for row in rows
            if (
                _da_assessment_grade_number_v1(
                    row["class_name"]
                )
                ==
                grade
            )
        ]


        if grade_rows:

            return grade_rows


    return []





# =====================================================================
# DARABILIM_BASELINE_DAILY_ANALYTICS_FINAL_V1
#
# ENTRY DIAGNOSTIC
#        ↓
# OBJECTIVE BASELINE
#        ↓
# DAILY DESCRIPTOR ASSESSMENT
#        ↓
# GAP RECOVERY
#        ↓
# LEARNING PATH
#        ↓
# PROGRESS
# =====================================================================


def _da_resolve_objective_v1(
    conn,
    objective_id=None,
    objective_code=""
):

    try:

        if objective_id:

            objective_id = int(
                objective_id
            )

        else:

            objective_id = None

    except (
        TypeError,
        ValueError
    ):

        objective_id = None


    objective_code = str(
        objective_code
        or ""
    ).strip()


    if objective_id:

        row = conn.execute(
            """
            SELECT
                id,
                objective_code,
                objective_text,
                subject,
                lesson_topic,
                section

            FROM learning_objectives

            WHERE id = ?
            """,
            (
                objective_id,
            )
        ).fetchone()


        if row:

            return row


    if objective_code:

        row = conn.execute(
            """
            SELECT
                id,
                objective_code,
                objective_text,
                subject,
                lesson_topic,
                section

            FROM learning_objectives

            WHERE objective_code = ?

            ORDER BY id

            LIMIT 1
            """,
            (
                objective_code,
            )
        ).fetchone()


        if row:

            return row


    return None


def _da_latest_entry_baseline_v1(
    conn,
    student_id,
    objective_id
):

    row = conn.execute(
        """
        SELECT
            r.percentage,
            r.earned_points,
            r.max_points,
            r.mastery_status,
            r.attempt_id,
            a.completed_at,
            d.subject,
            d.title

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        WHERE
            a.student_id = ?
            AND r.learning_objective_id = ?
            AND a.status = 'completed'
            AND COALESCE(
                d.diagnostic_type,
                'entry'
            ) = 'entry'

        ORDER BY
            a.completed_at DESC,
            a.id DESC

        LIMIT 1
        """,
        (
            student_id,
            objective_id
        )
    ).fetchone()


    if row:

        return row


    # Compatibility:
    # older diagnostics created before diagnostic_type.
    return conn.execute(
        """
        SELECT
            r.percentage,
            r.earned_points,
            r.max_points,
            r.mastery_status,
            r.attempt_id,
            a.completed_at,
            d.subject,
            d.title

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        WHERE
            a.student_id = ?
            AND r.learning_objective_id = ?
            AND a.status = 'completed'

        ORDER BY
            a.completed_at ASC,
            a.id ASC

        LIMIT 1
        """,
        (
            student_id,
            objective_id
        )
    ).fetchone()


def _da_daily_objective_evidence_v1(
    conn,
    student_id,
    objective_id
):

    rows = conn.execute(
        """
        SELECT
            la.assessment_date,
            la.lesson_topic,
            la.subject,

            SUM(
                COALESCE(
                    lad.earned_point,
                    0
                )
            ) AS earned,

            SUM(
                COALESCE(
                    lad.max_point,
                    0
                )
            ) AS maximum,

            COUNT(*) AS descriptor_count

        FROM lesson_assessment_descriptors lad

        JOIN lesson_assessments la
            ON la.id = lad.assessment_id

        WHERE
            lad.student_id = ?
            AND lad.learning_objective_id = ?

        GROUP BY
            la.assessment_date,
            la.lesson_topic,
            la.subject

        HAVING
            SUM(
                COALESCE(
                    lad.max_point,
                    0
                )
            ) > 0

        ORDER BY
            la.assessment_date,
            la.lesson_topic
        """,
        (
            student_id,
            objective_id
        )
    ).fetchall()


    result = []


    for row in rows:

        earned = int(
            row["earned"]
            or 0
        )

        maximum = int(
            row["maximum"]
            or 0
        )


        percentage = (
            round(
                earned
                / maximum
                * 100,
                1
            )
            if maximum
            else 0
        )


        result.append({

            "date":
                row["assessment_date"],

            "lesson_topic":
                row["lesson_topic"]
                or "",

            "subject":
                row["subject"]
                or "",

            "earned":
                earned,

            "maximum":
                maximum,

            "percentage":
                percentage,

            "descriptor_count":
                int(
                    row["descriptor_count"]
                    or 0
                )
        })


    return result


def _da_gap_recovery_status_v1(
    baseline,
    evidence
):

    baseline = float(
        baseline
        or 0
    )


    if baseline >= 80:

        return (
            "Меңгерілген",
            "🟢"
        )


    if not evidence:

        return (
            "Олқылық анықталды",
            "🔴"
        )


    percentages = [
        float(
            item.get(
                "percentage",
                0
            )
            or 0
        )
        for item in evidence
    ]


    current = percentages[-1]


    recent = percentages[-3:]


    # Stable mastery:
    # at least 2 latest daily evidences >= 80.
    if (
        len(recent) >= 2
        and
        all(
            value >= 80
            for value in recent[-2:]
        )
    ):

        return (
            "Олқылық толықты",
            "🟢"
        )


    if current >= 80:

        return (
            "Айтарлықтай жақсарды",
            "🟡"
        )


    if current > baseline:

        return (
            "Толықтырылуда",
            "🟠"
        )


    return (
        "Олқылық сақталуда",
        "🔴"
    )


def _da_student_objective_analytics_v1(
    conn,
    student_id,
    objective_id
):

    objective = (
        _da_resolve_objective_v1(
            conn,
            objective_id,
            ""
        )
    )


    if not objective:

        return None


    baseline_row = (
        _da_latest_entry_baseline_v1(
            conn,
            student_id,
            objective["id"]
        )
    )


    if not baseline_row:

        return None


    baseline = round(
        float(
            baseline_row["percentage"]
            or 0
        ),
        1
    )


    evidence = (
        _da_daily_objective_evidence_v1(
            conn,
            student_id,
            objective["id"]
        )
    )


    if evidence:

        current = float(
            evidence[-1][
                "percentage"
            ]
        )

    else:

        current = baseline


    growth = round(
        current - baseline,
        1
    )


    status, icon = (
        _da_gap_recovery_status_v1(
            baseline,
            evidence
        )
    )


    completed_descriptors = sum(
        int(
            item["earned"]
            or 0
        )
        for item in evidence
    )


    total_descriptors = sum(
        int(
            item["maximum"]
            or 0
        )
        for item in evidence
    )


    return {

        "student_id":
            student_id,

        "learning_objective_id":
            objective["id"],

        "objective_code":
            objective["objective_code"]
            or "",

        "objective_text":
            objective["objective_text"]
            or "",

        "subject":
            (
                baseline_row["subject"]
                or objective["subject"]
                or ""
            ),

        "lesson_topic":
            objective["lesson_topic"]
            or "",

        "section":
            objective["section"]
            or "",

        "baseline_percentage":
            baseline,

        "baseline_earned":
            int(
                baseline_row["earned_points"]
                or 0
            ),

        "baseline_max":
            int(
                baseline_row["max_points"]
                or 0
            ),

        "baseline_date":
            baseline_row["completed_at"],

        "baseline_title":
            baseline_row["title"]
            or "Кіріс бақылау",

        "current_percentage":
            round(
                current,
                1
            ),

        "growth":
            growth,

        "status":
            status,

        "status_icon":
            icon,

        "daily_evidence_count":
            len(evidence),

        "completed_descriptors":
            completed_descriptors,

        "total_descriptors":
            total_descriptors,

        "timeline":
            evidence
    }


def _da_all_baseline_daily_analytics_v1(
    conn
):

    baseline_rows = conn.execute(
        """
        SELECT DISTINCT
            a.student_id,
            r.learning_objective_id

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        WHERE
            a.status = 'completed'
        """
    ).fetchall()


    results = []


    for row in baseline_rows:

        item = (
            _da_student_objective_analytics_v1(
                conn,
                row["student_id"],
                row["learning_objective_id"]
            )
        )


        if item:

            student = conn.execute(
                """
                SELECT
                    full_name,
                    class_name

                FROM students

                WHERE id = ?
                """,
                (
                    row["student_id"],
                )
            ).fetchone()


            if student:

                item["full_name"] = (
                    student["full_name"]
                    or ""
                )

                item["class_name"] = (
                    student["class_name"]
                    or ""
                )


            results.append(
                item
            )


    return results


@app.route(
    "/api/baseline-daily-progress",
    methods=["GET"]
)
def da_baseline_daily_progress_api_v1():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    student_id = (
        request.args.get(
            "student_id",
            ""
        )
        .strip()
    )


    subject = (
        request.args.get(
            "subject",
            ""
        )
        .strip()
    )


    conn = get_db()


    results = (
        _da_all_baseline_daily_analytics_v1(
            conn
        )
    )


    conn.close()


    if student_id:

        try:

            wanted_student = int(
                student_id
            )

        except ValueError:

            return jsonify({
                "ok": False,
                "error": "student_id қате."
            }), 400


        results = [
            item
            for item in results
            if int(
                item["student_id"]
            ) == wanted_student
        ]


    if subject:

        results = [
            item
            for item in results
            if str(
                item["subject"]
                or ""
            ) == subject
        ]


    return jsonify({

        "ok":
            True,

        "results":
            results
    })


@app.route(
    "/api/baseline-daily-gaps",
    methods=["GET"]
)
def da_baseline_daily_gaps_api_v1():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    conn = get_db()


    results = (
        _da_all_baseline_daily_analytics_v1(
            conn
        )
    )


    conn.close()


    gaps = [
        item
        for item in results
        if float(
            item["baseline_percentage"]
            or 0
        ) < 80
    ]


    gaps.sort(
        key=lambda item: (
            float(
                item[
                    "current_percentage"
                ]
                or 0
            ),
            item["full_name"],
            item["objective_code"]
        )
    )


    return jsonify({
        "ok": True,
        "gaps": gaps
    })





# =====================================================================
# DARABILIM_OBJECTIVE_CODE_DEDUP_ANALYTICS_FINAL_V1
#
# ANALYTICS UNIQUE KEY:
#
#     student_id + objective_code
#
# learning_objectives кестесіндегі duplicate ID-лер өшірілмейді.
# Аналитикада бір code бір ғана оқу мақсаты ретінде қаралады.
# =====================================================================


def _da_objective_ids_by_code_v2(
    conn,
    objective_code
):

    objective_code = str(
        objective_code
        or ""
    ).strip()


    if not objective_code:

        return []


    rows = conn.execute(
        """
        SELECT id

        FROM learning_objectives

        WHERE
            TRIM(
                COALESCE(
                    objective_code,
                    ''
                )
            ) = ?

        ORDER BY id
        """,
        (
            objective_code,
        )
    ).fetchall()


    return [
        int(row["id"])
        for row in rows
    ]


def _da_canonical_objective_by_code_v2(
    conn,
    objective_code
):

    objective_code = str(
        objective_code
        or ""
    ).strip()


    if not objective_code:

        return None


    return conn.execute(
        """
        SELECT
            id,
            objective_code,
            objective_text,
            subject,
            lesson_topic,
            section

        FROM learning_objectives

        WHERE
            TRIM(
                COALESCE(
                    objective_code,
                    ''
                )
            ) = ?

        ORDER BY id

        LIMIT 1
        """,
        (
            objective_code,
        )
    ).fetchone()


def _da_latest_entry_baseline_by_code_v2(
    conn,
    student_id,
    objective_code
):

    ids = (
        _da_objective_ids_by_code_v2(
            conn,
            objective_code
        )
    )


    if not ids:

        return None


    placeholders = ",".join(
        "?"
        for _ in ids
    )


    params = [
        student_id,
        *ids
    ]


    # --------------------------------------------------------
    # Prefer explicit ENTRY diagnostic.
    # --------------------------------------------------------

    sql = f"""
        SELECT
            r.percentage,
            r.earned_points,
            r.max_points,
            r.mastery_status,
            r.attempt_id,

            a.completed_at,

            d.subject,
            d.title

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        WHERE
            a.student_id = ?

            AND
            r.learning_objective_id IN (
                {placeholders}
            )

            AND
            a.status = 'completed'

            AND
            COALESCE(
                d.diagnostic_type,
                'entry'
            ) = 'entry'

        ORDER BY
            a.completed_at DESC,
            a.id DESC,
            r.id DESC

        LIMIT 1
    """


    row = conn.execute(
        sql,
        params
    ).fetchone()


    if row:

        return row


    # --------------------------------------------------------
    # Compatibility with old diagnostics:
    # earliest completed result becomes baseline.
    # --------------------------------------------------------

    sql = f"""
        SELECT
            r.percentage,
            r.earned_points,
            r.max_points,
            r.mastery_status,
            r.attempt_id,

            a.completed_at,

            d.subject,
            d.title

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN diagnostics d
            ON d.id = a.diagnostic_id

        WHERE
            a.student_id = ?

            AND
            r.learning_objective_id IN (
                {placeholders}
            )

            AND
            a.status = 'completed'

        ORDER BY
            a.completed_at ASC,
            a.id ASC,
            r.id ASC

        LIMIT 1
    """


    return conn.execute(
        sql,
        params
    ).fetchone()


def _da_daily_evidence_by_code_v2(
    conn,
    student_id,
    objective_code
):

    ids = (
        _da_objective_ids_by_code_v2(
            conn,
            objective_code
        )
    )


    if not ids:

        return []


    placeholders = ",".join(
        "?"
        for _ in ids
    )


    params = [
        student_id,
        objective_code,
        *ids
    ]


    sql = f"""
        SELECT
            la.assessment_date,
            la.lesson_topic,
            la.subject,

            SUM(
                COALESCE(
                    lad.earned_point,
                    0
                )
            ) AS earned,

            SUM(
                COALESCE(
                    lad.max_point,
                    0
                )
            ) AS maximum,

            COUNT(*) AS descriptor_count

        FROM lesson_assessment_descriptors lad

        JOIN lesson_assessments la
            ON la.id = lad.assessment_id

        WHERE
            lad.student_id = ?

            AND
            (
                TRIM(
                    COALESCE(
                        lad.objective_code,
                        ''
                    )
                ) = ?

                OR

                lad.learning_objective_id IN (
                    {placeholders}
                )
            )

        GROUP BY
            la.assessment_date,
            la.lesson_topic,
            la.subject

        HAVING
            SUM(
                COALESCE(
                    lad.max_point,
                    0
                )
            ) > 0

        ORDER BY
            la.assessment_date,
            la.lesson_topic
    """


    rows = conn.execute(
        sql,
        params
    ).fetchall()


    result = []


    for row in rows:

        earned = int(
            row["earned"]
            or 0
        )


        maximum = int(
            row["maximum"]
            or 0
        )


        percentage = (
            round(
                earned
                / maximum
                * 100,
                1
            )
            if maximum
            else 0.0
        )


        result.append({

            "date":
                row["assessment_date"],

            "lesson_topic":
                row["lesson_topic"]
                or "",

            "subject":
                row["subject"]
                or "",

            "earned":
                earned,

            "maximum":
                maximum,

            "percentage":
                percentage,

            "descriptor_count":
                int(
                    row["descriptor_count"]
                    or 0
                )
        })


    return result


def _da_student_objective_code_analytics_v2(
    conn,
    student_id,
    objective_code
):

    objective = (
        _da_canonical_objective_by_code_v2(
            conn,
            objective_code
        )
    )


    if not objective:

        return None


    baseline_row = (
        _da_latest_entry_baseline_by_code_v2(
            conn,
            student_id,
            objective_code
        )
    )


    if not baseline_row:

        return None


    baseline = round(
        float(
            baseline_row["percentage"]
            or 0
        ),
        1
    )


    evidence = (
        _da_daily_evidence_by_code_v2(
            conn,
            student_id,
            objective_code
        )
    )


    if evidence:

        current = float(
            evidence[-1][
                "percentage"
            ]
        )

    else:

        current = baseline


    growth = round(
        current - baseline,
        1
    )


    status, icon = (
        _da_gap_recovery_status_v1(
            baseline,
            evidence
        )
    )


    completed_descriptors = sum(
        int(
            item["earned"]
            or 0
        )
        for item in evidence
    )


    total_descriptors = sum(
        int(
            item["maximum"]
            or 0
        )
        for item in evidence
    )


    return {

        "student_id":
            int(student_id),

        "learning_objective_id":
            int(
                objective["id"]
            ),

        "objective_code":
            objective["objective_code"]
            or "",

        "objective_text":
            objective["objective_text"]
            or "",

        "subject":
            (
                baseline_row["subject"]
                or objective["subject"]
                or ""
            ),

        "lesson_topic":
            objective["lesson_topic"]
            or "",

        "section":
            objective["section"]
            or "",

        "baseline_percentage":
            baseline,

        "baseline_earned":
            int(
                baseline_row["earned_points"]
                or 0
            ),

        "baseline_max":
            int(
                baseline_row["max_points"]
                or 0
            ),

        "baseline_date":
            baseline_row["completed_at"],

        "baseline_title":
            baseline_row["title"]
            or "Кіріс бақылау",

        "current_percentage":
            round(
                current,
                1
            ),

        "growth":
            growth,

        "status":
            status,

        "status_icon":
            icon,

        "daily_evidence_count":
            len(evidence),

        "completed_descriptors":
            completed_descriptors,

        "total_descriptors":
            total_descriptors,

        "timeline":
            evidence
    }


def _da_all_baseline_daily_analytics_v2(
    conn
):

    # --------------------------------------------------------
    # UNIQUE:
    #
    # student_id + objective_code
    # --------------------------------------------------------

    rows = conn.execute(
        """
        SELECT DISTINCT
            a.student_id,
            TRIM(
                COALESCE(
                    lo.objective_code,
                    ''
                )
            ) AS objective_code

        FROM objective_results r

        JOIN diagnostic_attempts a
            ON a.id = r.attempt_id

        JOIN learning_objectives lo
            ON lo.id =
               r.learning_objective_id

        WHERE
            a.status = 'completed'

            AND

            TRIM(
                COALESCE(
                    lo.objective_code,
                    ''
                )
            ) <> ''

        ORDER BY
            a.student_id,
            objective_code
        """
    ).fetchall()


    results = []


    for row in rows:

        item = (
            _da_student_objective_code_analytics_v2(
                conn,
                row["student_id"],
                row["objective_code"]
            )
        )


        if not item:
            continue


        student = conn.execute(
            """
            SELECT
                full_name,
                class_name

            FROM students

            WHERE id = ?
            """,
            (
                row["student_id"],
            )
        ).fetchone()


        if student:

            item["full_name"] = (
                student["full_name"]
                or ""
            )

            item["class_name"] = (
                student["class_name"]
                or ""
            )


        results.append(
            item
        )


    return results


# ------------------------------------------------------------
# Replace V1 analytics source with V2 source.
#
# Existing APIs call this name:
# _da_all_baseline_daily_analytics_v1
#
# Therefore APIs do not need to be duplicated.
# ------------------------------------------------------------

_da_all_baseline_daily_analytics_v1 = (
    _da_all_baseline_daily_analytics_v2
)





# =====================================================================
# DARABILIM_ANALYTICS_RESULTS_RESET_FINAL_V1
#
# mode = daily
#
# Deletes:
#   lesson_assessment_descriptors
#   lesson_assessment_items
#   lesson_assessments
#
# Keeps:
#   diagnostics
#   diagnostic questions
#   diagnostic attempts/results
#   students
#   objectives
#   smart lessons
#   materials
#   teacher tasks/descriptors
#
#
# mode = all
#
# Deletes:
#   daily assessment results
#   personal_tasks
#   objective_results
#   diagnostic_answers
#   diagnostic_attempts
#
# Keeps:
#   diagnostics
#   diagnostic_questions
#   diagnostic_objectives
#   students
#   learning_objectives
#   smart lessons
#   materials
#   teacher tasks/descriptors
# =====================================================================


@app.route(
    "/api/analytics-results/reset",
    methods=["POST"]
)
def da_analytics_results_reset_final_v1():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    data = (
        request.get_json(
            silent=True
        )
        or {}
    )


    mode = str(
        data.get(
            "mode",
            ""
        )
        or ""
    ).strip().lower()


    confirmation = str(
        data.get(
            "confirmation",
            ""
        )
        or ""
    ).strip()


    if mode not in (
        "daily",
        "all"
    ):

        return jsonify({
            "ok": False,
            "error": "Тазалау режимі қате."
        }), 400


    required_confirmation = (
        "DAILY"
        if mode == "daily"
        else "ALL"
    )


    if confirmation != required_confirmation:

        return jsonify({
            "ok": False,
            "error": "Тазалау расталмады."
        }), 400


    conn = get_db()


    try:

        before_counts = {}


        for table in (
            "lesson_assessment_descriptors",
            "lesson_assessment_items",
            "lesson_assessments",
            "personal_tasks",
            "objective_results",
            "diagnostic_answers",
            "diagnostic_attempts"
        ):

            before_counts[table] = (
                conn.execute(
                    f"""
                    SELECT COUNT(*)
                    FROM {table}
                    """
                ).fetchone()[0]
            )


        # -----------------------------------------------------
        # DAILY RESULTS
        #
        # Child -> parent order.
        # -----------------------------------------------------

        conn.execute(
            """
            DELETE FROM
                lesson_assessment_descriptors
            """
        )


        conn.execute(
            """
            DELETE FROM
                lesson_assessment_items
            """
        )


        conn.execute(
            """
            DELETE FROM
                lesson_assessments
            """
        )


        # -----------------------------------------------------
        # ALL RESULTS
        # -----------------------------------------------------

        if mode == "all":

            conn.execute(
                """
                DELETE FROM
                    personal_tasks
                """
            )


            conn.execute(
                """
                DELETE FROM
                    objective_results
                """
            )


            conn.execute(
                """
                DELETE FROM
                    diagnostic_answers
                """
            )


            conn.execute(
                """
                DELETE FROM
                    diagnostic_attempts
                """
            )


        conn.commit()


        after_counts = {}


        for table in (
            "lesson_assessment_descriptors",
            "lesson_assessment_items",
            "lesson_assessments",
            "personal_tasks",
            "objective_results",
            "diagnostic_answers",
            "diagnostic_attempts"
        ):

            after_counts[table] = (
                conn.execute(
                    f"""
                    SELECT COUNT(*)
                    FROM {table}
                    """
                ).fetchone()[0]
            )


        integrity = (
            conn.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )


        if integrity != "ok":

            raise RuntimeError(
                "Database integrity error."
            )


        return jsonify({

            "ok":
                True,

            "mode":
                mode,

            "before":
                before_counts,

            "after":
                after_counts,

            "message":
                (
                    "Күнделікті бағалау нәтижелері тазаланды."
                    if mode == "daily"
                    else
                    "Кіріс бақылау және күнделікті нәтижелер тазаланды."
                )
        })


    except Exception as error:

        conn.rollback()


        return jsonify({

            "ok":
                False,

            "error":
                str(error)

        }), 500


    finally:

        conn.close()





# =====================================================================
# DARABILIM_PARENT_MESSENGER_FINAL_V2
#
# students.parent_phone -> WhatsApp
# students.telegram     -> Telegram
#
# Parent-facing report does NOT show objective_code.
# Weak skills are taken from uncompleted descriptors.
# =====================================================================


def _da_parent_students_v2(
    conn
):

    rows = conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name,
            parent_name,
            parent_phone,
            telegram

        FROM students

        ORDER BY
            class_name,
            full_name
        """
    ).fetchall()


    return [
        {
            "id":
                row["id"],

            "full_name":
                row["full_name"]
                or "",

            "class_name":
                row["class_name"]
                or "",

            "parent_name":
                row["parent_name"]
                or "",

            "whatsapp":
                row["parent_phone"]
                or "",

            "telegram":
                row["telegram"]
                or ""
        }

        for row in rows
    ]


def _da_parent_student_v2(
    conn,
    student_id
):

    return conn.execute(
        """
        SELECT
            id,
            full_name,
            class_name,
            parent_name,
            parent_phone,
            telegram

        FROM students

        WHERE id = ?

        LIMIT 1
        """,
        (
            student_id,
        )
    ).fetchone()


def _da_parent_skill_evidence_v2(
    conn,
    student_id
):

    # ---------------------------------------------------------
    # Latest state of each actual descriptor.
    #
    # Parent sees descriptor TEXT, not objective code.
    # ---------------------------------------------------------

    rows = conn.execute(
        """
        SELECT
            lad.id,
            lad.descriptor_text,
            lad.is_completed,
            lad.earned_point,
            lad.max_point,
            lad.objective_code,

            la.assessment_date,
            la.lesson_topic

        FROM lesson_assessment_descriptors lad

        JOIN lesson_assessments la
            ON la.id = lad.assessment_id

        WHERE
            lad.student_id = ?

        ORDER BY
            la.assessment_date DESC,
            lad.id DESC
        """,
        (
            student_id,
        )
    ).fetchall()


    latest = {}


    for row in rows:

        text = str(
            row["descriptor_text"]
            or ""
        ).strip()


        if not text:

            continue


        # Same skill text should appear once.
        key = text.casefold()


        if key not in latest:

            latest[key] = row


    weak = []
    strong = []


    for row in latest.values():

        text = str(
            row["descriptor_text"]
            or ""
        ).strip()


        completed = (
            int(
                row["is_completed"]
                or 0
            )
            == 1
        )


        if completed:

            strong.append(
                text
            )

        else:

            weak.append(
                text
            )


    return (
        weak[:10],
        strong[:10]
    )


def _da_parent_weak_objective_text_v2(
    analytics
):

    # ---------------------------------------------------------
    # Fallback only when no uncompleted daily descriptor exists.
    #
    # Still no objective code is shown to parent.
    # ---------------------------------------------------------

    result = []


    for item in analytics:

        baseline = float(
            item.get(
                "baseline_percentage",
                0
            )
            or 0
        )


        status = str(
            item.get(
                "status",
                ""
            )
            or ""
        )


        if baseline >= 80:

            continue


        if status == "Олқылық толықты":

            continue


        text = str(
            item.get(
                "objective_text",
                ""
            )
            or ""
        ).strip()


        if (
            text
            and
            text not in result
        ):

            result.append(
                text
            )


    return result[:10]


def _da_parent_report_v2(
    conn,
    student_id
):

    student = (
        _da_parent_student_v2(
            conn,
            student_id
        )
    )


    if not student:

        return None


    all_analytics = (
        _da_all_baseline_daily_analytics_v1(
            conn
        )
    )


    analytics = [
        item

        for item in all_analytics

        if int(
            item.get(
                "student_id",
                0
            )
            or 0
        )
        ==
        int(
            student_id
        )
    ]


    weak_skills, strong_skills = (
        _da_parent_skill_evidence_v2(
            conn,
            student_id
        )
    )


    if not weak_skills:

        weak_skills = (
            _da_parent_weak_objective_text_v2(
                analytics
            )
        )


    latest_assessment = conn.execute(
        """
        SELECT
            total_score,
            max_score,
            assessment_date,
            lesson_topic

        FROM lesson_assessments

        WHERE student_id = ?

        ORDER BY
            assessment_date DESC,
            id DESC

        LIMIT 1
        """,
        (
            student_id,
        )
    ).fetchone()


    baseline_values = [
        float(
            item.get(
                "baseline_percentage",
                0
            )
            or 0
        )

        for item in analytics
    ]


    current_values = [
        float(
            item.get(
                "current_percentage",
                0
            )
            or 0
        )

        for item in analytics
    ]


    baseline_average = (
        round(
            sum(
                baseline_values
            )
            /
            len(
                baseline_values
            ),
            1
        )

        if baseline_values

        else None
    )


    current_average = (
        round(
            sum(
                current_values
            )
            /
            len(
                current_values
            ),
            1
        )

        if current_values

        else None
    )


    growth = (
        round(
            current_average
            -
            baseline_average,
            1
        )

        if (
            baseline_average
            is not None
            and
            current_average
            is not None
        )

        else None
    )


    statuses = [
        str(
            item.get(
                "status",
                ""
            )
            or ""
        )

        for item in analytics
    ]


    if (
        "Олқылық сақталуда"
        in statuses
        or
        "Олқылық анықталды"
        in statuses
    ):

        overall_status = (
            "🔴 Қосымша қолдау қажет"
        )


    elif (
        "Толықтырылуда"
        in statuses
    ):

        overall_status = (
            "🟠 Білім олқылығы толықтырылуда"
        )


    elif (
        "Айтарлықтай жақсарды"
        in statuses
    ):

        overall_status = (
            "🟡 Айтарлықтай ілгерілеу бар"
        )


    elif analytics:

        overall_status = (
            "🟢 Оқу нәтижесі тұрақты"
        )


    else:

        overall_status = (
            "Нәтиже әлі жеткіліксіз"
        )


    return {

        "student": {

            "id":
                student["id"],

            "full_name":
                student["full_name"]
                or "",

            "class_name":
                student["class_name"]
                or "",

            "parent_name":
                student["parent_name"]
                or "",

            # REAL DATABASE COLUMN:
            "whatsapp":
                student["parent_phone"]
                or "",

            "telegram":
                student["telegram"]
                or ""
        },


        "baseline_average":
            baseline_average,


        "current_average":
            current_average,


        "growth":
            growth,


        "overall_status":
            overall_status,


        "weak_skills":
            weak_skills,


        "strong_skills":
            strong_skills,


        "latest_assessment":
            (
                {

                    "score":
                        latest_assessment[
                            "total_score"
                        ],

                    "max_score":
                        latest_assessment[
                            "max_score"
                        ],

                    "date":
                        latest_assessment[
                            "assessment_date"
                        ],

                    "lesson_topic":
                        latest_assessment[
                            "lesson_topic"
                        ]
                        or ""
                }

                if latest_assessment

                else None
            )
    }


@app.route(
    "/api/parent-report/students-v2",
    methods=["GET"]
)
def da_parent_report_students_api_v2():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    conn = get_db()


    students = (
        _da_parent_students_v2(
            conn
        )
    )


    conn.close()


    return jsonify({

        "ok":
            True,

        "students":
            students
    })


@app.route(
    "/api/parent-report-v2/<int:student_id>",
    methods=["GET"]
)
def da_parent_report_api_v2(
    student_id
):

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    conn = get_db()


    report = (
        _da_parent_report_v2(
            conn,
            student_id
        )
    )


    conn.close()


    if not report:

        return jsonify({
            "ok": False,
            "error": "Оқушы табылмады."
        }), 404


    return jsonify({

        "ok":
            True,

        "report":
            report
    })





# =====================================================================
# DARABILIM_MY_LESSONS_SAVE_BACKEND_FINAL_V4
# =====================================================================

@app.route(
    "/api/smart-lessons/quick-save-final",
    methods=["POST"]
)
def darabilim_my_lessons_quick_save_final_v4():

    if not logged_in():

        return jsonify({
            "ok": False,
            "error": "Unauthorized"
        }), 401


    import json


    data = request.get_json(
        silent=True
    ) or {}


    plan = data.get("plan")


    if not isinstance(
        plan,
        dict
    ):

        plan = {}


    # ---------------------------------------------------------
    # FIELD RESOLVER
    #
    # Frontend sends explicit fields AND plan.
    # Explicit value wins.
    # ---------------------------------------------------------

    def first_text(*values):

        for value in values:

            if value is None:
                continue

            text = str(
                value
            ).strip()

            if text:
                return text

        return ""


    subject = first_text(
        data.get("subject"),
        plan.get("subject")
    )


    class_name = first_text(
        data.get("class_name"),
        data.get("className"),
        plan.get("class_name"),
        plan.get("className")
    )


    lesson_topic = first_text(
        data.get("lesson_topic"),
        data.get("lessonTopic"),
        data.get("topic"),
        plan.get("lesson_topic"),
        plan.get("lessonTopic"),
        plan.get("topic")
    )


    lesson_goal = first_text(
        data.get("lesson_goal"),
        data.get("lessonGoal"),
        plan.get("lesson_goal"),
        plan.get("lessonGoal")
    )


    # ---------------------------------------------------------
    # VALIDATE MASTER IDENTITY
    # ---------------------------------------------------------

    if not subject:

        return jsonify({
            "ok": False,
            "error": "Пән анықталмады."
        }), 400


    if not class_name:

        return jsonify({
            "ok": False,
            "error": "Сынып анықталмады."
        }), 400


    if not lesson_topic:

        return jsonify({
            "ok": False,
            "error": "Сабақ тақырыбы анықталмады."
        }), 400


    # ---------------------------------------------------------
    # DURATION
    # ---------------------------------------------------------

    raw_duration = (
        data.get("duration")
        or
        plan.get("duration")
        or
        45
    )


    try:

        duration = int(
            raw_duration
        )

    except Exception:

        duration = 45


    if duration <= 0:
        duration = 45


    # ---------------------------------------------------------
    # OBJECTIVES
    # ---------------------------------------------------------

    objective_codes = (
        data.get("objective_codes")
        or
        plan.get("objective_codes")
        or
        []
    )


    objective_texts = (
        data.get("objective_texts")
        or
        plan.get("objective_texts")
        or
        []
    )


    if not isinstance(
        objective_codes,
        list
    ):

        objective_codes = []


    if not isinstance(
        objective_texts,
        list
    ):

        objective_texts = []


    # Some lesson plans store objectives as object list.
    objectives = (
        data.get("objectives")
        or
        plan.get("objectives")
        or
        []
    )


    if isinstance(
        objectives,
        list
    ):

        if not objective_codes:

            for item in objectives:

                if not isinstance(
                    item,
                    dict
                ):
                    continue

                code = (
                    item.get("objective_code")
                    or
                    item.get("code")
                )

                if code:

                    objective_codes.append(
                        str(code).strip()
                    )


        if not objective_texts:

            for item in objectives:

                if not isinstance(
                    item,
                    dict
                ):
                    continue

                text = (
                    item.get("objective_text")
                    or
                    item.get("text")
                )

                if text:

                    objective_texts.append(
                        str(text).strip()
                    )


    # ---------------------------------------------------------
    # ENSURE PLAN CONTAINS MASTER CONTEXT
    # ---------------------------------------------------------

    plan["subject"] = subject
    plan["class_name"] = class_name
    plan["lesson_topic"] = lesson_topic
    plan["duration"] = duration

    plan["objective_codes"] = objective_codes
    plan["objective_texts"] = objective_texts

    if lesson_goal:

        plan["lesson_goal"] = lesson_goal


    # ---------------------------------------------------------
    # INSERT
    #
    # IMPORTANT:
    # New lesson ALWAYS gets a new row.
    # No UPDATE.
    # No REPLACE.
    # ---------------------------------------------------------

    conn = get_db()


    try:

        cursor = conn.execute(
            """
            INSERT INTO smart_lessons (
                subject,
                class_name,
                lesson_topic,
                duration,
                objective_codes,
                objective_texts,
                lesson_goal,
                plan_json,
                status,
                created_at,
                updated_at
            )

            VALUES (
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                'saved',
                CURRENT_TIMESTAMP,
                CURRENT_TIMESTAMP
            )
            """,
            (
                subject,
                class_name,
                lesson_topic,
                duration,

                json.dumps(
                    objective_codes,
                    ensure_ascii=False
                ),

                json.dumps(
                    objective_texts,
                    ensure_ascii=False
                ),

                lesson_goal,

                json.dumps(
                    plan,
                    ensure_ascii=False
                )
            )
        )


        lesson_id = cursor.lastrowid


        if not lesson_id:

            raise RuntimeError(
                "Жаңа lesson_id жасалмады."
            )


        conn.commit()


        # -----------------------------------------------------
        # VERIFY USING SAME CONNECTION
        # -----------------------------------------------------

        saved = conn.execute(
            """
            SELECT
                id,
                subject,
                class_name,
                lesson_topic,
                duration,
                status

            FROM smart_lessons

            WHERE id = ?

            LIMIT 1
            """,
            (
                lesson_id,
            )
        ).fetchone()


        if not saved:

            raise RuntimeError(
                "INSERT орындалды, бірақ сабақ қайта табылмады."
            )


        total = conn.execute(
            """
            SELECT COUNT(*)
            FROM smart_lessons
            """
        ).fetchone()[0]


        return jsonify({

            "ok":
                True,

            "saved":
                True,

            "lesson_id":
                lesson_id,

            "smart_lessons_count":
                total,

            "subject":
                subject,

            "class_name":
                class_name,

            "lesson_topic":
                lesson_topic,

            "redirect_url":
                "/smart-lessons/saved",

            "message":
                "Сабақ Менің сабақтарым бөліміне сақталды."
        })


    except Exception as error:

        conn.rollback()


        return jsonify({

            "ok":
                False,

            "error":
                str(error)

        }), 500


    finally:

        conn.close()



if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=int(
            os.environ.get(
                "PORT",
                5000
            )
        ),
        debug=False
    )
